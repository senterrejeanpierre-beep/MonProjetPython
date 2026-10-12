"""Liens du PR et pages choisies par l'utilisateur dans le cahier de son chantier."""
from copy import deepcopy
import hashlib
import html
import json
import os
from pathlib import Path, PurePosixPath
import posixpath
import re
import shutil
import sys
import tempfile
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string

import liens_horizon

CONFIG = 'cahier_pr.json'
S = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
P = 'http://schemas.openxmlformats.org/package/2006/relationships'
LABEL = 'Cahier technique · choisir les pages'


def empreinte(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def lire_config(chantier):
    p = Path(chantier) / CONFIG
    return json.loads(p.read_text(encoding='utf-8')) if p.exists() else {'articles': {}}


def chemin_relatif(chantier, relatif):
    """Un document partagé reste à l'intérieur de son chantier, même après déplacement."""
    path = PurePosixPath(str(relatif))
    if path.is_absolute() or '..' in path.parts or '\\' in str(relatif) or ':' in str(relatif):
        raise ValueError('Le cahier doit se trouver dans le dossier du chantier.')
    root = Path(chantier).resolve()
    fichier = (root / Path(*path.parts)).resolve()
    if not fichier.is_relative_to(root):
        raise ValueError('Le cahier doit se trouver dans le dossier du chantier.')
    return fichier


def trouver_cahier(chantier):
    chantier = Path(chantier)
    data = chantier / 'data'

    config = lire_config(chantier)

    if config.get('fichier'):
        # Un choix explicite manquant ne doit pas ouvrir un autre cahier.
        return chemin_relatif(chantier, config['fichier'])

    # La racine reste lisible pour reprendre les anciens chantiers.
    for dossier in (data, chantier):
        if not dossier.is_dir():
            continue
        fichiers = {p.name.casefold(): p for p in dossier.iterdir()
                    if p.is_file() and not p.is_symlink()}
        for nom in ('cahier_technique.pdf', 'cahier technique.pdf', 'cahier-technique.pdf'):
            if nom in fichiers:
                return fichiers[nom]
        candidats = [p for p in fichiers.values() if p.suffix.casefold() == '.pdf'
                     and 'cahier' in p.stem.casefold() and 'technique' in p.stem.casefold()]
        if candidats:
            return candidats[0] if len(candidats) == 1 else None
    return None


def ranger_cahier(chantier, pdf=None):
    """Copie le cahier dans data, sans écraser de document ni perdre les choix."""
    chantier = Path(chantier)
    pdf = Path(pdf) if pdf is not None else trouver_cahier(chantier)
    if pdf is None or not pdf.is_file():
        return pdf
    if pdf.is_symlink():
        raise ValueError('Choisissez le cahier PDF sans lien symbolique.')
    config_path = chantier / CONFIG
    avant = config_path.read_bytes() if config_path.exists() else None
    config = json.loads(avant) if avant is not None else {'articles': {}}
    nombre_pages(pdf)
    data = chemin_relatif(chantier, 'data')
    data.mkdir(exist_ok=True)
    signature = empreinte(pdf)
    cible = data / pdf.name
    numero = 2
    if cible.is_symlink():
        raise ValueError('Le fichier de destination ne doit pas être un lien symbolique.')
    while cible.exists() and empreinte(cible) != signature:
        cible = data / f'{pdf.stem}_{numero}{pdf.suffix}'
        numero += 1
        if cible.is_symlink():
            raise ValueError('Le fichier de destination ne doit pas être un lien symbolique.')
    if not cible.exists():
        cree = False
        try:
            with cible.open('xb') as sortie:
                cree = True
                with pdf.open('rb') as source:
                    shutil.copyfileobj(source, sortie)
        except FileExistsError:
            raise ValueError('Le cahier a été ajouté sur un autre poste. Recommencez.')
        except Exception:
            if cree:
                cible.unlink(missing_ok=True)
            raise
    if empreinte(cible) != signature or empreinte(pdf) != signature:
        raise ValueError('Le cahier a changé pendant sa copie. Recommencez.')
    relatif = cible.relative_to(chantier.resolve()).as_posix()
    for choix in config.get('articles', {}).values():
        if choix.get('sha256') == signature:
            choix['fichier'] = relatif
    config['fichier'] = relatif
    if config != (json.loads(avant) if avant is not None else {'articles': {}}):
        _ecrire_config(chantier, config, avant)
    return cible


def _ecrire_config(chantier, config, avant):
    fd, tmp = tempfile.mkstemp(prefix='.cahier_pr_', suffix='.json', dir=chantier)
    os.close(fd)
    try:
        Path(tmp).write_text(json.dumps(config, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        cible = Path(chantier) / CONFIG
        if (cible.read_bytes() if cible.exists() else None) != avant:
            raise ValueError('Les choix du cahier ont changé sur un autre poste. Recommencez.')
        os.replace(tmp, cible)
    finally:
        Path(tmp).unlink(missing_ok=True)


def lire_pages(texte, nombre_pages):
    """Numéros physiques du PDF, à partir de 1 ; intervalles et pages séparées."""
    texte = str(texte).strip().replace('–', '-').replace('—', '-')
    if not texte:
        raise ValueError('Indiquez les pages du PDF, par exemple 30-32, 35.')
    resultat = []
    for morceau in re.split(r'[,;]', texte):
        match = re.fullmatch(r'\s*(\d+)\s*(?:-\s*(\d+)\s*)?', morceau)
        if not match:
            raise ValueError('Pages incorrectes. Exemple : 30-32, 35.')
        debut, fin = int(match[1]), int(match[2] or match[1])
        if not 1 <= debut <= fin <= nombre_pages:
            raise ValueError(f'Indiquez des pages comprises entre 1 et {nombre_pages}, dans le bon ordre.')
        for page in range(debut, fin + 1):
            if page not in resultat:
                resultat.append(page)
    return resultat


def _lecteur(pdf):
    from pypdf import PdfReader
    reader = PdfReader(pdf)
    if reader.is_encrypted and not reader.decrypt(''):
        raise ValueError('Ce cahier PDF est protégé par un mot de passe.')
    return reader


def nombre_pages(pdf):
    with Path(pdf).open('rb') as stream:
        return len(_lecteur(stream).pages)

def trouver_pages_article(pdf, article):
    """Recherche automatiquement les pages contenant exactement le numéro d'article."""
    article = str(article).strip()
    if not article:
        return []

    reader = _lecteur(pdf)
    pages = []

    motif = re.compile(
        rf'(?<![0-9A-Za-z.]){re.escape(article)}(?![0-9A-Za-z.])',
        re.IGNORECASE
    )

    for numero, page in enumerate(reader.pages, start=1):
        texte = page.extract_text() or ''
        if motif.search(texte):
            pages.append(numero)

    return pages

def _pages_automatiques(pdf, articles):
    """Repère les rubriques du texte, hors sommaire et renvois en cours de phrase."""
    signature = empreinte(pdf)
    ressources = Path(getattr(sys, '_MEIPASS', Path(__file__).resolve().parent)) / 'ressources' / 'cahiers'
    correspondances = ressources / (signature + '.json')
    if correspondances.is_file():
        confirme = json.loads(correspondances.read_text(encoding='utf-8'))
        if confirme.get('pdf_sha256') == signature:
            return {code: str(pages[0]) if pages[0] == pages[1] else f'{pages[0]}-{pages[1]}'
                    for code, pages in confirme['articles'].items() if code in articles}
    with Path(pdf).open('rb') as stream:
        reader = _lecteur(stream)
        titres = []
        for numero, page in enumerate(reader.pages, 1):
            texte = page.extract_text() or ''
            if re.search(r'table des mati[eè]res|sommaire', texte, re.I):
                continue
            lignes = texte.splitlines()
            # Les pages suivantes du sommaire n'en répètent pas toujours le titre.
            entrees_sommaire = sum(bool(re.search(r'[.·…]{3,}|\S\s{3,}\d+\s*$', l)) for l in lignes)
            if entrees_sommaire >= 3:
                continue
            for ligne in lignes:
                if re.search(r'[.·…]{3,}', ligne):
                    continue
                match = re.match(r'^\s*(\d+(?:[.]\d+[A-Za-z]*)*)\s+([A-Za-zÀ-ÿ].*)$', ligne)
                if match and ('.' in match[1] or match[1] in articles):
                    titres.append((match[1], numero))
        resultat = {}
        for article in articles:
            code = article
            occurrences = []
            while True:
                occurrences = [i for i, (titre, _) in enumerate(titres) if titre == code]
                if occurrences or '.' not in code:
                    break
                code = code.rsplit('.', 1)[0]
                # Une rubrique générale (un seul nombre) n'est pas un poste.
                if '.' not in code:
                    break
            if len(occurrences) != 1:
                continue
            index = occurrences[0]
            debut = titres[index][1]
            fin = len(reader.pages)
            for suivant, page in titres[index + 1:]:
                if suivant != code and not suivant.startswith(code + '.'):
                    # Deux articles peuvent commencer sur la même page.
                    fin = page
                    break
            resultat[article] = str(debut) if fin == debut else f'{debut}-{fin}'
        return resultat


def pages_article_automatiques(pdf, article, articles_pr):
    if str(article) not in articles_pr:
        return ''
    return _pages_automatiques(pdf, [str(article)]).get(str(article), '')


def choix_article(chantier, article, pdf):
    choix = lire_config(chantier).get('articles', {}).get(str(article), {})
    if not pdf or not Path(pdf).is_file():
        return '', bool(choix)
    relatif = Path(pdf).resolve().relative_to(Path(chantier).resolve()).as_posix()
    actuel = choix.get('fichier') == relatif and choix.get('sha256') == empreinte(pdf)
    return (choix.get('saisie', '') if actuel else ''), bool(choix) and not actuel


def preparer_pages(chantier, article, pdf, saisie):
    """Extrait les pages actuelles puis mémorise le choix ; aucune écriture dans Excel."""
    from pypdf import PdfWriter
    chantier, pdf = Path(chantier), Path(pdf)
    relatif = pdf.resolve().relative_to(chantier.resolve()).as_posix()
    pdf = chemin_relatif(chantier, relatif)
    signature = empreinte(pdf)
    config_path = chantier / CONFIG
    avant = config_path.read_bytes() if config_path.exists() else None
    config = json.loads(avant) if avant is not None else {'articles': {}}
    with pdf.open('rb') as stream:
        reader = _lecteur(stream)
        pages = lire_pages(saisie, len(reader.pages))
        writer = PdfWriter()
        for page in pages:
            writer.add_page(reader.pages[page - 1])
        writer.add_metadata({'/Title': f'Article {article} — pages {saisie} du cahier'})
        cache = chantier / 'data' / '.cahier_pages'
        cache.mkdir(parents=True, exist_ok=True)
        cle = hashlib.sha256((signature + json.dumps(pages)).encode()).hexdigest()
        sortie = cache / (cle + '.pdf')
        fd, tmp = tempfile.mkstemp(suffix='.pdf', dir=cache)
        os.close(fd)
        try:
            writer.write(tmp)
            if empreinte(pdf) != signature:
                raise ValueError('Le cahier a changé pendant la lecture. Recommencez.')
            os.replace(tmp, sortie)
        finally:
            Path(tmp).unlink(missing_ok=True)
    config['fichier'] = relatif
    config.setdefault('articles', {})[str(article)] = {
        'fichier': relatif, 'sha256': signature, 'saisie': str(saisie).strip(), 'pages': pages}
    fd, tmp = tempfile.mkstemp(prefix='.cahier_pr_', suffix='.json', dir=chantier)
    os.close(fd)
    try:
        Path(tmp).write_text(json.dumps(config, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        actuel = config_path.read_bytes() if config_path.exists() else None
        if actuel != avant:
            raise ValueError('Les choix du cahier ont changé sur un autre poste. Recommencez.')
        os.replace(tmp, config_path)
    finally:
        Path(tmp).unlink(missing_ok=True)
    return sortie


def lire_postes(pr):
    wb = load_workbook(pr, read_only=True)
    try:
        if 'Chiffrage' not in wb:
            return []
        rows = list(wb['Chiffrage'].iter_rows(max_col=3, values_only=True))
        postes = []
        for i, values in enumerate(rows):
            article = str(values[1] or '').strip()
            if not re.fullmatch(r'\d+(?:[.][0-9A-Za-z]+)*', article):
                continue
            # Un code accompagné d'une désignation est un poste, pas un titre Excel.
            if not isinstance(values[2], str) or not values[2].strip() or values[2].startswith('='):
                continue
            cellule = f'B{i+1}'
            for j in range(i + 1, min(i + 9, len(rows))):
                if rows[j][1] and re.fullmatch(r'\d+(?:[.][0-9A-Za-z]+)*', str(rows[j][1]).strip()):
                    break
                texte = rows[j][2]
                if texte is None or str(texte).startswith(('Cahier technique', 'Cahier spécial')):
                    cellule = f'C{j+1}'
                    break
            postes.append({'article': article, 'designation': values[2], 'cellule': cellule, 'retour': f'B{i+1}'})
        return postes
    finally:
        wb.close()

def _ajouter_xml(xml, balise, ajout):
    """Insère dans un élément OOXML sans réécrire le reste du document."""
    fermeture = re.search(r'</(?:[\w]+:)?' + balise + r'\s*>', xml)
    if fermeture:
        return xml[:fermeture.start()] + ajout + xml[fermeture.start():]
    vide = re.search(r'<((?:[\w]+:)?' + balise + r')\b([^>]*?)/>', xml)
    if vide:
        return xml[:vide.start()] + '<' + vide[1] + vide[2] + '>' + ajout + '</' + vide[1] + '>' + xml[vide.end():]
    raise ValueError(f'Élément Excel introuvable : {balise}')


def _cellule_texte(xml, ref, texte):
    pattern = rf'<(?:\w+:)?c\b[^>]*\br="{re.escape(ref)}"[^>]*?(?:/>|>.*?</(?:\w+:)?c>)'
    match = re.search(pattern, xml, re.S)
    cellule = f'<c xmlns="{S}" r="{ref}" t="inlineStr"><is><t>{html.escape(texte)}</t></is></c>'
    if match:
        style = re.search(r'\bs="(\d+)"', match[0])
        if style:
            cellule = cellule.replace(' t="inlineStr"', f' s="{style[1]}" t="inlineStr"')
        return xml[:match.start()] + cellule + xml[match.end():]
    ligne = re.search(r'\d+', ref)[0]
    row = re.search(rf'<(?:\w+:)?row\b[^>]*\br="{ligne}"[^>]*?(?:/>|>.*?</(?:\w+:)?row>)', xml, re.S)
    if row:
        contenu = row[0]
        col = column_index_from_string(re.match(r'[A-Z]+', ref)[0])
        suivant = next((m for m in re.finditer(r'<(?:\w+:)?c\b[^>]*\br="([A-Z]+)\d+"', contenu)
                        if column_index_from_string(m[1]) > col), None)
        contenu = (contenu[:suivant.start()] + cellule + contenu[suivant.start():]
                   if suivant else _ajouter_xml(contenu, 'row', cellule))
        return xml[:row.start()] + contenu + xml[row.end():]
    ajout = f'<row xmlns="{S}" r="{ligne}">{cellule}</row>'
    suivant = next((m for m in re.finditer(r'<(?:\w+:)?row\b[^>]*\br="(\d+)"', xml)
                    if int(m[1]) > int(ligne)), None)
    return (xml[:suivant.start()] + ajout + xml[suivant.start():]
            if suivant else _ajouter_xml(xml, 'sheetData', ajout))


def _mettre_lien(xml, ref, destination, identifiant=None):
    attrs = (f'xmlns:r="{R}" r:id="{identifiant}"' if identifiant
             else f'location="{html.escape(destination, quote=True)}"')
    lien = f'<hyperlink xmlns="{S}" ref="{ref}" {attrs}/>'
    pattern = rf'<(?:\w+:)?hyperlink\b[^>]*\bref="{re.escape(ref)}"[^>]*?(?:/>|>.*?</(?:\w+:)?hyperlink>)'
    if re.search(pattern, xml):
        return re.sub(pattern, lambda _: lien, xml, count=1)
    if re.search(r'<(?:\w+:)?hyperlinks\b', xml):
        return _ajouter_xml(xml, 'hyperlinks', lien)
    # Ordre des éléments du format SpreadsheetML.
    pos = re.search(r'<(?:\w+:)?(?:printOptions|pageMargins|pageSetup|headerFooter|'
                    r'rowBreaks|colBreaks|customProperties|cellWatches|ignoredErrors|'
                    r'smartTags|drawing|legacyDrawing|picture|oleObjects|controls|'
                    r'webPublishItems|tableParts|extLst)\b|</(?:\w+:)?worksheet>', xml).start()
    return xml[:pos] + f'<hyperlinks xmlns="{S}">{lien}</hyperlinks>' + xml[pos:]


def _chemin_feuille(src, nom):
    wb = ET.fromstring(src.read('xl/workbook.xml'))
    feuille = next(s for s in wb.find(f'{{{S}}}sheets') if s.get('name') == nom)
    relations = ET.fromstring(src.read('xl/_rels/workbook.xml.rels'))
    cible = next(r.get('Target') for r in relations if r.get('Id') == feuille.get(f'{{{R}}}id'))
    return cible.lstrip('/') if cible.startswith('/') else posixpath.normpath('xl/' + cible)


def _creer_onglets(src, postes, selections, pdf):
    """Construit uniquement les nouveaux onglets ; aucun réenregistrement du PR."""
    import pymupdf
    D = 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing'
    A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
    wb = src.read('xl/workbook.xml').decode('utf-8')
    rel = src.read('xl/_rels/workbook.xml.rels').decode('utf-8')
    types = src.read('[Content_Types].xml').decode('utf-8')
    numero = max(int(s.get('sheetId')) for s in ET.fromstring(wb).find(f'{{{S}}}sheets')) + 1
    nouveaux = {}
    noms = set(src.namelist())
    images = {}
    ids = {r.get('Id') for r in ET.fromstring(rel)}
    with pymupdf.open(str(pdf)) as document:
        for poste in postes:
            article = poste['article']
            if article not in selections:
                continue
            while (f'xl/worksheets/sheet{numero}.xml' in noms or
                   f'xl/drawings/cahier{numero}.xml' in noms or f'rIdCahier{numero}' in ids):
                numero += 1
            titre = 'CSC ' + article
            if len(titre) > 31:
                continue
            rid = f'rIdCahier{numero}'
            feuille = f'xl/worksheets/sheet{numero}.xml'
            dessin = f'xl/drawings/cahier{numero}.xml'
            wb = _ajouter_xml(wb, 'sheets', f'<sheet xmlns="{S}" xmlns:r="{R}" name="{titre}" sheetId="{numero}" r:id="{rid}"/>')
            rel = _ajouter_xml(rel, 'Relationships', f'<Relationship xmlns="{P}" Id="{rid}" Type="{R}/worksheet" Target="worksheets/sheet{numero}.xml"/>')
            for fichier, type_partie in ((feuille, 'worksheet'), (dessin, 'drawing')):
                types = _ajouter_xml(types, 'Types', f'<Override PartName="/{fichier}" ContentType="application/vnd.openxmlformats-officedocument.{"spreadsheetml.worksheet" if type_partie == "worksheet" else "drawing"}+xml"/>')
            pages = selections[article]
            texte = lambda ref, valeur: f'<c r="{ref}" t="inlineStr"><is><t>{html.escape(valeur)}</t></is></c>'
            rows = [f'<row r="1">{texte("A1", titre)}</row>',
                    f'<row r="2">{texte("A2", "← Retour au poste dans le PR")}</row>',
                    f'<row r="3">{texte("A3", "Cahier technique — pages intégrées au classeur")}</row>']
            ancres, liens_images, ligne = [], [], 5
            for index, page in enumerate(pages, 1):
                if page not in images:
                    pix = document[page - 1].get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False)
                    png = pix.tobytes('png')
                    nom = f'cahier_{hashlib.sha256(png).hexdigest()}.png'
                    member = 'xl/media/' + nom
                    if member not in noms:
                        nouveaux[member] = png
                    images[page] = (nom, pix.width, pix.height)
                nom, largeur, hauteur = images[page]
                h = round(hauteur * 1000 / largeur)
                rows.append(f'<row r="{ligne}">{texte("A"+str(ligne), f"Page {page} du cahier technique")}</row>')
                ancres.append(f'<xdr:oneCellAnchor><xdr:from><xdr:col>0</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>{ligne}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:ext cx="9525000" cy="{h*9525}"/><xdr:pic><xdr:nvPicPr><xdr:cNvPr id="{index}" name="Cahier page {page}"/><xdr:cNvPicPr/></xdr:nvPicPr><xdr:blipFill><a:blip r:embed="rId{index}"/><a:stretch><a:fillRect/></a:stretch></xdr:blipFill><xdr:spPr><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr></xdr:pic><xdr:clientData/></xdr:oneCellAnchor>')
                liens_images.append(f'<Relationship Id="rId{index}" Type="{R}/image" Target="../media/{nom}"/>')
                ligne += int(h / 20) + 4
            retour = html.escape(f"'Chiffrage'!{poste['retour']}", quote=True)
            nouveaux[feuille] = (f'<worksheet xmlns="{S}" xmlns:r="{R}"><dimension ref="A1:P{ligne}"/><sheetViews><sheetView showGridLines="0" zoomScale="85" workbookViewId="0"><pane ySplit="3" topLeftCell="A4" activePane="bottomLeft" state="frozen"/></sheetView></sheetViews><sheetFormatPr defaultRowHeight="15"/><cols><col min="1" max="16" width="12" customWidth="1"/></cols><sheetData>'+''.join(rows)+f'</sheetData><hyperlinks><hyperlink ref="A2" location="{retour}"/></hyperlinks><pageMargins left="0.2" right="0.2" top="0.2" bottom="0.2" header="0" footer="0"/><drawing r:id="rId1"/></worksheet>').encode()
            nouveaux[f'xl/worksheets/_rels/sheet{numero}.xml.rels'] = f'<Relationships xmlns="{P}"><Relationship Id="rId1" Type="{R}/drawing" Target="../drawings/cahier{numero}.xml"/></Relationships>'.encode()
            nouveaux[dessin] = (f'<xdr:wsDr xmlns:xdr="{D}" xmlns:a="{A}" xmlns:r="{R}">'+''.join(ancres)+'</xdr:wsDr>').encode()
            nouveaux[f'xl/drawings/_rels/cahier{numero}.xml.rels'] = (f'<Relationships xmlns="{P}">'+''.join(liens_images)+'</Relationships>').encode()
            numero += 1
    if not any(n.get('Extension') == 'png' for n in ET.fromstring(types)):
        types = _ajouter_xml(types, 'Types', '<Default Extension="png" ContentType="image/png"/>')
    nouveaux.update({'xl/workbook.xml': wb.encode(), 'xl/_rels/workbook.xml.rels': rel.encode(), '[Content_Types].xml': types.encode()})
    return nouveaux


def integrer_cahier_dans_pr(chantier, pr, sauvegarder=None):
    """Point d'entrée conservé : même intégration protégée que l'ouverture du PR."""
    return assurer_liens(chantier, pr, sauvegarder)


def retirer_cahier_du_modele(pr):
    """Sur la copie d'un modèle seulement, retire les pages du chantier précédent."""
    pr = Path(pr)
    with ZipFile(pr) as src:
        wb = src.read('xl/workbook.xml').decode('utf-8')
        feuilles = ET.fromstring(wb).find(f'{{{S}}}sheets')
        csc = [s for s in feuilles if s.get('name', '').startswith('CSC ')]
        if not csc:
            return
        ids = {s.get(f'{{{R}}}id') for s in csc}
        rel = src.read('xl/_rels/workbook.xml.rels').decode('utf-8')
        supprimer = {_chemin_feuille(src, s.get('name')) for s in csc}
        # Retire les pièces de ces feuilles uniquement si aucune autre pièce ne les utilise.
        liens = {}
        for name in src.namelist():
            if not name.endswith('.rels') or '/_rels/' not in name:
                continue
            dossier, fichier = name.rsplit('/_rels/', 1)
            parent = dossier + '/' + fichier[:-5]
            liens[parent] = {r.get('Target').lstrip('/') if r.get('Target').startswith('/')
                             else posixpath.normpath(dossier + '/' + r.get('Target'))
                             for r in ET.fromstring(src.read(name)) if r.get('TargetMode') != 'External'}
        while True:
            candidates = {c for p in supprimer for c in liens.get(p, set())}
            partages = {c for p, valeurs in liens.items() if p not in supprimer and p != 'xl/workbook.xml' for c in valeurs}
            nouvelles = candidates - partages - supprimer
            if not nouvelles:
                break
            supprimer.update(nouvelles)
        supprimer.update(posixpath.join(posixpath.dirname(n), '_rels', posixpath.basename(n) + '.rels')
                          for n in list(supprimer))
        for s in csc:
            wb = re.sub(r'<(?:\w+:)?sheet\b[^>]*\bname="' + re.escape(s.get('name')) + r'"[^>]*/>', '', wb)
        for rid in ids:
            rel = re.sub(r'<(?:\w+:)?Relationship\b[^>]*\bId="' + re.escape(rid) + r'"[^>]*/>', '', rel)
        types = src.read('[Content_Types].xml').decode('utf-8')
        for name in supprimer:
            types = re.sub(r'<(?:\w+:)?Override\b[^>]*\bPartName="/' + re.escape(name) + r'"[^>]*/>', '', types)
        member = _chemin_feuille(src, 'Chiffrage')
        contenu = src.read(member).decode('utf-8')
        noms = {s.get('name') for s in csc}
        for poste in lire_postes(pr):
            if 'CSC ' + poste['article'] in noms and poste['cellule'].startswith('C'):
                contenu = _cellule_texte(contenu, poste['cellule'], LABEL)
        changes = {'xl/workbook.xml': wb.encode(), 'xl/_rels/workbook.xml.rels': rel.encode(),
                   '[Content_Types].xml': types.encode(), member: contenu.encode()}
        fd, tmp = tempfile.mkstemp(prefix='.cahier_modele_', suffix=pr.suffix, dir=pr.parent)
        os.close(fd)
        try:
            with ZipFile(tmp, 'w') as dst:
                for info in src.infolist():
                    if info.filename not in supprimer:
                        dst.writestr(deepcopy(info), changes.get(info.filename, src.read(info.filename)))
        except Exception:
            Path(tmp).unlink(missing_ok=True)
            raise
    try:
        os.replace(tmp, pr)
    finally:
        Path(tmp).unlink(missing_ok=True)


def assurer_liens(chantier, pr, sauvegarder=None):
    """Préserve le PR et ses CSC ; ajoute les pages manquantes et leurs liens."""
    pr, chantier = Path(pr), Path(chantier)
    signature = empreinte(pr)
    postes = lire_postes(pr)
    if not postes:
        return 0
    if pr.with_name('~$' + pr.name).exists():
        raise PermissionError('Enregistrez et fermez le PR pour activer les liens du cahier.')
    pdf = trouver_cahier(chantier)
    signature_pdf = empreinte(pdf) if pdf and pdf.is_file() else None
    with ZipFile(pr) as src:
        feuilles = {s.get('name') for s in ET.fromstring(src.read('xl/workbook.xml')).find(f'{{{S}}}sheets')}
        manquants = [p for p in postes if 'CSC ' + p['article'] not in feuilles]
        selections = {}
        if manquants and signature_pdf:
            config = lire_config(chantier)
            automatiques = _pages_automatiques(pdf, [p['article'] for p in postes])
            total_pages = nombre_pages(pdf)
            for poste in manquants:
                code = poste['article']
                choix = config.get('articles', {}).get(code, {})
                if choix:
                    saisie, change = choix_article(chantier, code, pdf)
                    if change:
                        continue
                else:
                    saisie = automatiques.get(code, '')
                if saisie and len('CSC ' + code) <= 31:
                    selections[code] = lire_pages(saisie, total_pages)
        changes = _creer_onglets(src, manquants, selections, pdf) if selections else {}
        feuilles.update('CSC ' + code for code in selections)
        member = _chemin_feuille(src, 'Chiffrage')
        rel_member = posixpath.join(posixpath.dirname(member), '_rels', posixpath.basename(member) + '.rels')
        original = src.read(member).decode('utf-8')
        contenu = original
        rel_original = src.read(rel_member) if rel_member in src.namelist() else None
        rels = ET.fromstring(rel_original) if rel_original else ET.Element(f'{{{P}}}Relationships')
        sheet_root = ET.fromstring(contenu)
        changed, cibles_dessins = 0, {}
        for index, poste in enumerate(postes):
            ref, article = poste['cellule'], poste['article']
            interne = 'CSC ' + article in feuilles
            destination = f"'CSC {article}'!A1" if interne else liens_horizon.lien_cahier(chantier, article)
            old = sheet_root.find(f'{{{S}}}hyperlinks/{{{S}}}hyperlink[@ref="{ref}"]')
            old_id = old.get(f'{{{R}}}id') if old is not None else None
            old_rel = next((r for r in rels if r.get('Id') == old_id), None)
            cible_dessin = '#' + destination if interne else destination
            if old_rel is not None:
                cibles_dessins[old_rel.get('Target')] = cible_dessin
            if not interne:
                cibles_dessins[f"#'CSC {article}'!A1"] = destination
            identique = (old is not None and old.get('location') == destination and not old_id) if interne else (old_rel is not None and old_rel.get('Target') == destination)
            cellule = sheet_root.find(f'.//{{{S}}}c[@r="{ref}"]')
            vide = cellule is None or not any(cellule.find(f'{{{S}}}{tag}') is not None for tag in ('v', 'is', 'f'))
            if identique and not vide:
                continue
            changed += 1
            if not identique:
                identifiant = None
                if not interne:
                    identifiant = f'rIdCahier{index + 1}'
                    while any(r.get('Id') == identifiant for r in rels):
                        identifiant += 'x'
                    ET.SubElement(rels, f'{{{P}}}Relationship', Id=identifiant, Type=R + '/hyperlink', Target=destination, TargetMode='External')
                contenu = _mettre_lien(contenu, ref, destination, identifiant)
            # Le numéro d'article n'est jamais remplacé si aucun espace n'existe.
            if ref.startswith('C') and (vide or article in selections):
                pages = selections.get(article)
                label = 'Cahier technique · p. ' + ', '.join(map(str, pages)) if pages else LABEL
                contenu = _cellule_texte(contenu, ref, label)
        if contenu != original:
            changes[member] = contenu.encode('utf-8')
        if rel_original is None or ET.tostring(rels) != ET.tostring(ET.fromstring(rel_original)):
            changes[rel_member] = ET.tostring(rels, encoding='utf-8', xml_declaration=True)
        # Les zones de texte superposées doivent pointer vers le même onglet.
        for name in src.namelist():
            if name.startswith('xl/drawings/_rels/') and name.endswith('.rels') and cibles_dessins:
                xml = src.read(name).decode('utf-8')
                def changer_cible(match):
                    cible = html.unescape(match[1])
                    if cible not in cibles_dessins or cibles_dessins[cible] == cible:
                        return match[0]
                    return 'Target="' + html.escape(cibles_dessins[cible], quote=True) + '"'
                modifie = re.sub(r'\bTarget="([^"]*)"', changer_cible, xml)
                if modifie != xml:
                    changes[name] = modifie.encode('utf-8')
        if not changes:
            return 0
        for name, contenu_xml in changes.items():
            if name.endswith(('.xml', '.rels')):
                ET.fromstring(contenu_xml)
        fd, tmp = tempfile.mkstemp(prefix='.cahier_pr_', suffix=pr.suffix, dir=pr.parent)
        os.close(fd)
        try:
            with ZipFile(tmp, 'w', compression=ZIP_DEFLATED) as dst:
                for info in src.infolist():
                    dst.writestr(deepcopy(info), changes.get(info.filename, src.read(info.filename)))
                for name, value in changes.items():
                    if name not in src.namelist():
                        dst.writestr(name, value)
            def verifier_sources():
                if pr.with_name('~$' + pr.name).exists():
                    raise PermissionError('Le PR a été rouvert. Fermez-le puis recommencez.')
                if empreinte(pr) != signature or (signature_pdf and empreinte(pdf) != signature_pdf):
                    raise ValueError('Le PR ou le cahier a changé pendant la préparation. Recommencez.')
            verifier_sources()
            if sauvegarder:
                sauvegarder(pr)
            verifier_sources()
        except Exception:
            Path(tmp).unlink(missing_ok=True)
            raise
    try:
        os.replace(tmp, pr)
    finally:
        Path(tmp).unlink(missing_ok=True)
    return changed

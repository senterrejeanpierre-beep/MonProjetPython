"""Import explicite des jours du PR dans la copie du planning de soumission."""

from html import escape
import os
from pathlib import Path
import re
import shutil
import tempfile
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from planning_soumission_contenu import changements_soumission


S = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
P = 'http://schemas.openxmlformats.org/package/2006/relationships'
NOM_FEUILLE = 'Postes PR'


def lire_jours_pr(pr):
    wb = load_workbook(pr, read_only=True, data_only=True)
    try:
        if 'Chiffrage' not in wb:
            raise ValueError('Feuille Chiffrage absente du PR.')
        lignes = list(wb['Chiffrage'].iter_rows(max_col=7, values_only=True))
        postes = []
        for i, ligne in enumerate(lignes):
            if len(ligne) < 2 or not str(ligne[1] or '').startswith('Article'):
                continue
            if i + 15 >= len(lignes):
                raise ValueError(f'Bloc PR incomplet à la ligne {i + 1}.')
            article, enonce = lignes[i + 1][1:3]
            if article in (None, ''):
                continue
            jours = lignes[i + 15][6]
            if jours is not None and (isinstance(jours, bool) or not isinstance(jours, (int, float)) or jours < 0):
                raise ValueError(f'Jours G{i + 16} illisibles pour {article}.')
            postes.append((str(article).strip(), str(enonce or '').strip(), jours))
        if not postes:
            raise ValueError('Aucun article reconnu dans le PR.')
        return postes
    finally:
        wb.close()


def _cellule_texte(ref, valeur):
    return f'<c r="{ref}" t="inlineStr"><is><t>{escape(str(valeur))}</t></is></c>'


def _cellule_nombre(ref, valeur):
    return f'<c r="{ref}"><v>{valeur}</v></c>'


def _feuille(postes, debuts):
    lignes = [('<row r="1">' + ''.join(_cellule_texte(f'{col}1', titre)
              for col, titre in zip('ABCDE', ('Numéro article', 'Énoncé du poste', 'Jours PR (G)',
                                               'Jour de début (saisie)', 'Jour de fin'))) + '</row>')]
    for numero, (article, enonce, jours) in enumerate(postes, 2):
        contenu = _cellule_texte(f'A{numero}', article) + _cellule_texte(f'B{numero}', enonce)
        if jours is not None:
            contenu += _cellule_nombre(f'C{numero}', jours)
        debut = debuts.get(article)
        if isinstance(debut, (int, float)) and not isinstance(debut, bool):
            contenu += _cellule_nombre(f'D{numero}', debut)
        elif isinstance(debut, str) and debut.startswith('='):
            contenu += f'<c r="D{numero}"><f>{escape(debut[1:])}</f></c>'
        contenu += f'<c r="E{numero}"><f>IF(OR(C{numero}="",D{numero}=""),"",D{numero}+C{numero}-1)</f></c>'
        lignes.append(f'<row r="{numero}">{contenu}</row>')
    fin = len(postes) + 1
    lignes.append(f'<row r="{fin + 2}">' + _cellule_texte(f'B{fin + 2}', 'Durée du chantier : dernier jour de fin')
                  + f'<c r="E{fin + 2}"><f>IF(COUNT(E2:E{fin})=0,"",MAX(E2:E{fin}))</f></c></row>')
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<worksheet xmlns="{S}"><dimension ref="A1:E{fin + 2}"/>'
            '<sheetViews><sheetView workbookViewId="0"/></sheetViews>'
            '<sheetFormatPr defaultRowHeight="15"/>'
            '<cols><col min="1" max="1" width="21" customWidth="1"/>'
            '<col min="2" max="2" width="80" customWidth="1"/>'
            '<col min="3" max="5" width="23" customWidth="1"/></cols>'
            '<sheetData>' + ''.join(lignes) + '</sheetData>'
            f'<dataValidations count="1"><dataValidation type="whole" operator="greaterThanOrEqual" '
            f'allowBlank="1" showErrorMessage="1" sqref="D2:D{fin}"><formula1>1</formula1>'
            '</dataValidation></dataValidations></worksheet>').encode('utf-8')


def importer_jours_pr(pr, planning, sauvegarder=None, *, nom='', client='', adresse=''):
    """Reporte les articles du PR dans le planning et conserve les débuts par article."""
    pr, planning = Path(pr), Path(planning)
    postes = lire_jours_pr(pr)
    if planning.with_name('~$' + planning.name).exists():
        raise PermissionError('Enregistrez et fermez le planning de soumission dans Excel avant l’import.')
    debuts = {}
    debuts_principaux = {}
    wb = load_workbook(planning, read_only=True, data_only=False)
    try:
        if NOM_FEUILLE in wb:
            for row in wb[NOM_FEUILLE].iter_rows(min_row=2, max_col=4, values_only=True):
                if row[0] not in (None, '') and row[3] not in (None, ''):
                    debuts[str(row[0]).strip()] = row[3]
        if 'PlanningProjet' in wb:
            for row in wb['PlanningProjet'].iter_rows(min_row=23, max_row=62,
                                                       min_col=2, max_col=5, values_only=True):
                if row[0] not in (None, '') and row[3] not in (None, ''):
                    debuts_principaux[str(row[0]).strip()] = row[3]
    finally:
        wb.close()
    nouveau = _feuille(postes, debuts)
    with ZipFile(planning) as archive:
        classeur = archive.read('xl/workbook.xml').decode('utf-8')
        relations = archive.read('xl/_rels/workbook.xml.rels').decode('utf-8')
        types = archive.read('[Content_Types].xml').decode('utf-8')
        feuilles = ET.fromstring(classeur).find(f'{{{S}}}sheets')
        existante = next((s for s in feuilles if s.get('name') == NOM_FEUILLE), None)
        changements = changements_soumission(archive, nom, client, adresse, postes,
                                             debuts_principaux)
        if existante is not None:
            rid = existante.get(f'{{{R}}}id')
            cible = next(r.get('Target') for r in ET.fromstring(relations) if r.get('Id') == rid)
            membre = cible.lstrip('/') if cible.startswith('/') else 'xl/' + cible
        else:
            ids = [int(s.get('sheetId')) for s in feuilles]
            rid_n = max(int(m.group(1)) for m in re.finditer(r'Id="rId(\d+)"', relations)) + 1
            numero = max((int(m.group(1)) for n in archive.namelist()
                          if (m := re.fullmatch(r'xl/worksheets/sheet(\d+)\.xml', n))), default=0) + 1
            rid, membre = f'rId{rid_n}', f'xl/worksheets/sheet{numero}.xml'
            entete_classeur = re.search(r'<workbook\b[^>]*>', classeur).group()
            if 'xmlns:r=' not in entete_classeur:
                classeur = re.sub(r'<workbook\b', f'<workbook xmlns:r="{R}"', classeur, count=1)
            classeur = classeur.replace('</sheets>',
                f'<sheet name="{NOM_FEUILLE}" sheetId="{max(ids) + 1}" r:id="{rid}"/></sheets>', 1)
            relations = relations.replace('</Relationships>',
                f'<Relationship Id="{rid}" Type="{R}/worksheet" Target="worksheets/sheet{numero}.xml"/></Relationships>', 1)
            types = types.replace('</Types>',
                f'<Override PartName="/{membre}" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>', 1)
            changements.update({'xl/workbook.xml': classeur.encode('utf-8'),
                                'xl/_rels/workbook.xml.rels': relations.encode('utf-8'),
                                '[Content_Types].xml': types.encode('utf-8')})
        vue = re.search(r'<workbookView\b[^>]*/>', classeur)
        if vue is None:
            raise ValueError('Vue du classeur de planning introuvable.')
        nouvelle_vue = re.sub(r'\sactiveTab="[^"]*"', '', vue.group())
        nouvelle_vue = nouvelle_vue[:-2] + ' activeTab="0"/>'
        classeur = classeur.replace(vue.group(), nouvelle_vue, 1)
        changements['xl/workbook.xml'] = classeur.encode('utf-8')
        for nom in archive.namelist():
            if nom == membre or not re.fullmatch(r'xl/worksheets/sheet\d+\.xml', nom):
                continue
            original = changements.get(nom, archive.read(nom))
            nettoye = re.sub(rb'(<sheetView\b[^>]*?)\s+tabSelected="1"', rb'\1', original)
            if nom == 'xl/worksheets/sheet1.xml':
                nettoye = re.sub(rb'(<sheetView\b[^>]*?)\s*/>', rb'\1 tabSelected="1"/>', nettoye, count=1)
            if nettoye != original:
                changements[nom] = nettoye
        if existante is None or archive.read(membre) != nouveau:
            changements[membre] = nouveau
        noms_existants = set(archive.namelist())
        if all(nom in noms_existants and archive.read(nom) == donnees
               for nom, donnees in changements.items()):
            return len(postes), sum(j is None for _, _, j in postes), False
        if sauvegarder:
            sauvegarder(planning)
        fd, temporaire = tempfile.mkstemp(prefix='.import_jours_pr_', suffix='.xlsx', dir=planning.parent)
        os.close(fd)
        try:
            with ZipFile(temporaire, 'w') as sortie:
                for info in archive.infolist():
                    sortie.writestr(info, changements.pop(info.filename, archive.read(info.filename)))
                for nom, donnees in changements.items():
                    sortie.writestr(nom, donnees)
            shutil.copymode(planning, temporaire)
        except Exception:
            Path(temporaire).unlink(missing_ok=True)
            raise
    try:
        os.replace(temporaire, planning)
    finally:
        Path(temporaire).unlink(missing_ok=True)
    return len(postes), sum(j is None for _, _, j in postes), True

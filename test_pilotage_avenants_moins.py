import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from openpyxl import Workbook

from main import (
    _chemin_fichier_chantier,
    _lire_revision_globale_pilotage,
    _lire_synthese_avenants_pilotage,
    _lire_avenants_moins_publics,
    _montant_diminution_avenant,
    _reporter_quantites_cloture,
    _soumission_avant_cloture,
    _presentation_ecarts_pilotage,
)


def nombre_pilotage(val):
    if val in (None, "", "-", "—"):
        return 0
    try:
        texte_nombre = str(val)
        texte_nombre = texte_nombre.replace("€", "")
        texte_nombre = texte_nombre.replace("−", "-")
        texte_nombre = texte_nombre.replace(" ", "")
        texte_nombre = texte_nombre.replace("\xa0", "")
        texte_nombre = texte_nombre.replace("\u202f", "")
        texte_nombre = texte_nombre.replace(",", ".")
        return float(texte_nombre)
    except:
        return 0


class TestPilotageAvenantsMoins(unittest.TestCase):
    def test_moins_publics_lus_dans_fichier_du_chantier(self):
        with TemporaryDirectory() as tmp:
            dossier = Path(tmp) / "Chantier déplacé"
            dossier.mkdir()
            wb = Workbook()
            ws = wb.active
            ws["B62"] = "Montant cumulé à reporter dans état"
            ws["E62"] = -1250.5
            wb.save(dossier / "Avenants.xlsx")
            wb.close()
            self.assertEqual(_lire_avenants_moins_publics(dossier / "Etat.xlsx"), 1250.5)
            wb = Workbook()
            ws = wb.active
            ws["B62"], ws["E62"] = "Montant cumulé à reporter dans état", 0
            wb.save(dossier / "Avenants.xlsx")
            wb.close()
            self.assertEqual(_lire_avenants_moins_publics(dossier / "Etat.xlsx"), 0)
            wb = Workbook()
            wb.active["B62"] = "Montant cumulé à reporter dans état"
            wb.save(dossier / "Avenants.xlsx")
            wb.close()
            with self.assertRaisesRegex(ValueError, "résultat absent"):
                _lire_avenants_moins_publics(dossier / "Etat.xlsx")

    def test_ecarts_explicites_avec_deux_references(self):
        libelle, montant, marche, initial = _presentation_ecarts_pilotage(
            171360.60, 170240.1058, 181177.84025)
        self.assertEqual(libelle, "Production au-delà du marché total")
        self.assertAlmostEqual(montant, 10937.73445)
        self.assertEqual(marche, "+10,937.73 €")
        self.assertEqual(initial, "+9,817.24 € au-dessus")
        libelle, montant, marche, initial = _presentation_ecarts_pilotage(100, 120, 80)
        self.assertEqual(libelle, "Reste à facturer sur le marché total")
        self.assertEqual(montant, 40)
        self.assertEqual(marche, "-40.00 €")
        self.assertEqual(initial, "-20.00 € en dessous")
        self.assertEqual(_presentation_ecarts_pilotage(100, 100, 100)[2:],
                         ("0.00 €", "0.00 € : égal à la soumission"))

    def test_cloture_preserve_soumission_et_synthese(self):
        ws = Workbook().active
        ws["R24"] = 12
        ws["Q24"] = 3
        ws["Q25"] = "Total soumission hors TVA"
        ws["R25"] = 171360.60
        ws["P26"] = "Synthèse"
        _reporter_quantites_cloture(ws)
        self.assertEqual(ws["P24"].value, 12)
        self.assertEqual(ws["Q24"].value, 0)
        self.assertEqual(ws["Q25"].value, "Total soumission hors TVA")
        self.assertEqual(ws["R25"].value, 171360.60)
        self.assertEqual(ws["P26"].value, "Synthèse")

    def test_soumission_recuperee_avant_cloture(self):
        with TemporaryDirectory() as tmp:
            chemin = Path(tmp) / "Etat_avancement_02.xlsx"
            dossier = chemin.parent / "Sauvegarde"
            dossier.mkdir()
            wb = Workbook()
            wb.active.title = "Bordereau"
            wb.active["Q25"] = "Total soumission hors TVA"
            wb.active["R25"] = 171360.60
            wb.save(dossier / "Etat_avancement_02_2026-09-10.xlsx")
            self.assertEqual(_soumission_avant_cloture(chemin, nombre_pilotage), 171360.60)

    def test_avenants_moins_negatif_ou_positif_reste_une_diminution(self):
        attendu = 20439.02

        self.assertAlmostEqual(
            _montant_diminution_avenant(nombre_pilotage("-20 439,02 €")),
            attendu,
            places=2,
        )
        self.assertAlmostEqual(
            _montant_diminution_avenant(nombre_pilotage("20 439,02 €")),
            attendu,
            places=2,
        )

    def test_pilotage_lit_la_synthese_du_fichier_avenants(self):
        class Cellule:
            def __init__(self, value):
                self.value = value

        class FeuilleAvenants:
            max_row = 62
            valeurs = {
                "E62": Cellule("12 000,00 €"),
                "G62": Cellule("-20 439,02 €"),
            }

            def __getitem__(self, cellule):
                return self.valeurs.get(cellule, Cellule(None))

        avenants_plus, avenants_moins = _lire_synthese_avenants_pilotage(
            FeuilleAvenants(),
            nombre_pilotage,
        )

        self.assertAlmostEqual(avenants_plus, 12000.00, places=2)
        self.assertAlmostEqual(avenants_moins, 20439.02, places=2)

    def test_pilotage_lit_la_revision_globale_en_colonne_e(self):
        class Cellule:
            def __init__(self, value):
                self.value = value

        class FeuilleRevisionGlobale:
            max_row = 62
            valeurs = {
                "E62": Cellule("3 250,50 €"),
            }

            def __getitem__(self, cellule):
                return self.valeurs.get(cellule, Cellule(None))

        revision_globale = _lire_revision_globale_pilotage(
            FeuilleRevisionGlobale(),
            nombre_pilotage,
        )

        self.assertAlmostEqual(revision_globale, 3250.50, places=2)

    def test_ouverture_revision_retrouve_nom_unicode_existant(self):
        with TemporaryDirectory() as tmp:
            dossier = Path(tmp)
            fichier_existant = dossier / "Formule_révision.xlsm"
            fichier_existant.touch()

            chemin = _chemin_fichier_chantier(dossier, "Formule_révision.xlsm")

            self.assertEqual(chemin, fichier_existant)


if __name__ == "__main__":
    unittest.main()

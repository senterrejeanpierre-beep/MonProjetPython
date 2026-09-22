import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from main import _creer_documents_chantier, _modeles_creation_chantier


class TestCreationChantier(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.base = Path(self.tmp.name)
        self.modeles = self.base / 'Modele'
        self.modeles.mkdir()
        (self.modeles / 'Etat_avancement_02.xlsm').write_bytes(b'ancien modele')
        (self.modeles / 'prix_de_revient model.xlsx').write_bytes(b'prix')
        (self.modeles / 'Planning_soumission.xlsx').write_bytes(b'planning soumission')

    def test_reprise_dossier_vide_et_documents(self):
        (self.base / 'Chantier').mkdir()
        (self.modeles / 'Documents_installation/098_Montage_echafaudage/').mkdir(parents=True)
        (self.modeles / 'Fiches_techniques/014_Chaux').mkdir(parents=True)
        (self.modeles / 'Administratif/Personnel').mkdir(parents=True)
        (self.modeles / 'Administratif/Personnel/diplome.pdf').write_bytes(b'commun')
        (self.modeles / 'Factures').mkdir()
        (self.modeles / 'Factures/facture_autre_chantier.pdf').write_bytes(b'prive')
        resultat = _creer_documents_chantier(self.base, 'Chantier', {'type_etat': 'Modèle actuel'})
        self.assertTrue(resultat.is_file())
        self.assertEqual((self.base / 'Chantier/data/prix_de_revient.xlsx').read_bytes(), b'prix')
        self.assertEqual((self.base / 'Chantier/Planning_soumission.xlsx').read_bytes(), b'planning soumission')
        self.assertEqual((self.base / 'Chantier/Etat_avancement_02.xlsm').read_bytes(), b'ancien modele')
        self.assertTrue((self.base / 'Chantier/Documents_installation').is_dir())
        self.assertTrue((self.base / 'Chantier/Documents_installation/098_Montage_echafaudage').is_dir())
        self.assertTrue((self.base / 'Chantier/Technique').is_dir())
        self.assertFalse((self.base / 'Chantier/Fiches_techniques').exists())
        self.assertTrue((self.base / 'Chantier/Administratif').is_dir())
        self.assertFalse(list((self.base / 'Chantier/Administratif').iterdir()))
        self.assertTrue((self.base / 'Chantier/Factures').is_dir())
        self.assertEqual({p.name for p in (self.base / 'Chantier/Factures').iterdir()},
                         {'Entreprise', 'Achats_et_sous_traitants', 'Achats_materiaux'})
        self.assertEqual((self.modeles / 'Factures/facture_autre_chantier.pdf').read_bytes(), b'prive')

    def test_nouveau_chantier_liens_dossiers_relatifs(self):
        from openpyxl import Workbook, load_workbook
        from module_engins_cout import TOTAL_INITIAL
        w = Workbook()
        w.active['A98'] = '098 · Montage échafaudage'
        w.active['K17'] = '=' + TOTAL_INITIAL
        for col, texte in zip('ABCDEFGHIJK', ('Définition', 'Entreprise', 'Unité', 'Quantité',
                                              'Pièce', 'M1', 'M2', 'M3', 'Heures', 'Prix u', 'Prix total')):
            w.active[f'{col}162'] = texte
        w.active['A163'] = 'FIN DE CHANTIER'
        w.active['K163'] = '=SUM($K164:$K524)'
        for col in 'ABCDEFGHIJK':
            w.active[f'{col}164'] = 0
        w.active['A164'] = '164 · Démontage échafaudages'
        w.active['K164'] = '=IF($J164="","",$J164*IF(ISNUMBER($E164),$E164,1))'
        w.save(self.modeles / 'Cout_securite.xlsx')
        _creer_documents_chantier(self.base, 'Nouveau', {})
        dossier = self.base / 'Nouveau'
        w = load_workbook(dossier / 'Cout_securite.xlsx')
        self.assertEqual(w.active['B98'].value, 'Docs')
        from liens_horizon import analyser
        self.assertEqual(analyser(w.active['B98'].hyperlink.target)[1], '098')
        self.assertTrue((dossier / 'Documents_installation/098_Montage_echafaudage').is_dir())
        self.assertEqual(w.active['A525'].value, '193 · Protection des accès pour engin')

    def test_pas_de_substitution_public_prive(self):
        import json
        resultat = _creer_documents_chantier(self.base, 'Public', {'type_etat': 'Public'})
        fiche = json.loads(resultat.read_text())
        self.assertEqual(fiche['type_etat'], 'Public')
        self.assertTrue(fiche['modele_etat_manquant'])
        self.assertFalse(list((self.base / 'Public').glob('Etat_avancement*')))
        self.assertTrue((self.base / 'Public/data/prix_de_revient.xlsx').exists())

    def test_selection_accents_et_casse(self):
        (self.modeles / 'Etat_avancement_prive\u0301.xlsm').write_bytes(b'prive')
        (self.modeles / 'Etat_avancement_public.xlsm').write_bytes(b'public')
        self.assertIn('Privé', _modeles_creation_chantier(self.base)[2])
        _creer_documents_chantier(self.base, 'Prive', {'type_etat': 'Privé'})
        self.assertEqual(len(list((self.base / 'Prive').glob('Etat_avancement*'))), 2)
        self.assertEqual((self.base / 'Prive/Etat_avancement_public.xlsm').read_bytes(), b'public')
        self.assertEqual((self.base / 'Prive/Etat_avancement_prive\u0301.xlsm').read_bytes(), b'prive')

    def test_echec_copie_ne_laisse_aucun_chantier(self):
        with patch('main.shutil.copy2', side_effect=OSError('copie interrompue')):
            with self.assertRaises(OSError):
                _creer_documents_chantier(self.base, 'Erreur', {})
        self.assertEqual(list(self.base.iterdir()), [self.modeles])

    def test_dossier_existant_preserve(self):
        dossier = self.base / 'Existant'
        dossier.mkdir()
        document = dossier / 'important.txt'
        document.write_text('conserver')
        with self.assertRaises(FileExistsError):
            _creer_documents_chantier(self.base, 'Existant', {})
        self.assertEqual(document.read_text(), 'conserver')

    def test_nom_invalide(self):
        for nom in ('.', '..', '../ailleurs'):
            with self.assertRaises(ValueError):
                _creer_documents_chantier(self.base, nom, {})


if __name__ == '__main__':
    unittest.main()

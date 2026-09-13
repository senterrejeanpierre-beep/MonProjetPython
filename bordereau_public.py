"""Copie de travail du métré Mariemont, sans modification du document source."""
from pathlib import Path
import hashlib
import json
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

CONFIG = 'bordereau_public.json'

def est_metre(ws):
    return (ws['B1'].value == 'N°' and ws['K1'].value == 'QUANTITÉ'
            and ws['L1'].value == 'PRIX UNITAIRES')

def feuille(wb):
    if 'Bordereau' in wb.sheetnames:
        return wb['Bordereau']
    candidates = [s for s in wb if est_metre(s)]
    if len(candidates) != 1:
        raise ValueError('Feuille du bordereau introuvable ou ambiguë.')
    return candidates[0]

def lignes_postes(ws):
    return [r for r in range(2, ws.max_row + 1)
            if ws[f'B{r}'].value and ws[f'D{r}'].value in {'QP', 'QF', 'PG'}
            and ws[f'L{r}'].value is not None]

def lire_config(dossier):
    p = Path(dossier) / CONFIG
    if not p.exists():
        return None
    config = json.loads(p.read_text(encoding='utf-8'))
    for key in ('original', 'copie'):
        name = config[key]
        if not isinstance(name, str) or any(c in name for c in ('/', '\\', ':')) or name in ('', '.', '..'):
            raise ValueError('Chemin de bordereau invalide.')
    return config

def preparer(original, destination, modele=None):
    original, destination = Path(original), Path(destination)
    destination.mkdir(parents=True, exist_ok=True)
    copie = destination / (original.stem + '_avancement.xlsx')
    config_path = destination / CONFIG
    if copie.exists() or config_path.exists():
        raise FileExistsError('Une copie de travail existe déjà. Elle est conservée.')
    wb = load_workbook(original)
    try:
        ws = feuille(wb)
        if not est_metre(ws):
            raise ValueError('Structure du métré non reconnue.')
        if any(c.value is not None for row in ws.iter_rows(min_col=16) for c in row):
            raise ValueError('La zone avancement P:W est déjà occupée.')
        if any(m.max_col >= 16 for m in ws.merged_cells.ranges):
            raise ValueError('La zone avancement contient des cellules fusionnées.')
        postes = lignes_postes(ws)
        if not postes:
            raise ValueError('Aucun poste chiffrable trouvé.')
        habiller_avancement(wb, modele)
        wb.save(copie)
        config = {'original': original.name, 'copie': copie.name,
                  'sha256_original': hashlib.sha256(original.read_bytes()).hexdigest()}
        config_path.write_text(json.dumps(config, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        return copie
    finally:
        wb.close()

def habiller_entete(ws):
    """En-tête du suivi conforme au tableau métier, sans toucher aux saisies."""
    from openpyxl.styles import Border, Side
    for fusion in list(ws.merged_cells.ranges):
        if fusion.min_col >= 16 and fusion.max_col <= 23 and 3 <= fusion.min_row <= fusion.max_row <= 7:
            ws.unmerge_cells(str(fusion))
    thin = Side(style='thin', color='000000')
    for row in ws.iter_rows(min_row=3, max_row=7, min_col=16, max_col=23):
        for cell in row:
            cell.value = None
            cell.font = Font(name='Arial', size=14, bold=True, color='000000')
            cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
            cell.fill = PatternFill('solid', fgColor='F2F2F2' if cell.row <= 5 else 'FFFFFF')
            cell.border = Border()
    for r in range(3, 8):
        ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 15, 20)
    for plage in ('P3:T5', 'U3:W5'):
        ws.merge_cells(plage)
    ws['P3'] = 'Historique des états'
    ws['U3'] = ('="État d’avancement n° "&Q9&CHAR(10)&'
                'IF(OR(Q10="",Q11=""),"Période à renseigner",'
                '"Période du "&TEXT(Q10,"dd/mm/yyyy")&" au "&TEXT(Q11,"dd/mm/yyyy"))')
    titres = ('Quantités\nprécédentes', 'Quantités\nétat', 'Quantités\ncumulées',
              'Global en %', 'Prix total', 'Quantités état', 'Prix unitaire', 'Prix global')
    for col, titre in enumerate(titres, 16):
        ws.merge_cells(start_row=6, start_column=col, end_row=7, end_column=col)
        ws.cell(6, col, titre)
    ws['S6'].fill = PatternFill('solid', fgColor='DDEBF7')
    for fusion in list(ws.merged_cells.ranges):
        if fusion.min_col >= 16 and fusion.max_col <= 23 and 3 <= fusion.min_row <= fusion.max_row <= 7:
            ws.cell(fusion.min_row, fusion.min_col).border = Border(left=thin, right=thin, top=thin, bottom=thin)
            fusion.format()


def habiller_avancement(wb, modele=None):
    """Reprend le tableau complet du modèle, adapté aux lignes du métré public."""
    from copy import copy
    from openpyxl.styles import Border, Side, Protection
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.workbook.external_link.external import ExternalLink, ExternalBook, ExternalSheetNames
    from openpyxl.packaging.relationship import Relationship
    from openpyxl.workbook.properties import CalcProperties
    ws = feuille(wb)
    if not est_metre(ws):
        raise ValueError('Structure du métré non reconnue.')
    postes = lignes_postes(ws)
    if not postes or min(postes) < 12:
        raise ValueError('Le métré ne laisse pas assez de place pour les en-têtes du suivi.')
    # Garder les saisies déjà effectuées dans la copie de travail.
    saisies = {(r, c): ws.cell(r, c).value for r in postes for c in (16, 17)}
    complet = ws['P9'].value == 'Numéro de l’état'
    parametres = {a: ws[a].value for a in ('Q9', 'Q10', 'Q11', 'W9')} if complet else {}
    fin = max(c.row for row in ws.iter_rows(max_col=15) for c in row if c.value is not None)
    debut = fin + 3
    for fusion in list(ws.merged_cells.ranges):
        if fusion.min_col >= 16 and fusion.max_col <= 23:
            ws.unmerge_cells(str(fusion))
    for row in ws.iter_rows(min_col=16, max_col=23):
        for c in row:
            c.value = None
            c._style = None
            c.comment = None
    source = None
    if modele is not None:
        source = load_workbook(modele)
    try:
        modele_ws = source['Bordereau'] if source else None
        thin = Side(style='thin', color='000000')
        euro = '#,##0.00 "€";[Red]-#,##0.00 "€";"-   €"'
        def style(c, r, col):
            if modele_ws:
                sc = modele_ws.cell(r, col)
                c.font, c.fill, c.border = copy(sc.font), copy(sc.fill), copy(sc.border)
                c.alignment, c.protection = copy(sc.alignment), copy(sc.protection)
                c.number_format = sc.number_format
            else:
                c.font = Font(name='Arial', size=12, bold=r < 18 or r >= 676)
                c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            c.alignment = Alignment(horizontal='right' if c.column in (20, 23) else 'center',
                                    vertical='center', wrap_text=True)
        # Les dimensions Excel peuvent être regroupées : les remplacer individuellement.
        from openpyxl.worksheet.dimensions import ColumnDimension
        for col in range(16, 24):
            lettre = ws.cell(1, col).column_letter
            largeur = modele_ws.column_dimensions[lettre].width if modele_ws else 20
            ws.column_dimensions[lettre] = ColumnDimension(ws, index=lettre, width=largeur)
        habiller_entete(ws)
        for label, adresse, texte, valeur in (
            ('P9', 'Q9', 'Numéro de l’état', 1),
            ('P10', 'Q10', 'Début de période', None),
            ('P11', 'Q11', 'Fin de période', None),
            ('V9', 'W9', 'Taux de TVA', 0.06),
        ):
            ws[label] = texte
            ws[adresse] = parametres.get(adresse, valeur)
            ws[adresse].fill = PatternFill('solid', fgColor='FFF2CC')
            ws[adresse].protection = Protection(locked=False)
        ws['W9'].number_format = '0%'
        for a in ('Q10', 'Q11'):
            ws[a].number_format = 'dd-mm-yyyy'
        ws.data_validations.dataValidation = [d for d in ws.data_validations.dataValidation if str(d.sqref) != 'Q9']
        dv = DataValidation(type='whole', operator='between', formula1=1, formula2=20,
                            allow_blank=False, showErrorMessage=True)
        dv.error = 'Indiquez un numéro d’état de 1 à 20.'
        ws.add_data_validation(dv); dv.add(ws['Q9'])
        for r in range(12, fin + 1):
            for col in range(16, 24):
                style(ws.cell(r, col), 24, col)
        for r in postes:
            for col in (16, 17):
                ws.cell(r, col, saisies[r, col] if saisies[r, col] is not None else 0)
                ws.cell(r, col).protection = Protection(locked=False)
            for col, formule in {'R': f'P{r}+Q{r}', 'S': f'IFERROR(R{r}/K{r},0)',
                                 'T': f'R{r}*L{r}', 'U': f'Q{r}', 'V': f'L{r}',
                                 'W': f'Q{r}*L{r}'}.items():
                ws[f'{col}{r}'] = '=' + formule
                ws[f'{col}{r}'].number_format = '0.00%' if col == 'S' else (euro if col in 'TVW' else '0.00')
        # Reprendre les styles de toutes les lignes du bas, y compris les séparateurs.
        for offset in range(9):
            for col in range(16, 24):
                style(ws.cell(debut + offset, col), 676 + offset, col)
            ws.row_dimensions[debut + offset].height = 34
        marche = [r for r in postes if ws[f'M{r}'].value is not None]
        options = [r for r in postes if ws[f'M{r}'].value is None]
        def somme(col, rows):
            return '=SUM(' + ','.join(f'{col}{r}' for r in rows) + ')' if rows else '=0'
        # Liens relatifs aux documents de CE chantier, sans caches d'un ancien chantier.
        if wb._external_links:
            attendus = ['Avenants.xlsx', 'Revision_global.xlsx']
            if not complet or [l.file_link.Target for l in wb._external_links] != attendus:
                raise ValueError('La copie possède des liaisons externes non reconnues.')
            wb._external_links = []
        for nom in ('Avenants.xlsx', 'Revision_global.xlsx'):
            link = ExternalLink(externalBook=ExternalBook(id='rId1', sheetNames=ExternalSheetNames(sheetName=['Feuil1'])))
            link.file_link = Relationship(Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLinkPath',
                                          Target=nom, TargetMode='External', Id='rId1')
            wb._external_links.append(link)
        def ligne_gauche(offset, texte, formule):
            r = debut + offset
            ws.cell(r, 16, texte)
            style(ws.cell(r, 16), 676 + offset if 0 <= offset <= 8 else 676, 19)
            ws.merge_cells(start_row=r, start_column=16, end_row=r, end_column=19)
            ws.cell(r, 20, formule).number_format = euro
        # Soumission au-dessus du bloc de synthèse, pour conserver les deux tableaux côte à côte.
        ligne_gauche(-2, 'Total soumission hors TVA', somme('M', marche))
        for c in range(16, 24):
            if ws.cell(debut - 2, c).__class__.__name__ != 'MergedCell':
                style(ws.cell(debut - 2, c), 676, c)
        ligne_gauche(0, 'Total état cumulé hors TVA', somme('T', marche))
        ligne_gauche(2, 'Total du mois hors tva', f'=W{debut}')
        ligne_gauche(4, 'Total des avenants cumulé', "=SUM('[1]Feuil1'!$D$3:$D$60)-SUM('[1]Feuil1'!$E$3:$E$60)")
        ligne_gauche(6, '=IF(Q11="","Total exécuté au","Total exécuté au "&TEXT(Q11,"dd-mm-yyyy"))', f'=T{debut}')
        ligne_gauche(8, 'Total des révision cumulé', "=SUM('[2]Feuil1'!$D$3:$D$60)")
        # N() convertit les tirets décoratifs des fichiers métier en zéro, sans masquer les erreurs de lien.
        avenant = "=N(INDEX('[1]Feuil1'!$D$3:$D$60,1+3*($Q$9-1)))-N(INDEX('[1]Feuil1'!$E$3:$E$60,1+3*($Q$9-1)))"
        revision = "=N(INDEX('[2]Feuil1'!$D$3:$D$60,1+3*($Q$9-1)))"
        droite = [
            ('="Montant de l’état "&Q9', somme('W', marche)),
            ('="Montant de l’avenant "&Q9', avenant),
            ('Montant global état hors TVA', f'=SUM(W{debut}:W{debut+1})'),
            ('="Montant de la TVA de "&TEXT(W9,"0%")', f'=W{debut+2}*$W$9'),
            ('="Montant de la révision sur l’état "&Q9&" hors avenant"', revision),
            ('="Montant Global à facturer état "&Q9', f'=SUM(W{debut+2}:W{debut+4})'),
        ]
        for offset, (texte, formule) in enumerate(droite):
            r = debut + offset
            ws.cell(r, 21, texte)
            style(ws.cell(r, 21), 676 + offset, 22)
            ws.merge_cells(start_row=r, start_column=21, end_row=r, end_column=22)
            ws.cell(r, 23, formule).number_format = euro
        # Même repère visuel que la capture fournie.
        for col in range(21, 24):
            ws.cell(debut + 5, col).fill = PatternFill('solid', fgColor='FFFF00')
        for r, col in ((debut + 5, 23), (debut + 3, 23), (debut + 2, 16)):
            font = copy(ws.cell(r, col).font); font.color = '000080' if col == 23 else '0070C0'
            ws.cell(r, col).font = font
        ligne_gauche(10, 'Avenants cumulés en plus', "=SUM('[1]Feuil1'!$D$3:$D$60)")
        ligne_gauche(11, 'Avenants cumulés en moins', "=SUM('[1]Feuil1'!$E$3:$E$60)")
        if options:
            ligne_gauche(13, 'Options hors TVA (hors marché)', somme('N', options))
            ws.row_dimensions[debut + 13].height = 30
        ws.print_area = f'A1:W{debut+13}'
        ws.sheet_view.zoomScale = 60
        wb.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True, calcMode='auto')
        return debut
    finally:
        if source:
            source.close()


def completer_copie(copie, modele, sortie):
    wb = load_workbook(copie)
    try:
        debut = habiller_avancement(wb, modele)
        wb.save(sortie)
        return debut
    finally:
        wb.close()

def exporter_prix(original, copie, sortie, empreinte):
    """Ne change que les cellules L des postes dans le XML du fichier officiel."""
    import math
    import re
    import zipfile
    from xml.etree import ElementTree as ET
    original, copie, sortie = map(Path, (original, copie, sortie))
    if hashlib.sha256(original.read_bytes()).hexdigest() != empreinte:
        raise ValueError("L'original a changé depuis la création de la copie. Vérifiez sa version.")
    a = load_workbook(original)
    b = load_workbook(copie)
    try:
        sa, sb = feuille(a), feuille(b)
        ra, rb = lignes_postes(sa), lignes_postes(sb)
        def index(ws, rows):
            result = {}
            for r in rows:
                code = str(ws[f'B{r}'].value).strip()
                if code in result:
                    raise ValueError(f'Article en double : {code}')
                result[code] = r
            return result
        ia, ib = index(sa, ra), index(sb, rb)
        if ia.keys() != ib.keys():
            raise ValueError('Les postes de la copie et de l’original diffèrent.')
        prix = {}
        for code, r in ia.items():
            t = ib[code]
            if any(sa[f'{col}{r}'].value != sb[f'{col}{t}'].value for col in ('C', 'D', 'E')):
                raise ValueError(f'Désignation, mesurage ou unité modifiés : {code}')
            valeur = sb[f'L{t}'].value
            if isinstance(valeur, bool) or not isinstance(valeur, (int, float)) or not math.isfinite(valeur):
                raise ValueError(f'Prix non numérique pour {code}.')
            if sa[f'L{r}'].data_type == 'f':
                raise ValueError(f'Formule protégée en L{r}.')
            prix[f'L{r}'] = valeur
        with zipfile.ZipFile(original) as zin:
            ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
            book = ET.fromstring(zin.read('xl/workbook.xml'))
            rid = next(s for s in book.find('s:sheets', ns) if s.attrib['name'] == sa.title).attrib[
                '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']
            rels = ET.fromstring(zin.read('xl/_rels/workbook.xml.rels'))
            target = next(r.attrib['Target'] for r in rels if r.attrib['Id'] == rid)
            sheet_path = target.lstrip('/') if target.startswith('/') else 'xl/' + target
            xml = zin.read(sheet_path).decode('utf-8')
            for adresse, valeur in prix.items():
                pattern = rf'<c\b[^>]*\br="{adresse}"[^>]*>.*?</c>'
                match = re.search(pattern, xml, re.S)
                if not match:
                    raise ValueError(f'Cellule officielle introuvable : {adresse}')
                cell = match.group()
                if '<f' in cell:
                    raise ValueError(f'Formule officielle protégée : {adresse}')
                opening = cell[:cell.index('>') + 1]
                opening = re.sub(r'\s+t="[^"]*"', '', opening)
                replacement = opening + '<v>' + repr(valeur) + '</v></c>'
                xml = xml[:match.start()] + replacement + xml[match.end():]
            with zipfile.ZipFile(sortie, 'w') as zout:
                for info in zin.infolist():
                    zout.writestr(info, xml.encode('utf-8') if info.filename == sheet_path else zin.read(info.filename))
        return len(prix)
    finally:
        a.close()
        b.close()

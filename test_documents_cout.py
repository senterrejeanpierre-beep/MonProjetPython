import shutil
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from documents_cout import (lire_postes, actualiser_liens, chemin_document, dossier_poste,
                            renommer_dossier_poste, renommer_libelle_poste)
from liens_horizon import analyser, resoudre


class DocumentsCoutTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.racine = Path(self.temp.name)
        self.chantier = self.racine / 'Chantier'
        self.chantier.mkdir()
        self.classeur = self.chantier / 'Cout_securite.xlsx'
        w = Workbook()
        s = w.active
        s.title = 'Coût'
        s['A21'] = '021 · Échafaudage'
        s['B21'] = 'Docs'
        s['B21'].hyperlink = 'Documents_installation/021/'
        s['A22'] = '022 · Manutention'
        s['B22'] = 'Docs'
        s['B22'].hyperlink = 'Documents_installation/022/'
        s['A23'] = '023 · Définition'
        s['B23'] = 'Entreprise'
        s['B23'].hyperlink = 'Documents_installation/023/'
        s['J21'] = 52
        s['K21'] = '=J21*2'
        w.save(self.classeur)
        with ZipFile(self.classeur, 'a') as z:
            z.writestr('xl/media/temoin.bin', b'graphique intact')

    def test_dossiers_vides_et_conservation(self):
        from urllib.parse import unquote
        with ZipFile(self.classeur) as z:
            avant = {n: z.read(n) for n in z.namelist()}
        self.assertTrue(actualiser_liens(self.classeur))
        w = load_workbook(self.classeur)
        self.assertEqual(w.active['K21'].value, '=J21*2')
        self.assertEqual(w.active['J21'].value, 52)
        self.assertEqual(w.active['B23'].value, 'Entreprise')
        self.assertIsNone(w.active['B23'].hyperlink)
        for ref in ('B21', 'B22'):
            self.assertEqual(w.active[ref].value, 'Docs')
            cible = w.active[ref].hyperlink.target
            self.assertTrue(cible.startswith('horizonchantier://documents/'))
            self.assertEqual(resoudre(self.racine, cible)[0], self.chantier)
        with ZipFile(self.classeur) as z:
            for nom, data in avant.items():
                if nom not in ('xl/worksheets/sheet1.xml', 'xl/worksheets/_rels/sheet1.xml.rels'):
                    self.assertEqual(z.read(nom), data, nom)
        self.assertFalse(actualiser_liens(self.classeur))

    def test_renommage_du_poste_garde_ses_documents(self):
        ancien = self.chantier / 'Documents_installation/193_Ancien_nom'
        ancien.mkdir(parents=True)
        (ancien / 'offre.pdf').write_bytes('offre conservée'.encode())
        nouveau = renommer_dossier_poste(
            self.classeur, {'libelle': '193 · Protection des accès pour engin'})
        self.assertEqual(nouveau.name, '193_Protection_des_accès_pour_engin')
        self.assertEqual((nouveau / 'offre.pdf').read_bytes(), 'offre conservée'.encode())
        self.assertFalse(ancien.exists())

    def test_renommage_excel_preserve_valeurs_et_autres_membres(self):
        self.assertTrue(renommer_libelle_poste(self.classeur, '021', 'Nouvel échafaudage'))
        w = load_workbook(self.classeur)
        self.assertEqual(w.active['A21'].value, '021 · Nouvel échafaudage')
        self.assertEqual(w.active['J21'].value, 52)
        self.assertEqual(w.active['K21'].value, '=J21*2')
        with ZipFile(self.classeur) as z:
            self.assertEqual(z.read('xl/media/temoin.bin'), b'graphique intact')
        self.assertFalse(renommer_libelle_poste(self.classeur, '021', 'Nouvel échafaudage'))

    def test_cellule_vide_autofermante_preserve_cellules_suivantes(self):
        from xml.etree import ElementTree as ET
        w = load_workbook(self.classeur)
        w.active['C21'] = 12
        w.active['D21'] = '=C21*5'
        w.save(self.classeur)
        actualiser_liens(self.classeur)
        with ZipFile(self.classeur) as z:
            contenus = {n: z.read(n) for n in z.namelist()}
        nom = 'xl/worksheets/sheet1.xml'
        import re
        xml = contenus[nom].decode()
        xml = re.sub(r'<c r="B21".*?</c>', '<c r="B21" s="0"/>', xml)
        # Simule l'état sans liens laissé par la première correction.
        xml = re.sub(r'<hyperlinks.*?</hyperlinks>', '', xml)
        contenus[nom] = xml.encode()
        with ZipFile(self.classeur, 'w') as z:
            for n, data in contenus.items():
                z.writestr(n, data)
        actualiser_liens(self.classeur)
        w = load_workbook(self.classeur)
        self.assertEqual(w.active['C21'].value, 12)
        self.assertEqual(w.active['D21'].value, '=C21*5')
        self.assertEqual(w.active['B21'].value, 'Docs')

    def test_repare_relations_prefixees_des_classeurs_deja_modifies(self):
        from xml.etree import ElementTree as ET
        actualiser_liens(self.classeur)
        with ZipFile(self.classeur) as z:
            contenus = {n: z.read(n) for n in z.namelist()}
        rel = 'xl/worksheets/_rels/sheet1.xml.rels'
        contenus[rel] = ET.tostring(ET.fromstring(contenus[rel]), encoding='utf-8')
        self.assertIn(b'<ns0:Relationships', contenus[rel])
        with ZipFile(self.classeur, 'w') as z:
            for n, data in contenus.items():
                z.writestr(n, data)
        self.assertTrue(actualiser_liens(self.classeur))
        with ZipFile(self.classeur) as z:
            self.assertIn(b'<Relationships xmlns=', z.read(rel))
            self.assertNotIn(b'ns0:', z.read(rel))
            self.assertNotIn(b'ns0:', z.read('xl/worksheets/sheet1.xml'))
        self.assertFalse(actualiser_liens(self.classeur))

    def test_relation_avec_attribut_xml_qualifie_reste_valide(self):
        from xml.etree import ElementTree as ET
        with ZipFile(self.classeur) as z:
            contenus = {n: z.read(n) for n in z.namelist()}
        rel = 'xl/worksheets/_rels/sheet1.xml.rels'
        arbre = ET.fromstring(contenus[rel])
        arbre[0].set('{urn:test}indicateur', 'présent')
        contenus[rel] = ET.tostring(arbre, encoding='utf-8')
        with ZipFile(self.classeur, 'w') as z:
            for nom, data in contenus.items():
                z.writestr(nom, data)
        self.assertTrue(actualiser_liens(self.classeur))
        with ZipFile(self.classeur) as z:
            verifie = ET.fromstring(z.read(rel))
        self.assertEqual(verifie[0].get('{urn:test}indicateur'), 'présent')

    def test_reutilise_dossier_et_plusieurs_documents_apres_deplacement(self):
        dossier = self.chantier / 'Documents_installation/021_Echafaudage'
        dossier.mkdir(parents=True)
        (dossier / 'offre.pdf').write_bytes(b'offre')
        (dossier / 'photo.jpg').write_bytes(b'photo')
        actualiser_liens(self.classeur)
        cible = lire_postes(self.classeur)[0]['cible']
        self.assertEqual(analyser(cible)[1], '021')
        nouveau = self.racine / 'Partage déplacé'
        shutil.move(self.chantier, nouveau)
        self.assertEqual((nouveau / 'Documents_installation/021_Echafaudage' / 'offre.pdf').read_bytes(), b'offre')
        self.assertEqual((nouveau / 'Documents_installation/021_Echafaudage' / 'photo.jpg').read_bytes(), b'photo')
        self.assertFalse(actualiser_liens(nouveau / self.classeur.name))

    def test_classeur_ouvert_aucune_modification_excel(self):
        avant = self.classeur.read_bytes()
        self.classeur.with_name('~$' + self.classeur.name).touch()
        with self.assertRaises(PermissionError):
            actualiser_liens(self.classeur)
        self.assertEqual(self.classeur.read_bytes(), avant)

    def test_cellule_absente_et_ancien_lien_direct(self):
        dossier = self.chantier / 'Documents_installation/021'
        dossier.mkdir(parents=True)
        (dossier / 'offre.pdf').write_bytes(b'offre')
        w = load_workbook(self.classeur)
        w.active['B21'].hyperlink = 'Documents_installation/021/offre.pdf'
        w.active['A24'] = '024 · Transport'
        w.save(self.classeur)
        actualiser_liens(self.classeur)
        w = load_workbook(self.classeur)
        self.assertEqual(w.active['B24'].value, 'Docs')
        self.assertEqual(analyser(w.active['B21'].hyperlink.target)[1], '021')
        self.assertEqual((dossier / 'offre.pdf').read_bytes(), b'offre')


if __name__ == '__main__':
    unittest.main()

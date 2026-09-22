import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from listes_pr import reparer_listes_pr, NS
from main import open_pr
from listes_pr import RUBRIQUES_SECURITE


class ListesPR(unittest.TestCase):
    def test_ancien_ajout_en_fin_remonte_sans_doublon(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'prix_de_revient.xlsx'
            wb = Workbook()
            ws = wb.active
            ws.title = 'Chiffrage'
            ws['B1'], ws['B14'] = 'MATIÈRE', 'MAIN-D’ŒUVRE'
            ws['D2'] = ws['D15'] = 'Désignation matière'
            for nom in ('Bibliothèque_Matière', 'Bibliothèque_MainOeuvre'):
                lib = wb.create_sheet(nom)
                lib.append(['Choix manuel'])
                for titre in RUBRIQUES_SECURITE[:-1]:
                    lib.append([titre])
                lib.append(['11 — CONTAINERS, ENGINS DE LEVAGE ET VÉGÉTATION'])
            wb.save(path)
            wb.close()
            self.assertTrue(reparer_listes_pr(path))
            self.assertFalse(reparer_listes_pr(path))
            wb = load_workbook(path, read_only=True)
            for nom in ('Bibliothèque_Matière', 'Bibliothèque_MainOeuvre'):
                valeurs = [r[0] for r in wb[nom].iter_rows(max_col=1, values_only=True)]
                self.assertEqual(valeurs, list(RUBRIQUES_SECURITE) + ['Choix manuel'])
            wb.close()

    def test_ouverture_pr_ne_modifie_pas_le_classeur(self):
        with tempfile.TemporaryDirectory() as tmp:
            chantier = Path(tmp)
            data = chantier / 'data'
            data.mkdir()
            path = data / 'prix_de_revient.xlsx'
            wb = Workbook()
            ws = wb.active
            ws.title = 'Chiffrage'
            ws['B1'], ws['B14'] = 'MATIÈRE', 'MAIN-D’ŒUVRE'
            ws['D2'] = ws['D15'] = 'Désignation matière'
            ws['B27'] = 'RÉCAPITULATIF POSTE'
            ws['P13'] = '=SUM(P3:P12)'
            for name in ('Bibliothèque_Matière', 'Bibliothèque_MainOeuvre'):
                wb.create_sheet(name).append(['Choix'])
            wb.save(path)
            wb.close()
            avant = path.read_bytes()
            with patch('main.ouvrir_chemin') as ouvrir, patch('main.reparer_listes_pr') as reparer:
                self.assertFalse(open_pr(chantier))
            ouvrir.assert_called_once_with(path)
            reparer.assert_not_called()
            self.assertEqual(path.read_bytes(), avant)

    def test_pr_verrouille_s_ouvre_sans_ecriture(self):
        with tempfile.TemporaryDirectory() as tmp:
            chantier = Path(tmp)
            data = chantier / 'data'
            data.mkdir()
            path = data / 'prix_de_revient.xlsx'
            wb = Workbook()
            wb.save(path)
            wb.close()
            avant = path.read_bytes()
            (data / '~$prix_de_revient.xlsx').write_bytes(b'verrou Excel')
            with patch('main.ouvrir_chemin') as ouvrir, patch('main.reparer_listes_pr') as reparer:
                self.assertFalse(open_pr(chantier))
            ouvrir.assert_called_once_with(path)
            reparer.assert_not_called()
            self.assertEqual(path.read_bytes(), avant)

    def test_echec_reparation_listes_ne_bloque_pas_ouverture(self):
        with tempfile.TemporaryDirectory() as tmp:
            chantier = Path(tmp)
            data = chantier / 'data'
            data.mkdir()
            path = data / 'prix_de_revient.xlsx'
            path.write_bytes(b'contenu de test')
            with patch('main.ouvrir_chemin') as ouvrir, patch(
                'main.reparer_listes_pr', side_effect=PermissionError('fichier inaccessible')
            ):
                self.assertFalse(open_pr(chantier))
            ouvrir.assert_called_once_with(path)

    def test_colonnes_et_sauvegardes(self):
        for col in ('C', 'D'):
            with self.subTest(col=col), tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / 'prix_de_revient.xlsx'
                wb = Workbook()
                ws = wb.active
                ws.title = 'Chiffrage'
                ws['B1'], ws['B14'] = 'MATIÈRE', 'MAIN-D’ŒUVRE'
                ws[f'{col}2'] = ws[f'{col}15'] = 'Désignation matière'
                ws['B27'] = 'RÉCAPITULATIF POSTE'
                ws['P13'] = '=SUM(P3:P12)'
                ws[f'{col}3'] = 'Choix existant'
                for name in ('Bibliothèque_Matière', 'Bibliothèque_MainOeuvre'):
                    lib = wb.create_sheet(name)
                    lib.append(['Premier choix'])
                    lib.append(['Dernier choix'])
                dv = DataValidation(type='list', formula1='"h/u,u/h"')
                ws.add_data_validation(dv)
                dv.add('E16:E25')
                wb.save(path)
                with ZipFile(path) as z:
                    before = {n: z.read(n) for n in z.namelist()}
                self.assertTrue(reparer_listes_pr(path))
                self.assertFalse(reparer_listes_pr(path))
                with ZipFile(path) as z:
                    for name, data in before.items():
                        if name not in ('xl/worksheets/sheet1.xml', 'xl/worksheets/sheet2.xml', 'xl/worksheets/sheet3.xml'):
                            self.assertEqual(z.read(name), data)
                    after = ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
                    old = ET.fromstring(before['xl/worksheets/sheet1.xml'])
                    self.assertEqual(ET.tostring(after.find(f'{{{NS}}}sheetData')),
                                     ET.tostring(old.find(f'{{{NS}}}sheetData')))
                for _ in range(2):
                    wb = load_workbook(path)
                    dvs = wb['Chiffrage'].data_validations.dataValidation
                    self.assertEqual(len(dvs), 3)
                    self.assertEqual(str(dvs[1].sqref), f'{col}3:{col}12')
                    self.assertEqual(str(dvs[2].sqref), f'{col}16:{col}25')
                    self.assertIn('Bibliothèque_MainOeuvre', dvs[2].formula1)
                    self.assertFalse(dvs[1].showDropDown)
                    for nom in ('Bibliothèque_Matière', 'Bibliothèque_MainOeuvre'):
                        valeurs = [cell[0] for cell in wb[nom].iter_rows(max_col=1, values_only=True)]
                        self.assertEqual(valeurs[:11], list(RUBRIQUES_SECURITE))
                        self.assertEqual(valeurs[11:], ['Premier choix', 'Dernier choix'])
                    wb.save(path)
                    wb.close()


if __name__ == '__main__':
    unittest.main()

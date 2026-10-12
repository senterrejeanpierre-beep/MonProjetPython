"""Zéros des cases numériques vides du PR, sans réenregistrer le classeur."""
from copy import deepcopy
import html
import os
from pathlib import Path
import re
import shutil
import tempfile
import unicodedata
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
from pourcentages_pr import _feuille_chiffrage, _signature, _recalcul_excel

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'


def cle(value):
    texte = str(value or '').casefold().replace('œ', 'oe')
    return ''.join(c for c in unicodedata.normalize('NFKD', texte) if c.isalnum())


ENTETES_NUMERIQUES = {
    'qtepcsbordereau', 'quantite', 'qtetotale', 'puachat', 'pertes',
    'pucorrige', 'totalmatiere', 'transporttotal', 'putransport',
    'totaltransport', 'totalrevient', 'nbpers', 'qteposte', 'jours',
    'heurescalc', 'heuresprevues', 'couthoraire', 'coutmo',
    'qterealisee', 'avancement', 'heuresreelles', 'prixumo', 'coutreelmo',
    'matierereelle', 'coutreelmatiere', 'coutreeltotal', 'prtotalref',
    'ecartcout', 'pvhtref', 'ecartmarge',
}


def cases_numeriques(lignes):
    """Repère les zones métier par leurs en-têtes, sans remplir les textes."""
    cases = set()
    for i, row in enumerate(lignes):
        titre = cle(row[1])
        if titre in ('matiere', 'maindoeuvre') and i + 12 < len(lignes):
            header = lignes[i+1]
            attendue = 'nbpers' if titre == 'maindoeuvre' else 'designationmatiere'
            if attendue not in {cle(v) for v in header}:
                continue
            colonnes = [j+1 for j,v in enumerate(header) if cle(v) in ENTETES_NUMERIQUES]
            for n in range(i+2, i+12):
                detail = lignes[n]
                # Les totaux placés dans la grille ne sont pas des lignes de saisie.
                if any(cle(v).startswith(('totaljours', 'totalheures')) for v in detail):
                    cases.add((n+1, 7))
                    if not detail[3]:
                        continue
                if all(v in (None, '') for v in detail[:15]) and detail[15] not in (None, ''):
                    cases.add((n+1, 16))
                    continue
                cases.update((n+1, col) for col in colonnes)
            # Sous-total de la rubrique, juste après les lignes de détail.
            if cle(lignes[i+12][14]).startswith('soustotal'):
                cases.add((i+13, 16))
        if titre.startswith(('prmatiere', 'prmo', 'prtotal')):
            cases.add((i+1, 3))
        if cle(row[4]) in ('ic', 'margevente', 'pvht', 'tva', 'pvttc'):
            cases.add((i+1, 6))
        if cle(row[8]) in ('ic', 'marge'):
            cases.add((i+1, 10))
        if titre == 'date' and cle(row[2]) == 'commentairesituation' and i+1 < len(lignes):
            cases.update((i+2,j+1) for j,v in enumerate(row) if cle(v) in ENTETES_NUMERIQUES)
        for j,v in enumerate(row[:-1]):
            if cle(v) in ('gainpertemo', 'gainpertematiere', 'gainpertetotal'):
                cases.add((i+1,j+2))
    return cases


def proteger_division(formule):
    """Garde le calcul d'origine ; seule une division par une case vide/nulle vaut zéro."""
    if not isinstance(formule, str) or not formule.startswith('='):
        return formule, ()
    expression = formule[1:]
    simple = re.sub(r'\bSUM\b', '', expression, flags=re.I)
    if not re.fullmatch(r'[A-Z0-9$.:()+*/\s-]+', simple, re.I):
        return formule, ()
    denominateurs = tuple(dict.fromkeys(re.findall(r'/\s*(\$?[A-Z]{1,3}\$?\d+)\b', expression, re.I)))
    if not denominateurs:
        return formule, ()
    conditions = ','.join(f'{ref}=0,{ref}=""' for ref in denominateurs)
    return f'=IF(OR({conditions}),0,{expression})', denominateurs


def _cellule(xml, ref, contenu):
    """Remplace le contenu seulement ; garde les styles et les attributs Excel."""
    pattern = rf'<(?:\w+:)?c\b[^>]*\br="{re.escape(ref)}"[^>]*?(?:/>|>.*?</(?:\w+:)?c>)'
    match = re.search(pattern, xml, re.S)
    if match:
        nom, attrs = re.match(r'<((?:\w+:)?c)\b([^>]*?)/?>', match[0]).groups()
        attrs = re.sub(r'\s+t="[^"]*"', '', attrs)
        nouveau = f'<{nom}{attrs}>{contenu}</{nom}>'
        return xml[:match.start()] + nouveau + xml[match.end():]
    numero = int(re.search(r'\d+', ref)[0])
    row = re.search(rf'<(?:\w+:)?row\b[^>]*\br="{numero}"[^>]*?(?:/>|>.*?</(?:\w+:)?row>)', xml, re.S)
    if row is None:
        raise ValueError(f'Ligne de saisie absente : {ref}')
    ajout = f'<c xmlns="{NS}" r="{ref}">{contenu}</c>'
    # Ordre des cellules dans la ligne : A, B, ...
    from openpyxl.utils import column_index_from_string
    colonne = column_index_from_string(re.match(r'[A-Z]+', ref)[0])
    suivant = next((m for m in re.finditer(r'<(?:\w+:)?c\b[^>]*\br="([A-Z]+)\d+"', row[0])
                    if column_index_from_string(m[1]) > colonne), None)
    if suivant:
        contenu_row = row[0][:suivant.start()] + ajout + row[0][suivant.start():]
    elif row[0].endswith('/>'):
        nom = re.match(r'<([^\s>]+)', row[0])[1]
        contenu_row = row[0][:-2] + '>' + ajout + f'</{nom}>'
    else:
        pos = row[0].rfind('</')
        contenu_row = row[0][:pos] + ajout + row[0][pos:]
    return xml[:row.start()] + contenu_row + xml[row.end():]


def completer_zeros_pr(path, sauvegarder=None):
    path = Path(path)
    signature = _signature(path)
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        if 'Chiffrage' not in wb:
            return 0
        lignes = list(wb['Chiffrage'].iter_rows(max_col=16, values_only=True))
    finally:
        wb.close()
    cases = cases_numeriques(lignes)
    if not cases:
        return 0
    with ZipFile(path) as src:
        membre = _feuille_chiffrage(src)
        original = src.read(membre).decode('utf-8')
        root = ET.fromstring(original)
        cellules = {n.get('r'): n for n in root.iter('{'+NS+'}c')}
        fusions = root.find('{'+NS+'}mergeCells')
        ignorees = set()
        for fusion in fusions if fusions is not None else ():
            c1,r1,c2,r2 = range_boundaries(fusion.get('ref'))
            ignorees.update((r,col) for r in range(r1,r2+1) for col in range(c1,min(c2,16)+1)
                            if (r,col)!=(r1,c1))
        changes = {}
        for row,col in sorted(cases-ignorees):
            ref = f'{get_column_letter(col)}{row}'
            valeur = lignes[row-1][col-1]
            if valeur in (None, ''):
                changes[ref] = f'<v xmlns="{NS}">0</v>'
                continue
            formule, denominateurs = proteger_division(valeur)
            if formule == valeur:
                continue
            cellule = cellules[ref]
            f = cellule.find('{'+NS+'}f')
            if f is None or f.attrib:
                continue  # Ne pas transformer isolément une formule matricielle/partagée.
            cache = cellule.find('{'+NS+'}v')
            resultat = deepcopy(cache) if cache is not None else None
            # Le nouveau résultat vaut zéro uniquement si une case du diviseur
            # est réellement vide/nulle ; les autres caches sont préservés.
            for denom in denominateurs:
                d = cellules.get(denom.replace('$','').upper())
                v = d.find('{'+NS+'}v') if d is not None else None
                if ((d is None or (d.find('{'+NS+'}f') is None and v is None))
                        or (d is not None and d.get('t') not in ('e','s','str','inlineStr')
                            and v is not None and v.text is not None and float(v.text)==0)):
                    resultat = ET.Element('{'+NS+'}v');resultat.text='0'
                    break
            formule_xml = f'<f xmlns="{NS}">{html.escape(formule[1:])}</f>'
            cache_xml = ET.tostring(resultat,encoding='unicode') if resultat is not None else ''
            # Les erreurs autres que la division vide restent des erreurs Excel.
            if cellule.get('t') == 'e' and resultat is not None and resultat.text != '0':
                continue
            changes[ref] = formule_xml + cache_xml
        if not changes:
            return 0
        if path.with_name('~$'+path.name).exists():
            raise PermissionError('Enregistrez et fermez le PR pour compléter ses cases numériques à zéro.')
        par_ligne = {}
        for ref, contenu in changes.items():
            par_ligne.setdefault(int(re.search(r'\d+', ref)[0]), []).append((ref,contenu))
        data = re.search(r'<((?:\w+:)?sheetData)\b[^>]*>(.*?)</(?:\w+:)?sheetData>',original,re.S)
        lignes_xml = list(re.finditer(r'<(?:\w+:)?row\b[^>]*?(?:/>|>.*?</(?:\w+:)?row>)',data[2],re.S))
        existantes = {int(re.search(r'\br="(\d+)"',m[0])[1]) for m in lignes_xml}
        nouvelles = iter(sorted(set(par_ligne)-existantes))
        prochaine = next(nouvelles,None)
        def remplir(contenu, numero):
            for ref, valeur in par_ligne.get(numero,()):
                contenu = _cellule(contenu,ref,valeur)
            return contenu
        morceaux, position = [], 0
        for ligne_xml in lignes_xml:
            numero = int(re.search(r'\br="(\d+)"',ligne_xml[0])[1])
            morceaux.append(data[2][position:ligne_xml.start()])
            while prochaine is not None and prochaine < numero:
                morceaux.append(remplir(f'<row xmlns="{NS}" r="{prochaine}"></row>',prochaine))
                prochaine = next(nouvelles,None)
            morceaux.append(remplir(ligne_xml[0],numero))
            position = ligne_xml.end()
        while prochaine is not None:
            morceaux.append(remplir(f'<row xmlns="{NS}" r="{prochaine}"></row>',prochaine))
            prochaine = next(nouvelles,None)
        morceaux.append(data[2][position:])
        xml = original[:data.start(2)] + ''.join(morceaux) + original[data.end(2):]
        ET.fromstring(xml)
        fd, nom = tempfile.mkstemp(prefix='.zeros_pr_',suffix=path.suffix,dir=path.parent)
        os.close(fd)
        temp = Path(nom)
        try:
            with ZipFile(temp,'w') as dst:
                for info in src.infolist():
                    data = src.read(info.filename)
                    if info.filename == membre:
                        data = xml.encode('utf-8')
                    elif info.filename == 'xl/workbook.xml':
                        data = _recalcul_excel(data.decode('utf-8')).encode('utf-8')
                    dst.writestr(deepcopy(info),data)
            # Windows ne permet pas le remplacement d'un fichier encore ouvert.
            src.close()
            shutil.copymode(path,temp)
            def verifier():
                if _signature(path)!=signature or path.with_name('~$'+path.name).exists():
                    raise ValueError('Le PR a changé ou a été ouvert pendant la préparation. Recommencez.')
            verifier()
            if sauvegarder:
                sauvegarder(path)
            verifier()
            os.replace(temp,path)
        finally:
            temp.unlink(missing_ok=True)
    return len(changes)

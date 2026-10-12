"""Pilotage des options depuis l'état, avec exports PDF réels sur jeux d'essai."""
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from main import HorizonChantierApp, _recalculer_etat, _reporter_quantites_cloture
from options_pilotage import lire_options_pilotage, _coefficient


POSITIONS = {'marche': (30, 16), 'cumule': (32, 16), 'mois': (34, 16)}


def classeur_option(active=True, public=True):
    w = Workbook()
    s = w.active
    s.title = 'Bordereau'
    if public:
        s['B1'], s['K1'], s['L1'] = 'N°', 'QUANTITÉ', 'PRIX UNITAIRES'
    else:
        s['J1'], s['L1'] = 'Quantité', 'P.U.'
    for r, code, libelle, quantite, prix in ((24, '01', 'Travaux de base', 10, 100),
                                           (25, '02', 'OPTION 1 : Maçonnerie', 10, 50)):
        s[f'B{r}'], s[f'C{r}'], s[f'D{r}'] = code, libelle, 'QP'
        s[f'{"K" if public else "J"}{r}'], s[f'L{r}'] = quantite, prix
        s[f'{"N" if r == 25 and public else "M"}{r}'] = quantite*prix
        s[f'P{r}'], s[f'Q{r}'] = (2, 1) if r == 24 or active else (0, 0)
        s[f'R{r}'] = s[f'P{r}'].value+s[f'Q{r}'].value
        s[f'T{r}'], s[f'W{r}'] = s[f'R{r}'].value*prix, s[f'Q{r}'].value*prix
    for r, libelle, montant in ((30, 'Total soumission hors TVA', 1000),
                                (32, 'Total état cumulé hors TVA', 300),
                                (34, 'Total du mois hors TVA', 100),
                                (36, 'Options hors TVA (hors marché)', 500)):
        s.merge_cells(start_row=r, start_column=16, end_row=r, end_column=19)
        s[f'P{r}'], s[f'T{r}'] = libelle, montant
    return w


def enregistrer_formules(w, chemin, public=True, deja_inclus=False):
    w.save(chemin)
    f = load_workbook(chemin)
    s = f.active
    s['T30'] = '=SUM(M24' + (',N25' if public else ',M25') * deja_inclus + ')'
    s['T32'] = '=SUM(T24:T25)' if deja_inclus else '=SUM(T24)'
    s['W32'] = '=SUM(W24:W25)' if deja_inclus else '=SUM(W24)'
    s['T34'] = '=W32'
    f.save(chemin)
    f.close()


class OptionsPilotageTests(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.base = Path(self.tmp.name)
        self.chemin = self.base / 'Etat_avancement_02.xlsm'

    def test_option_chiffree_non_portee_vaut_zero(self):
        w = classeur_option(active=False)
        enregistrer_formules(w, self.chemin)
        r = lire_options_pilotage(self.chemin, w.active, POSITIONS)
        self.assertTrue(r['presentes'])
        self.assertTrue(all(v == 0 for k, v in r.items() if k != 'presentes'))
        w.close()

    def test_option_portee_marche_et_execution_distincts_public_prive(self):
        for public in (True, False):
            with self.subTest(public=public):
                w = classeur_option(public=public)
                enregistrer_formules(w, self.chemin, public=public)
                avant = self.chemin.read_bytes()
                r = lire_options_pilotage(self.chemin, w.active, POSITIONS)
                for cle, attendu in (('marche', 500), ('mois', 50), ('cumule', 150)):
                    self.assertEqual(r[cle], attendu)
                    self.assertEqual(r['ajout_'+cle], attendu)
                self.assertEqual(self.chemin.read_bytes(), avant)
                w.close()

    def test_synthese_incluant_options_ne_compte_pas_deux_fois(self):
        w = classeur_option()
        enregistrer_formules(w, self.chemin, deja_inclus=True)
        r = lire_options_pilotage(self.chemin, w.active, POSITIONS)
        self.assertEqual((r['marche'], r['mois'], r['cumule']), (500, 50, 150))
        self.assertEqual((r['ajout_marche'], r['ajout_mois'], r['ajout_cumule']), (0, 0, 0))
        w.close()

    def test_synthese_manuelle_fait_foi(self):
        w = classeur_option()
        w.active['T30'], w.active['T32'], w.active['T34'] = 1500, 450, 150
        w.save(self.chemin)
        r = lire_options_pilotage(self.chemin, w.active, POSITIONS)
        self.assertEqual((r['ajout_marche'], r['ajout_mois'], r['ajout_cumule']), (0, 0, 0))
        w.close()

    def test_total_manuel_repris_par_formule_ne_double_pas_option(self):
        w = classeur_option()
        enregistrer_formules(w, self.chemin)
        f = load_workbook(self.chemin)
        f.active['W32'] = 150
        f.save(self.chemin)
        f.close()
        self.assertEqual(lire_options_pilotage(self.chemin, w.active, POSITIONS)['ajout_mois'], 0)
        w.close()

    def test_generation_future_et_deplacement_du_chantier(self):
        from bordereau_public import habiller_avancement
        w = classeur_option()
        debut = habiller_avancement(w)
        w.save(self.chemin)
        s = w.active
        for r, prix in ((24, 100), (25, 50)):
            s[f'R{r}'], s[f'T{r}'], s[f'W{r}'] = 3, 3*prix, prix
        s[f'T{debut-2}'], s[f'T{debut}'], s[f'T{debut+2}'] = 1000, 300, 100
        positions = {'marche': (debut-2, 16), 'cumule': (debut, 16), 'mois': (debut+2, 16)}
        dossier_deplace = self.base / 'Autre chantier déplacé'
        dossier_deplace.mkdir()
        deplace = dossier_deplace / self.chemin.name
        self.chemin.rename(deplace)
        r = lire_options_pilotage(deplace, s, positions)
        self.assertEqual((r['ajout_marche'], r['ajout_mois'], r['ajout_cumule']), (500, 50, 150))
        w.close()

    def test_cloture_conserve_marche_cumule_et_remet_option_du_mois_a_zero(self):
        w = classeur_option()
        enregistrer_formules(w, self.chemin)
        _reporter_quantites_cloture(w.active)
        _recalculer_etat(w.active)
        r = lire_options_pilotage(self.chemin, w.active, POSITIONS)
        self.assertEqual((r['marche'], r['mois'], r['cumule']), (500, 0, 150))
        self.assertEqual((r['ajout_marche'], r['ajout_mois'], r['ajout_cumule']), (500, 0, 150))
        self.assertEqual(w.active['P25'].value, 3)
        w.close()

    def test_vrai_zero_accepte_mais_cache_manquant_refuse(self):
        w = classeur_option()
        enregistrer_formules(w, self.chemin)
        w.active['N25'] = 0
        self.assertEqual(lire_options_pilotage(self.chemin, w.active, POSITIONS)['marche'], 0)
        w.active['T25'] = None
        with self.assertRaisesRegex(ValueError, 'T25'):
            lire_options_pilotage(self.chemin, w.active, POSITIONS)
        w.close()

    def test_plusieurs_options_seule_celle_portee_compte(self):
        w = classeur_option()
        s = w.active
        s['B26'], s['C26'], s['D26'] = '03', 'Option 2', 'QP'
        s['K26'], s['L26'], s['N26'] = 20, 100, 2000
        s['P26'], s['Q26'], s['R26'], s['T26'], s['W26'] = 0, 0, 0, 0, 0
        enregistrer_formules(w, self.chemin)
        self.assertEqual(lire_options_pilotage(self.chemin, s, POSITIONS)['marche'], 500)
        w.close()

    def test_formule_somme_moins_option_et_double_comptage(self):
        w = classeur_option()
        s = w.active
        s['T32'] = '=SUM(T24:T25)-T25'
        self.assertEqual(_coefficient(s, 'T32', 'T25'), 0)
        s['T32'] = '=SUM(T24:T25)+T25'
        self.assertEqual(_coefficient(s, 'T32', 'T25'), 2)
        s['T32'] = '=IF(A1=1,T25,0)'
        with self.assertRaises(ValueError):
            _coefficient(s, 'T32', 'T25')
        w.close()

    def test_pilotage_pdf_reel_et_popup_chiffres_sans_modifier_etat(self):
        import matplotlib
        matplotlib.use('Agg')
        from pypdf import PdfReader
        for active in (False, True):
            with self.subTest(active=active):
                w = classeur_option(active=active)
                enregistrer_formules(w, self.chemin)
                avant = self.chemin.read_bytes()
                app = SimpleNamespace(_nom_chantier_selectionne=lambda: 'Chantier options',
                                      _dossier_chantier_selectionne=lambda: self.base)
                with patch('main._charger_valeurs_excel', return_value=w), \
                        patch('main.dossier_base', return_value=self.base), \
                        patch('main.ouvrir_chemin') as ouvrir, \
                        patch('main.messagebox.showinfo') as information, \
                        patch('main.messagebox.showerror') as erreur:
                    HorizonChantierApp.pilotage_chantier(app)
                erreur.assert_not_called()
                ouvrir.assert_called_once()
                pdf = PdfReader(self.base / 'pilotage_chantier.pdf')
                self.assertEqual(len(pdf.pages), 1)
                texte = pdf.pages[0].extract_text()
                self.assertIn('Options portées', texte)
                self.assertIn('Dont options du mois', texte)
                popup = information.call_args.args[1]
                for libelle, montant in (('Montant soumission', 1000),
                                         ('Marché total', 1500 if active else 1000),
                                         ('Réalisé du mois', 150 if active else 100),
                                         ('Réalisé cumulé', 450 if active else 300)):
                    ligne = next(l for l in popup.splitlines() if libelle in l)
                    self.assertIn(f'{montant:,.2f} €', ligne)
                self.assertEqual(self.chemin.read_bytes(), avant)


if __name__ == '__main__':
    unittest.main()

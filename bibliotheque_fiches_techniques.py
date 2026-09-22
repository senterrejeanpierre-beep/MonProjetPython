"""Vue classée et portable de la bibliothèque commune de fiches techniques."""

import hashlib
import re
import shutil
import tempfile
import unicodedata
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


SOURCE = "FICHES TECHNIQUES"
CLASSEE = "Bibliotheque_Technique"
ANCIENNES_CLASSEES = ("FICHES TECHNIQUES CLASSEES", "Bibliothèque Technique", "Bibliotheque_technique")
EXTENSIONS = {".pdf", ".doc", ".docx", ".xls", ".xlsx", ".ods", ".odt"}
ORDRE_GROUPES = ("Bâtiment", "Échafaudage", "Pierre", "Peinture", "Acier", "Toiture", "Autres")


def _sans_accents(texte: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", texte.casefold())
                   if not unicodedata.combining(c))


def _nom_portable(nom: str) -> str:
    nom = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", nom).rstrip(" .")
    base = Path(nom).stem.upper()
    if base in {"CON", "PRN", "AUX", "NUL"} or re.fullmatch(r"(?:COM|LPT)[1-9]", base):
        nom = "_" + nom
    suffixe = Path(nom).suffix
    if len(nom) > 150:
        nom = nom[:150 - len(suffixe)].rstrip(" .") + suffixe
    return nom or "Document"


def _categorie(chemin: Path) -> tuple[str, str]:
    morceaux = [_sans_accents(p) for p in chemin.parts]
    nom = morceaux[-1]
    parent = " ".join(morceaux[:-1])
    if "nit 249" in parent or "nit 249" in nom:
        return "Peinture", "Peintures et protections"
    if "toiture" in nom or "couverture" in nom:
        return "Toiture", "Couverture et ardoises"
    if "ceco-" in nom or "scierie" in nom:
        return "Bâtiment", "Charpente et poutres"
    if "pouzzolane" in nom:
        return "Bâtiment", "Béton et granulats"
    if "cabine" in nom or "modulco" in nom:
        return "Autres", "Chapiteaux et abris"
    if "boehm" in nom and "chaux" in nom:
        return "Bâtiment", "Chaux"
    if "pierre" in nom and "durcisseur" in nom:
        return "Pierre", "Protection et restauration"
    if "echafaud" in nom or "etais" in nom or "offre adria" in parent:
        return "Échafaudage", "Échafaudages et étais"
    if "travail de la pierres" in parent:
        return "Pierre", "Travail de la pierre"
    if "pierres" in parent or "pierre bleue" in nom:
        return "Pierre", "Pierres et restauration"
    if "peinture antirouille" in parent:
        return "Peinture", "Peinture antirouille"
    if "peinture et durcisseur" in parent:
        return "Peinture", "Peintures et protections"
    if "anti-mousse" in parent or "algicid" in nom or "biocid" in nom:
        return "Peinture", "Traitements biocides"
    if "acier maison" in parent or "acier" in nom:
        return "Acier", "Acier et métaux"
    if "couverture" in parent or "ardoise" in nom or "nit 219" in parent:
        return "Toiture", "Couverture et ardoises"
    if any(t in parent for t in ("zinc", "rive en alu")) or "zinc" in nom:
        return "Toiture", "Zinc et rives"
    if "toles bardage" in parent or "panneaux tuiles" in nom:
        return "Toiture", "Bardage"
    if "ondulee" in parent:
        return "Toiture", "Plaques ondulées"
    if "briques" in parent or "brique" in nom:
        return "Bâtiment", "Briques"
    if "chaux" in parent or "chaux" in nom:
        return "Bâtiment", "Chaux"
    if "mortier de resine" in parent:
        return "Bâtiment", "Mortiers de résine"
    if "mortier colore" in parent:
        return "Bâtiment", "Mortiers de jointoiement"
    if "mortier" in parent or "mortier" in nom:
        return "Bâtiment", "Mortiers"
    if "barre en fibre" in parent or "barres en fibre" in parent:
        return "Bâtiment", "Renforts en fibre de verre"
    if "murfor" in parent:
        return "Bâtiment", "Renforts de maçonnerie"
    if "crochets" in parent:
        return "Bâtiment", "Crochets pour murs"
    if "charpente" in parent:
        return "Bâtiment", "Charpente et poutres"
    if "bois" in parent or "osb" in parent or "multiplex" in parent or "resineux" in nom:
        return "Bâtiment", "Bois et panneaux"
    if "foamglas" in parent or "rocwool" in parent or "isolant" in parent:
        return "Bâtiment", "Isolants"
    if "menbanes" in parent or "axo" in parent or "etancheite" in parent:
        return "Bâtiment", "Étanchéité"
    if "film pe" in parent:
        return "Bâtiment", "Films de protection"
    if "chassis" in parent:
        return "Bâtiment", "Châssis"
    if "tuyaux pvc" in parent or "drain" in parent:
        return "Bâtiment", "Tuyaux et drainage"
    if "gravier" in nom or "beton" in nom:
        return "Bâtiment", "Béton et granulats"
    if "catalogue et fiches technique materiel" in parent:
        return "Bâtiment", "Matériel de chantier"
    if "chapiteau" in parent:
        return "Autres", "Chapiteaux et abris"
    if "signalisation" in parent or "protection anti poussiere" in parent:
        return "Autres", "Sécurité et protections"
    return "Autres", "Documents divers"


def _documents(source: Path):
    fichiers = (p for p in source.rglob("*")
                if p.is_file() and p.suffix.lower() in EXTENSIONS and not p.name.startswith("~$"))
    def priorite(p):
        relatif = p.relative_to(source)
        premier = _sans_accents(relatif.parts[0])
        rang = 0 if premier == "fiches techniques 1" else 1 if len(relatif.parts) == 1 else 2
        return rang, relatif.as_posix().casefold()
    vus = set()
    for fichier in sorted(fichiers, key=priorite):
        empreinte = hashlib.sha256(fichier.read_bytes()).digest()
        if empreinte in vus:
            continue
        vus.add(empreinte)
        yield fichier


def creer_bibliotheque_classee(base: Path, alimenter_depuis_source: bool = False) -> tuple[Path, int, int]:
    """Réutilise la bibliothèque ; la recherche dans les sources est explicite."""
    base = Path(base)
    source = base / SOURCE
    sortie = base / CLASSEE
    anciennes = [base / nom for nom in ANCIENNES_CLASSEES if (base / nom).exists()
                 and not (sortie.exists() and (base / nom).samefile(sortie))]
    if anciennes:
        if sortie.exists() or len(anciennes) > 1:
            raise ValueError(
                f"Plusieurs bibliothèques techniques existent dans {base}. "
                "Leur regroupement est nécessaire ; aucun document n'a été remplacé.")
        anciennes[0].rename(sortie)
    if sortie.is_dir():
        for dossier in base.iterdir():
            if dossier.name != CLASSEE and dossier.samefile(sortie):
                dossier.rename(sortie)
                break
        return sortie, 0, 0
    if not alimenter_depuis_source or not source.is_dir():
        sortie.mkdir()
        return sortie, 0, 0
    groupes = defaultdict(list)
    for document in _documents(source):
        groupes[_categorie(document.relative_to(source))].append(document)
    categories = sorted(groupes, key=lambda c: (ORDRE_GROUPES.index(c[0]), c[1].casefold()))
    with tempfile.TemporaryDirectory(prefix=".fiches_classees_", dir=base) as tmp:
        preparation = Path(tmp) / CLASSEE
        preparation.mkdir()
        wb = Workbook()
        ws = wb.active
        ws.title = "Catégories"
        ws.append(["N°", "Métier", "Catégorie", "Documents", "Accès"])
        wd = wb.create_sheet("Documents")
        wd.append(["N°", "Métier", "Catégorie", "Fiche technique", "Accès"])
        total = 0
        for numero, (groupe, categorie) in enumerate(categories, 1):
            dossier_groupe = preparation / groupe
            dossier_groupe.mkdir(exist_ok=True)
            dossier = dossier_groupe / f"{numero:03d}_{_nom_portable(categorie).replace(' ', '_')}"
            dossier.mkdir()
            utilises = set()
            for document in sorted(groupes[(groupe, categorie)], key=lambda p: p.name.casefold()):
                nom = _nom_portable(document.name)
                racine, suffixe = Path(nom).stem, Path(nom).suffix
                candidat = nom
                index = 2
                while candidat.casefold() in utilises:
                    candidat = f"{racine}_{index}{suffixe}"
                    index += 1
                utilises.add(candidat.casefold())
                cible = dossier / candidat
                shutil.copy2(document, cible)
                wd.append([f"{numero:03d}", groupe, categorie, candidat, "Ouvrir"])
                lien = wd.cell(wd.max_row, 5)
                lien.hyperlink = cible.relative_to(preparation).as_posix()
                lien.style = "Hyperlink"
                total += 1
            ws.append([f"{numero:03d}", groupe, categorie, len(groupes[(groupe, categorie)]), "Ouvrir"])
            lien = ws.cell(ws.max_row, 5)
            lien.hyperlink = dossier.relative_to(preparation).as_posix() + "/"
            lien.style = "Hyperlink"
        for feuille in (ws, wd):
            for cellule in feuille[1]:
                cellule.font = Font(bold=True, color="FFFFFF")
                cellule.fill = PatternFill("solid", fgColor="174A72")
            feuille.freeze_panes = "B2"
            feuille.auto_filter.ref = f"A1:E{feuille.max_row}"
            for col, width in {"A": 8, "B": 18, "C": 34, "D": 55, "E": 16}.items():
                feuille.column_dimensions[col].width = width
        wb.save(preparation / "Index_fiches_techniques.xlsx")
        wb.close()
        preparation.rename(sortie)
    return sortie, len(categories), total


def copier_fiche_au_chantier(fiche: Path, bibliotheque: Path, chantier: Path) -> tuple[Path, bool]:
    """Copie une fiche choisie dans Technique sans écraser un document du chantier."""
    fiche, bibliotheque, chantier = Path(fiche), Path(bibliotheque), Path(chantier)
    fiche.resolve().relative_to(bibliotheque.resolve())
    if not fiche.is_file() or fiche.suffix.lower() not in EXTENSIONS:
        raise ValueError("Choisissez un document de la bibliothèque classée.")
    destination = chantier / "Technique"
    destination.mkdir(parents=True, exist_ok=True)
    nom = _nom_portable(fiche.name)
    candidat = destination / nom
    empreinte = hashlib.sha256(fiche.read_bytes()).digest()
    for existant in destination.iterdir():
        if existant.is_file() and hashlib.sha256(existant.read_bytes()).digest() == empreinte:
            return existant, False
    compteur = 2
    while candidat.exists():
        candidat = destination / f"{Path(nom).stem}_{compteur}{Path(nom).suffix}"
        compteur += 1
    shutil.copy2(fiche, candidat)
    return candidat, True

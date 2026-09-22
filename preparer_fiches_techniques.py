"""Prépare des dossiers de fiches numérotés et leur index par chantier."""

import re
import unicodedata
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


def nom_dossier(numero: str, categorie: str) -> str:
    texte = unicodedata.normalize("NFKD", categorie)
    texte = "".join(c for c in texte if not unicodedata.combining(c))
    texte = re.sub(r"[^A-Za-z0-9]+", "_", texte).strip("_")
    return f"{numero}_{texte or 'Fiches'}"


def lire_categories(index_general: Path) -> list[tuple[str, str]]:
    classeur = load_workbook(index_general, read_only=True, data_only=True)
    try:
        return [(str(numero), str(nom)) for numero, nom, *_ in
                classeur["Catégories"].iter_rows(min_row=5, values_only=True)
                if numero and nom]
    finally:
        classeur.close()


def preparer_chantier(dossier_chantier: Path, categories: list[tuple[str, str]]) -> Path:
    """Ne remplace aucun dossier ni index existant."""
    racine = dossier_chantier / "Fiches_techniques"
    racine.mkdir(exist_ok=True)
    dossiers = []
    for numero, categorie in categories:
        candidats = list(racine.glob(f"{numero}_*"))
        dossier = candidats[0] if candidats else racine / nom_dossier(numero, categorie)
        dossier.mkdir(exist_ok=True)
        dossiers.append((numero, categorie, dossier))

    index = racine / "Index_fiches_techniques.xlsx"
    if index.exists():
        return index
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Fiches du chantier"
    feuille.append(["N°", "Fiche technique", "Accès"])
    for cellule in feuille[1]:
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor="174A72")
    for numero, categorie, dossier in dossiers:
        feuille.append([numero, categorie, "Ouvrir le dossier"])
        lien = feuille.cell(feuille.max_row, 3)
        lien.hyperlink = dossier.name + "/"
        lien.style = "Hyperlink"
    feuille.column_dimensions["A"].width = 9
    feuille.column_dimensions["B"].width = 52
    feuille.column_dimensions["C"].width = 23
    feuille.freeze_panes = "B2"
    feuille.auto_filter.ref = f"A1:C{feuille.max_row}"
    try:
        classeur.save(index)
    finally:
        classeur.close()
    return index

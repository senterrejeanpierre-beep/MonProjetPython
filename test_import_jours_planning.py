import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

from import_jours_planning import importer_jours_pr


class ImportJoursPlanning(unittest.TestCase):
    def test_import_explicit_preserve_planning_et_debuts(self):
        with tempfile.TemporaryDirectory() as tmp:
            dossier = Path(tmp)
            pr = dossier / 'prix_de_revient.xlsx'
            wb = Workbook()
            ws = wb.active
            ws.title = 'Chiffrage'
            for debut, numero, jours in ((2, '01', 3.048), (44, '02', 2)):
                ws.cell(debut, 2, 'Article : N° code CSTC')
                ws.cell(debut + 1, 2, numero)
                ws.cell(debut + 1, 3, f'Poste {numero}')
                ws.cell(debut + 15, 7, jours)
            wb.save(pr)
            wb.close()
            planning = dossier / 'Planning_soumission.xlsx'
            wb = Workbook()
            wb.active.title = 'PlanningProjet'
            wb.active['B15'] = 'Ancien chantier'
            wb.active['B22'] = 'Ancienne phase'
            for ligne in range(23, 63):
                for col in 'BCEFG':
                    wb.active[f'{col}{ligne}'] = None
            wb.active['G23'] = '=F23-E23+1'
            wb.active['B23'].hyperlink = 'https://example.invalid/document'
            wb.create_sheet('À propos de')['A1'] = 'Texte conservé'
            wb.save(planning)
            wb.close()
            with ZipFile(planning) as z:
                origine = {n: z.read(n) for n in z.namelist()}
            self.assertEqual(importer_jours_pr(pr, planning, nom='Chantier test', client='Client test'),
                             (2, 0, True))
            with ZipFile(planning) as z:
                for nom, donnees in origine.items():
                    if nom not in ('xl/workbook.xml', 'xl/_rels/workbook.xml.rels', '[Content_Types].xml',
                                   'xl/worksheets/sheet1.xml'):
                        self.assertEqual(z.read(nom), donnees, nom)
            wb = load_workbook(planning)
            self.assertEqual(wb.active.title, 'PlanningProjet')
            self.assertEqual(wb['Postes PR']['A2'].value, '01')
            self.assertEqual(wb['Postes PR']['C2'].value, 3.048)
            self.assertEqual(wb['PlanningProjet']['B15'].value, 'Planning de soumission - Chantier test')
            self.assertEqual(wb['PlanningProjet']['B23'].value, '01')
            self.assertEqual(wb['PlanningProjet']['C23'].value, 'Poste 01')
            self.assertEqual(wb['PlanningProjet']['E23'].value, '=Début_Projet+1')
            self.assertEqual(wb['PlanningProjet']['F23'].value, '=E23+2.048')
            self.assertEqual(wb['PlanningProjet']['G23'].value, '=F23-E23+1')
            self.assertEqual(wb['PlanningProjet']['B23'].hyperlink.target,
                             'https://example.invalid/document')
            wb['Postes PR']['D2'] = 4
            wb.save(planning)
            wb.close()
            wb['PlanningProjet']['E23'] = 4
            wb.save(planning)
            wb.close()
            self.assertEqual(importer_jours_pr(pr, planning, nom='Chantier test', client='Client test'),
                             (2, 0, True))
            wb = load_workbook(planning, read_only=True)
            self.assertEqual(wb.active.title, 'PlanningProjet')
            self.assertEqual(wb['Postes PR']['D2'].value, 4)
            self.assertEqual(wb['PlanningProjet']['E23'].value, 4)
            self.assertIn('MAX(E2:E3)', wb['Postes PR']['E5'].value)
            wb.close()
            self.assertEqual(importer_jours_pr(pr, planning, nom='Chantier test', client='Client test'),
                             (2, 0, False))


if __name__ == '__main__':
    unittest.main()

"""Options portées dans l'état : lecture seule et complément des synthèses.

Les synthèses saisies manuellement restent la référence. Pour les synthèses
calculées, seule la part d'option absente de la formule est ajoutée au pilotage.
"""
import ast
import math
import re
import unicodedata

from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import range_boundaries
from bordereau_public import est_metre, lignes_postes


def _normaliser(valeur):
    return ''.join(c for c in unicodedata.normalize('NFKD', str(valeur or '').lower())
                   if not unicodedata.combining(c))


def _nombre(ws, adresse, obligatoire=False):
    valeur = ws[adresse].value
    if valeur in (None, '', '-', '—'):
        if obligatoire:
            raise ValueError(f"Option : montant ou quantité absent en {adresse}. "
                             "Recalculez et enregistrez l’état dans Excel.")
        return 0.0
    try:
        nombre = float(re.sub(r'\s+', '', str(valeur)).replace('€', '').replace('−', '-').replace(',', '.'))
        if not math.isfinite(nombre):
            raise ValueError
        return nombre
    except (TypeError, ValueError):
        raise ValueError(f"Option : valeur illisible en {adresse} : {valeur!r}.") from None


def _coefficient(ws, adresse, cible, visites=None):
    """Part d'une cellule dans une synthèse SUM, avec références intermédiaires.

    Pas d'évaluation des montants Excel : cette lecture sert seulement à éviter
    un double comptage. Une construction ambiguë n'est jamais assimilée à zéro.
    """
    adresse = adresse.replace('$', '').upper()
    if adresse == cible:
        return 1
    visites = set() if visites is None else visites
    if adresse in visites:
        raise ValueError(f"Référence circulaire dans la synthèse en {adresse}.")
    formule = ws[adresse].value
    if not isinstance(formule, str) or not formule.startswith('='):
        # Un total manuel peut aussi être repris par une simple référence (=W32).
        cellule = ws[adresse]
        if cellule.column >= 16 and any(
            _normaliser(ws.cell(cellule.row, col).value).startswith((
                'total soumission', 'total etat cumule', 'total du mois',
                'total execute', 'montant de l’etat', "montant de l'etat",
                'montant global etat hors tva')) for col in range(16, cellule.column)):
            return 1
        return 0
    visites = visites | {adresse}
    termes, references = [], {}
    for token in Tokenizer(formule).items:
        if token.type == 'OPERAND' and token.subtype == 'RANGE':
            reference = token.value.replace('$', '')
            if '!' in reference:
                feuille, reference = reference.rsplit('!', 1)
                feuille = feuille.strip("'").replace("''", "'")
                if feuille.casefold() != ws.title.casefold() and '[' not in feuille:
                    raise ValueError(f"La synthèse {adresse} renvoie à la feuille {feuille}. "
                                     "La présence de l’option dans ce total ne peut pas être vérifiée.")
                if feuille.casefold() != ws.title.casefold():
                    reference = None
            coefficient = 0
            if reference is not None:
                if not re.fullmatch(r'[A-Za-z]+[1-9]\d*(?::[A-Za-z]+[1-9]\d*)?', reference):
                    raise ValueError(f"Référence non reconnue dans la synthèse {adresse} : {reference}.")
                c1, r1, c2, r2 = range_boundaries(reference)
                if (r2-r1+1)*(c2-c1+1) > 100000:
                    raise ValueError(f"Plage trop large dans la synthèse {adresse}.")
                for ligne in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2):
                    coefficient += sum(_coefficient(ws, c.coordinate, cible, visites) for c in ligne)
            nom = f'ref{len(references)}'
            references[nom] = coefficient
            termes.append(nom)
        elif token.type == 'OPERAND':
            # Les constantes ne contiennent aucune contribution de l'option.
            termes.append('0')
        else:
            termes.append(token.value)

    def lire(noeud):
        if isinstance(noeud, ast.Expression):
            return lire(noeud.body)
        if isinstance(noeud, ast.Constant):
            return 0
        if isinstance(noeud, ast.Name):
            return references[noeud.id]
        if isinstance(noeud, ast.UnaryOp) and isinstance(noeud.op, (ast.UAdd, ast.USub)):
            return lire(noeud.operand) * (-1 if isinstance(noeud.op, ast.USub) else 1)
        if isinstance(noeud, ast.BinOp):
            gauche, droite = lire(noeud.left), lire(noeud.right)
            if isinstance(noeud.op, ast.Add):
                return gauche + droite
            if isinstance(noeud.op, ast.Sub):
                return gauche - droite
            if gauche == droite == 0:
                return 0
        if isinstance(noeud, ast.Call) and isinstance(noeud.func, ast.Name):
            valeurs = [lire(a) for a in noeud.args]
            if noeud.func.id.upper() in ('SUM', 'SOMME'):
                return sum(valeurs)
            if not any(valeurs):
                return 0
        raise ValueError(f"Impossible de vérifier si l’option est déjà comprise dans {adresse}.")

    try:
        return lire(ast.parse(''.join(termes).strip(), mode='eval'))
    except (SyntaxError, KeyError) as erreur:
        raise ValueError(f"Formule de synthèse non reconnue en {adresse}.") from erreur


def _adresse_montant(ws, position):
    ligne, colonne = position
    if ligne is None:
        return None
    for fusion in ws.merged_cells.ranges:
        if fusion.min_row <= ligne <= fusion.max_row and fusion.min_col <= colonne <= fusion.max_col:
            colonne = fusion.max_col
            break
    for col in range(colonne+1, min(colonne+3, ws.max_column)+1):
        c = ws.cell(ligne, col)
        if c.value not in (None, '', '-', '—'):
            return c.coordinate
    raise ValueError(f"Montant de synthèse absent à la ligne {ligne}.")


def lire_options_pilotage(chemin, ws, positions):
    """Une option sans exécution vaut zéro, même si son prix est renseigné."""
    public = est_metre(ws)
    lignes = ([r for r in lignes_postes(ws) if ws[f'M{r}'].value is None] if public else
              [r for r in range(24, ws.max_row+1) if ws[f'B{r}'].value is not None
               and re.match(r'^\s*options?(?=\b|\d)', _normaliser(ws[f'C{r}'].value))])
    resultat = {'presentes': bool(lignes), 'marche': 0.0, 'mois': 0.0, 'cumule': 0.0,
                'ajout_marche': 0.0, 'ajout_mois': 0.0, 'ajout_cumule': 0.0}
    actives = []
    for r in lignes:
        # Le cumul conserve l'option après clôture, quand le mois revient à zéro.
        quantites = [_nombre(ws, f'{col}{r}') for col in ('P', 'Q', 'R')]
        if not any(quantites):
            continue
        if public:
            col_montant = 'N'
        elif ws['E8'].value == 'Quantités calculées par' and ws['H8'].value == 'Prix unitaires':
            col_montant = 'J'
        else:
            col_montant = 'M'
        adresse = f'{col_montant}{r}'
        montant = _nombre(ws, adresse, obligatoire=True)
        mois = _nombre(ws, f'W{r}', obligatoire=True)
        cumule = _nombre(ws, f'T{r}', obligatoire=True)
        actives.append((r, adresse, montant, mois, cumule))
        resultat['marche'] += montant
        resultat['mois'] += mois
        resultat['cumule'] += cumule
    if not actives:
        return resultat
    formules = load_workbook(chemin, data_only=False)
    try:
        source = formules[ws.title]
        adresses = {cle: _adresse_montant(ws, position) for cle, position in positions.items()}
        for r, adresse, montant, mois, cumule in actives:
            for cle, cible, valeur in (('marche', adresse, montant), ('mois', f'W{r}', mois),
                                       ('cumule', f'T{r}', cumule)):
                synthese = adresses[cle]
                # Une synthèse manuelle est déjà le montant métier faisant foi.
                if synthese is None or source[synthese].data_type != 'f':
                    continue
                coefficient = _coefficient(source, synthese, cible)
                if coefficient not in (0, 1):
                    raise ValueError(f"L’option de la ligne {r} est comptée {coefficient} fois dans {synthese}.")
                resultat['ajout_'+cle] += (1-coefficient)*valeur
    finally:
        formules.close()
    return resultat

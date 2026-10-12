import hashlib
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from openpyxl import Workbook, load_workbook
from synthese_pr import lire_synthese_pr, generer_pdf_synthese, est_option

class SynthesePRTests(unittest.TestCase):
    def creer(self, path, manquant=False):
        w=Workbook();s=w.active;s.title='Chiffrage'
        for debut,article in ((3,'A'),(37,'B')):
            s.cell(debut-1,2,'Article : code');s.cell(debut,2,article);s.cell(debut,3,'Poste');s.cell(debut,5,2)
            for offset,label in ((25,'PR matière (revient)'),(26,'PR MO (revient)'),(27,'PR total')):
                s.cell(debut+offset,2,label)
            s.cell(debut+27,5,'PV HT')
            for offset,col,val in ((25,3,10),(26,3,20),(27,3,30),(25,10,0),(26,10,3),(27,6,33)):
                s.cell(debut+offset,col,val)
            s.cell(debut+26,6,0.1);s.cell(debut+13,10,55)
            s.cell(debut+13,4,'Équipe');s.cell(debut+13,9,16)
            s.cell(debut+13,7,80)
            s.cell(debut+14,6,'Total jours P');s.cell(debut+14,7,0)
            s.cell(debut+29,5,'PV TTC')
        if manquant:
            s['C28']='=P13'
            s['I16']='=H16*C16'
            s['G16']='=SUM(I16:I25)'
        w.save(path)
    def test_totaux_zeros_et_aucune_modification(self):
        with TemporaryDirectory() as d:
            p=Path(d)/'pr.xlsx';self.creer(p);avant=hashlib.sha256(p.read_bytes()).digest()
            r=lire_synthese_pr(p)
            self.assertEqual(r['totaux']['matiere'],40)
            self.assertEqual(r['totaux']['marge'],12)
            self.assertEqual(r['totaux']['jours'],0)
            self.assertEqual(r['totaux']['heures'],160)
            self.assertEqual(r['postes'][0]['heures'],80)
            self.assertEqual(r['taux_horaire'],55)
            self.assertAlmostEqual(r['taux_marge'],0.1)
            self.assertEqual(r['alertes'],[])
            self.assertEqual(r['options']['postes'],[])
            pdf=Path(d)/'rapport.pdf';generer_pdf_synthese(r,pdf,'Test')
            self.assertTrue(pdf.read_bytes().startswith(b'%PDF'))
            self.assertEqual(hashlib.sha256(p.read_bytes()).digest(),avant)
    def test_taux_appliques_ponderes_sans_utiliser_les_montants_mo(self):
        with TemporaryDirectory() as d:
            p=Path(d)/'pr.xlsx';self.creer(p)
            w=load_workbook(p);s=w.active
            s['I16']=8;s['J16']=40;s['I50']=24;s['J50']=60
            s['F29']=0.1;s['F63']=0.2;s['C64']=90
            w.save(p);w.close()
            r=lire_synthese_pr(p)
            self.assertAlmostEqual(r['taux_horaire'],55)
            self.assertAlmostEqual(r['taux_marge'],0.175)

    def test_cache_absent_ne_devient_pas_zero(self):
        with TemporaryDirectory() as d:
            p=Path(d)/'pr.xlsx';self.creer(p,True);r=lire_synthese_pr(p)
            self.assertIsNone(r['totaux']['matiere'])
            self.assertIsNone(r['totaux']['heures'])
            self.assertTrue(any('C28' in a for a in r['alertes']))

    def test_simulation_maconnerie_avec_sous_traitance_non_renseignee(self):
        with TemporaryDirectory() as d:
            p=Path(d)/'pr.xlsx';self.creer(p)
            w=load_workbook(p);s=w.active
            s['C3']='Maçonnerie';s['E3']=1
            for cellule,valeur in {'C28':6000,'C29':4000,'C30':10000,
                                  'J28':0,'J29':0,'F30':10000,'F29':0}.items():
                s[cellule]=valeur
            s['C37']='Sous-traitance à chiffrer'
            for cellule in ('E37','C62','C63','C64','J62','J63','F64','F63',
                            'D50','I50','J50','G50','G51'):
                s[cellule]=None
            w.save(p);w.close()
            avant=hashlib.sha256(p.read_bytes()).digest()
            r=lire_synthese_pr(p)
            self.assertEqual(r['totaux']['vente'],10000)
            self.assertEqual(r['totaux']['matiere'],6000)
            self.assertEqual(r['totaux']['mo'],4000)
            self.assertTrue(all(r['postes'][1][cle]==0 for cle in r['totaux']))
            self.assertEqual(r['alertes'],[])
            pdf=Path(d)/'simulation.pdf'
            generer_pdf_synthese(r,pdf,'Simulation partielle')
            self.assertTrue(pdf.read_bytes().startswith(b'%PDF'))
            self.assertEqual(hashlib.sha256(p.read_bytes()).digest(),avant)

    def test_pr_vide_et_ancien_bloc_sans_jours_donnent_zero(self):
        with TemporaryDirectory() as d:
            p=Path(d)/'pr.xlsx';self.creer(p)
            w=load_workbook(p);s=w.active
            for debut in (3,37):
                for offset,col in ((0,5),(25,3),(26,3),(27,3),(25,10),(26,10),
                                   (27,6),(26,6),(13,4),(13,9),(13,10),(13,7),(14,7),(14,6)):
                    s.cell(debut+offset,col).value=None
            w.save(p);w.close()
            r=lire_synthese_pr(p)
            self.assertTrue(all(v==0 for v in r['totaux'].values()))
            self.assertEqual((r['taux_horaire'],r['taux_marge']),(0,0))
            self.assertEqual(r['alertes'],[])

    def test_erreur_excel_ne_devient_pas_zero(self):
        with TemporaryDirectory() as d:
            p=Path(d)/'pr.xlsx';self.creer(p)
            w=load_workbook(p);w.active['C28']='#VALUE!';w.save(p);w.close()
            r=lire_synthese_pr(p)
            self.assertIsNone(r['totaux']['matiere'])
            self.assertTrue(any('C28' in a for a in r['alertes']))

    def test_jours_division_par_equipe_vide_sans_masquer_heures_saisies(self):
        from zipfile import ZipFile
        from xml.etree import ElementTree as ET
        ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
        for heures, attendu in ((0,0),(16,None)):
            with self.subTest(heures=heures), TemporaryDirectory() as d:
                p=Path(d)/'pr.xlsx';self.creer(p)
                w=load_workbook(p);s=w.active
                s['C16']=None;s['G16']=None;s['I16']=heures
                s['G17']='=I16/C16/8'
                w.save(p);w.close()
                # Résultat réellement enregistré par Excel dans le PR 005.
                copie=Path(d)/'cache.xlsx'
                with ZipFile(p) as src, ZipFile(copie,'w') as dst:
                    for info in src.infolist():
                        data=src.read(info.filename)
                        if info.filename=='xl/worksheets/sheet1.xml':
                            root=ET.fromstring(data)
                            cell=root.find(f'.//{ns}c[@r="G17"]')
                            cell.set('t','e');cell.find(ns+'v').text='#DIV/0!'
                            data=ET.tostring(root)
                        dst.writestr(info,data)
                r=lire_synthese_pr(copie)
                self.assertEqual(r['totaux']['jours'],attendu)
                self.assertEqual(any('G17' in a for a in r['alertes']),heures!=0)

    def test_option_separee_des_montants_jours_et_moyennes(self):
        with TemporaryDirectory() as d:
            p=Path(d)/'pr.xlsx';self.creer(p)
            w=load_workbook(p);s=w.active
            s['C37']='OPTION 1 : Ragréage des maçonneries'
            s['G17']=2;s['G51']=21
            s['J16']=40;s['J50']=80;s['F63']=0.3
            w.save(p);w.close()
            avant=hashlib.sha256(p.read_bytes()).digest()
            r=lire_synthese_pr(p)
            self.assertEqual(r['totaux']['vente'],66)
            self.assertEqual(r['totaux']['revient'],60)
            self.assertEqual(r['totaux']['jours'],2)
            self.assertEqual(r['totaux']['heures'],80)
            self.assertEqual(r['taux_horaire'],40)
            self.assertAlmostEqual(r['taux_marge'],0.1)
            o=r['options']
            self.assertEqual([v['article'] for v in o['postes']],['B'])
            self.assertEqual(o['totaux']['vente'],66)
            self.assertEqual(o['totaux']['jours'],21)
            self.assertEqual(o['taux_horaire'],80)
            self.assertAlmostEqual(o['taux_marge'],0.3)
            # Le consommateur planning garde les deux postes ; l'option n'est pas effacée.
            self.assertEqual([v['article'] for v in r['postes']],['A','B'])
            from unittest.mock import patch
            from reportlab.platypus import Table
            with patch('reportlab.platypus.Table',wraps=Table) as tables:
                generer_pdf_synthese(r,Path(d)/'synthese.pdf','Test')
            self.assertEqual(len(tables.call_args_list),2)
            base=tables.call_args_list[0].args[0]
            option=tables.call_args_list[1].args[0]
            self.assertEqual(len(base),2)
            self.assertEqual(base[1][0].getPlainText(),'A — Poste')
            self.assertIn('OPTION 1',option[1][0].getPlainText())
            self.assertEqual(option[-1][0].getPlainText(),'Total des options')
            self.assertEqual(option[-1][-1],'21,00')
            self.assertEqual(hashlib.sha256(p.read_bytes()).digest(),avant)

    def test_option_incomplete_ne_fausse_pas_la_base(self):
        with TemporaryDirectory() as d:
            p=Path(d)/'pr.xlsx';self.creer(p)
            w=load_workbook(p);s=w.active
            s['C37']='Option : mur'
            s['C62']='=P47';s['G51']=None;s['J50']=None
            w.save(p);w.close()
            r=lire_synthese_pr(p)
            self.assertEqual(r['totaux']['matiere'],20)
            self.assertEqual(r['totaux']['jours'],0)
            self.assertEqual(r['taux_horaire'],55)
            self.assertIsNone(r['options']['totaux']['matiere'])
            self.assertEqual(r['options']['totaux']['jours'],0)
            self.assertEqual(r['options']['taux_horaire'],0)
            self.assertTrue(r['alertes'])

    def test_reconnaissance_des_options_explicites(self):
        for texte in ('OPTION 1 : mur',' Option : mur','Options — cour','OPTION2 plateau'):
            self.assertTrue(est_option(texte),texte)
        for texte in ('Pose avec option de finition','Maçonnerie','Optionnellement renforcé'):
            self.assertFalse(est_option(texte),texte)

    def test_total_general_affiche_base_et_options_sans_masquer_un_resultat_absent(self):
        from unittest.mock import patch
        from reportlab.pdfgen.canvas import Canvas
        cas = (
            (False, 33, 11, '88,00 €'),
            (True, 33, 11, '88,00 €'),
            (True, 33, 0, '66,00 €'),
            (True, 0, 0, '0,00 €'),
            (True, 33, '=C64', 'Non disponible'),
            (True, '=C30', 11, 'Non disponible'),
        )
        for option, vente_base, vente_second_poste, attendu in cas:
            with self.subTest(option=option, base=vente_base, second=vente_second_poste), TemporaryDirectory() as d:
                p=Path(d)/'pr.xlsx';self.creer(p)
                w=load_workbook(p);s=w.active
                if option:
                    s['C37']='OPTION 1 : mur'
                s['F30']=vente_base;s['F64']=vente_second_poste
                w.save(p);w.close()
                avant=hashlib.sha256(p.read_bytes()).digest()
                r=lire_synthese_pr(p)
                original=Canvas.drawString
                original_droite=Canvas.drawRightString
                textes=[]
                def enregistrer(canvas,x,y,texte,*args,**kwargs):
                    textes.append(texte)
                    return original(canvas,x,y,texte,*args,**kwargs)
                def enregistrer_droite(canvas,x,y,texte,*args,**kwargs):
                    textes.append(texte)
                    return original_droite(canvas,x,y,texte,*args,**kwargs)
                with patch.object(Canvas,'drawString',enregistrer), patch.object(Canvas,'drawRightString',enregistrer_droite):
                    generer_pdf_synthese(r,Path(d)/'synthese.pdf','Test')
                indice=textes.index('Total général des prix de vente HT (options comprises)')
                self.assertEqual(textes[indice+1],attendu)
                if option and isinstance(vente_base,(int,float)):
                    self.assertEqual(r['totaux']['vente'],2*vente_base)
                self.assertEqual(hashlib.sha256(p.read_bytes()).digest(),avant)

    def test_graphique_pourcentage_options_sur_total_vente_ht(self):
        from unittest.mock import patch
        from reportlab.pdfgen.canvas import Canvas
        with TemporaryDirectory() as d:
            p=Path(d)/'pr.xlsx';self.creer(p)
            w=load_workbook(p);s=w.active
            s['C37']='OPTION 1 : mur';s['F64']=11
            w.save(p);w.close()
            r=lire_synthese_pr(p)
            self.assertEqual(r['totaux']['vente'],66)
            self.assertEqual(r['options']['totaux']['vente'],22)
            original=Canvas.drawString
            textes=[]
            def enregistrer(canvas,x,y,texte,*args,**kwargs):
                textes.append(texte)
                return original(canvas,x,y,texte,*args,**kwargs)
            with patch.object(Canvas,'drawString',enregistrer):
                generer_pdf_synthese(r,Path(d)/'graphique.pdf','Test')
            self.assertIn('Base : 75,0 %',textes)
            self.assertIn('Options : 25,0 %',textes)
            r['options']['totaux']['vente']=None
            textes.clear()
            with patch.object(Canvas,'drawString',enregistrer):
                generer_pdf_synthese(r,Path(d)/'incomplet.pdf','Test')
            self.assertIn('Répartition HT non disponible.',textes)
            self.assertFalse(any(t.startswith('Options : ') for t in textes))

if __name__=='__main__':unittest.main()

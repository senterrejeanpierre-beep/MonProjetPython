"""Une ligne par poste, avec les jours de début et de fin de chaque période."""
from copy import copy

from openpyxl.formatting.rule import FormulaRule
from openpyxl.formula.translate import Translator
from openpyxl.styles import PatternFill
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.filters import AutoFilter

COULEURS = ('4F81BD', '548235', 'C55A11', '8064A2')


def est_compact(ws):
    return ws['D8'].value == 'Période' and ws['F8'].value == 'Durée'


def est_sur_une_ligne(ws):
    return ws['D8'].value == 'Du jour' and ws['E8'].value == 'Au jour'


def remplissage_automatique(couleur):
    # Un remplissage conditionnel (DXF) utilise bgColor dans Excel,
    # contrairement au remplissage ordinaire des cellules (fgColor).
    return PatternFill(bgColor='FF' + couleur)


def assurer_couleurs_automatiques(ws):
    """Répare les couleurs des périodes déjà créées sans déplacer les saisies."""
    if not est_sur_une_ligne(ws):
        return False
    nombre = 0
    while ws.cell(8, 4+2*nombre).value == 'Du jour':
        nombre += 1
    lignes = [r for r in range(9, ws.max_row+1)
              if ws.cell(r, 1).value is not None
              and not str(ws.cell(r, 1).value).startswith('Source :')]
    if not lignes:
        return False
    premier = 4+2*nombre
    lettre = get_column_letter(premier)
    plage = f'{lettre}9:{get_column_letter(premier+139)}{max(lignes)}'
    modifie = False
    for n in range(nombre):
        d, f = get_column_letter(4+2*n), get_column_letter(5+2*n)
        formule = (f'AND(ISNUMBER(${d}9),ISNUMBER(${f}9),${d}9>=1,'
                   f'${f}9>${d}9-1,{lettre}$8>=${d}9,{lettre}$8<${f}9+1)')
        fill = remplissage_automatique(COULEURS[n % 4])
        regles = [regle for zone, liste in ws.conditional_formatting._cf_rules.items()
                  if str(zone.sqref) == plage for regle in liste
                  if regle.type == 'expression' and regle.formula == [formule]]
        if not regles:
            ws.conditional_formatting.add(plage, FormulaRule(formula=[formule], fill=fill))
            modifie = True
        for regle in regles:
            if regle.dxf is None or regle.dxf.fill != fill:
                regle.dxf = copy(regle.dxf) if regle.dxf else DifferentialStyle()
                regle.dxf.fill = copy(fill)
                modifie = True
    return modifie


def lire_plages(ws):
    nombre = 0
    while ws.cell(8, 4 + 2*nombre).value == 'Du jour':
        nombre += 1
    return {str(ws.cell(r, 1).value): [
        (ws.cell(r, c).coordinate, ws.cell(r, c).value)
        for c in range(4, 4 + 2*nombre)]
        for r in range(9, ws.max_row+1)
        if ws.cell(r, 1).value is not None
        and not str(ws.cell(r, 1).value).startswith('Source :')}


def restaurer_plages(ws, plages):
    lignes = {str(ws.cell(r, 1).value): r for r in range(9, ws.max_row+1)}
    if set(plages) - set(lignes):
        raise ValueError('Un poste du planning est absent du PR. Le planning est conservé.')
    for article, cellules in plages.items():
        for col, (origine, valeur) in enumerate(cellules, 4):
            cible = ws.cell(lignes[article], col)
            cible.value = (Translator(valeur, origin=origine).translate_formula(cible.coordinate)
                           if isinstance(valeur, str) and valeur.startswith('=') else valeur)


def centraliser_periodes(ws):
    """Reprend les anciens formats sans filtre ni lignes supplémentaires par poste."""
    if est_sur_une_ligne(ws):
        return False
    compact = est_compact(ws)
    nombre = 1
    if not compact:
        while ws.cell(8, 4+3*nombre).value == f'Durée {nombre+1}':
            nombre += 1
    ancien_calendrier = 8 if compact else 4+3*nombre
    jour = ws.cell(4, ancien_calendrier).value or 1
    postes = {}
    for r in range(9, ws.max_row+1):
        article = ws.cell(r, 1).value
        if article is None or str(article).startswith('Source :'):
            continue
        if compact:
            periode = ws.cell(r, 4).value
            if not isinstance(periode, int) or periode < 1:
                raise ValueError('Période non reconnue. Le planning est conservé.')
            periodes = [(periode, ws.cell(r, 5), ws.cell(r, 6))]
            nombre = max(nombre, periode)
        else:
            if ws.cell(r, 6).data_type != 'f':
                continue
            periodes = [(n+1, ws.cell(r, 5+3*n), ws.cell(r, 4+3*n)) for n in range(nombre)]
        cle = str(article)
        if cle not in postes:
            postes[cle] = (r, [ws.cell(r, c).value for c in range(1, 4)], {})
        for n, debut, duree in periodes:
            if n in postes[cle][2]:
                raise ValueError('Article et période en double. Le planning est conservé.')
            valeur = duree.value
            if valeur == f'=C{r}':
                valeur = ws.cell(r, 3).value
            if valeur is not None and not isinstance(valeur, (int, float)):
                raise ValueError('Durée avec formule personnalisée : le planning est conservé.')
            if debut.value is not None and not isinstance(debut.value, (int, float)):
                raise ValueError('Début avec formule personnalisée : le planning est conservé.')
            postes[cle][2][n] = (debut.value, valeur)
    if not postes:
        raise ValueError('Aucun poste reconnu dans le planning.')
    nombre = max(4, nombre)
    source = next((ws.cell(r, 1).value for r in range(9, ws.max_row+1)
                   if str(ws.cell(r, 1).value).startswith('Source :')), '')
    source_style = copy(ws.cell(ws.max_row, 1)._style)
    styles = [copy(ws.cell(9, c)._style) for c in (1, 2, 3)]
    saisie_style = copy(ws['E9']._style)
    entete_style = copy(ws['A8']._style)
    hauteurs = {a: ws.row_dimensions[p[0]].height for a, p in postes.items()}
    # Conserve les annotations et couleurs posées directement sur le calendrier.
    grille = {a: [copy(ws.cell(p[0], c)) for c in range(ancien_calendrier, ws.max_column+1)]
              for a, p in postes.items()}
    for plage in list(ws.merged_cells.ranges):
        if plage.min_row >= 7 or plage.min_col >= 7:
            ws.unmerge_cells(str(plage))
    ws.delete_rows(7, ws.max_row-6)
    for r in list(ws.row_dimensions):
        if r >= 7:
            del ws.row_dimensions[r]
    for c in list(ws.column_dimensions):
        if c not in ('A', 'B', 'C'):
            del ws.column_dimensions[c]
    for row in ws.iter_rows(min_row=1, max_row=6, min_col=7):
        for cell in row:
            cell.value = None
    ws.conditional_formatting._cf_rules.clear()
    ws.data_validations.dataValidation.clear()
    ws.auto_filter = AutoFilter()
    ws.sheet_properties.filterMode = False
    ws['A6'] = 'Sur la ligne du poste, saisissez les jours de début et de fin de chaque période.'
    for c, titre in enumerate(('Article', 'Poste', 'Jours PR'), 1):
        ws.cell(8, c, titre)._style = copy(entete_style)
    fin = 8 + len(postes)
    controles, fins = [], []
    for n in range(nombre):
        d, f = get_column_letter(4+2*n), get_column_letter(5+2*n)
        ws.merge_cells(f'{d}7:{f}7')
        ws[f'{d}7'] = f'Période {n+1}'
        for ref in (f'{d}7', f'{d}8', f'{f}8'):
            ws[ref]._style = copy(entete_style)
            ws[ref].fill = PatternFill('solid', fgColor=COULEURS[n % 4])
        ws[f'{d}8'], ws[f'{f}8'] = 'Du jour', 'Au jour'
        for c in (d, f):
            ws.column_dimensions[c].width = 11
            v = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1=0, allow_blank=True)
            v.showErrorMessage = True
            v.error = 'Saisissez un numéro de jour.'
            ws.add_data_validation(v)
            v.add(f'{c}9:{c}{fin}')
        # Une paire vide est facultative ; une seule borne renseignée reste incomplète.
        controles.append(f'SUMPRODUCT(--(({d}9:{d}{fin}="")<>({f}9:{f}{fin}="")))')
        controles.append(f'SUMPRODUCT(--({d}9:{d}{fin}<>""),--({d}9:{d}{fin}<1))')
        controles.append(f'SUMPRODUCT(--({f}9:{f}{fin}<>""),--({f}9:{f}{fin}<{d}9:{d}{fin}-1))')
        fins.append(f'{f}9:{f}{fin}')
    for r, (article, (_, valeurs, periodes)) in enumerate(postes.items(), 9):
        for c, valeur in enumerate(valeurs, 1):
            ws.cell(r, c, valeur)._style = copy(styles[c-1])
        ws.cell(r, 1).data_type = ws.cell(r, 2).data_type = 's'
        ws.row_dimensions[r].height = hauteurs[article]
        for n in range(nombre):
            d, f = ws.cell(r, 4+2*n), ws.cell(r, 5+2*n)
            debut, duree = periodes.get(n+1, (None, None))
            d.value = debut
            if duree is not None:
                f.value = (debut+duree-1 if debut is not None
                           else f'=IF({d.coordinate}="","",{d.coordinate}+{duree}-1)')
            for cell in (d, f):
                cell._style = copy(saisie_style)
                cell.number_format = '0.##'
                cell.fill = PatternFill('solid', fgColor='FFF2CC')
    ws['F4'] = f'=IF({"+".join(controles)}>0,"À compléter",MAX({",".join(fins)}))'
    ws.column_dimensions['F'].width = 15  # Le délai final partage cette colonne.
    premier = 4+2*nombre
    lettre = get_column_letter(premier)
    dernier = premier+139
    ws.merge_cells(start_row=3, start_column=premier, end_row=3, end_column=premier+6)
    ws.cell(3, premier, 'À PARTIR DU JOUR')._style = copy(ws['E3']._style)
    ws.merge_cells(start_row=4, start_column=premier, end_row=4, end_column=premier+2)
    ws.cell(4, premier, jour)._style = copy(saisie_style)
    v = DataValidation(type='whole', operator='greaterThanOrEqual', formula1=1)
    v.showErrorMessage = True
    ws.add_data_validation(v)
    v.add(f'{lettre}4')
    for c in range(premier, dernier+1):
        col = get_column_letter(c)
        ws.column_dimensions[col].width = 3.5
        ws.cell(8, c, f'=${lettre}$4' if c == premier else f'={get_column_letter(c-1)}8+1')._style = copy(entete_style)
        ws.cell(8, c).number_format = '0'
        for r, article in enumerate(postes, 9):
            cell = ws.cell(r, c)
            anciens = grille[article]
            if c-premier < len(anciens):
                ancien = anciens[c-premier]
                cell._style = copy(ancien._style)
                cell.value = ancien.value
                cell.comment = copy(ancien.comment)
            else:
                cell.border = copy(ws.cell(r, 4).border)
                cell.fill = PatternFill('solid', fgColor='F3F7FB' if r % 2 else 'FFFFFF')
    assurer_couleurs_automatiques(ws)
    if source:
        ws.merge_cells(start_row=fin+2, start_column=1, end_row=fin+2, end_column=3+2*nombre)
        ws.cell(fin+2, 1, source.replace('les durées retenues sont modifiables', 'les jours des périodes sont modifiables'))._style = source_style
        ws.row_dimensions[fin+2].height = 28
    ws.freeze_panes = 'D9'
    ws.print_area = f'A1:{get_column_letter(dernier)}{fin+2}'
    ws.print_title_cols = 'A:C'
    return True

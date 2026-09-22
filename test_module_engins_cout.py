import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile
from openpyxl import Workbook, load_workbook
from openpyxl.formula.translate import Translator
from documents_cout import actualiser_liens
from liens_horizon import analyser
from module_engins_cout import ajouter_module_engins, TITRE, TOTAL_INITIAL, FIN


class ModuleEnginsTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.p = Path(self.tmp.name) / 'Cout_securite.xlsx'
        w = Workbook()
        s = w.active
        s.title = 'Feuille1'
        s['K17'] = '=' + TOTAL_INITIAL
        for col, texte in zip('ABCDEFGHIJK', ('162 · Définition', 'Entreprise', 'location', 'jours', 'Pièce', 'M1', 'M2', 'M3', 'Heures', 'Prix u', 'Prix total')):
            s[f'{col}162'] = texte
        s['A163'] = 'FIN DE CHANTIER'
        s['J163'] = 'total rubrique'
        s['K163'] = '=SUM($K164:$K524)'
        s['A161'] = 'Containers = PR → pas ici.'
        for col in 'ABCDEFGHIJK':
            s[f'{col}164'] = 0
        s['A164'] = '164 · Démontage échafaudages'
        s['B164'] = 'Docs'
        s['E164'] = 2
        s['J164'] = 100
        s['K164'] = '=IF($J164="","",$J164*IF(ISNUMBER($E164),$E164,1))'
        s['A524'] = 'Note conservée'
        w.create_sheet('Feuille2')['K17'] = 71244.5
        w.save(self.p)
        with ZipFile(self.p, 'a') as z:
            z.writestr('xl/media/temoin.bin', b'image')

    def test_preserve_existant_et_formules_sans_double_compte(self):
        with ZipFile(self.p) as z:
            avant = {n: z.read(n) for n in z.namelist()}
        self.assertTrue(ajouter_module_engins(self.p))
        w = load_workbook(self.p)
        s = w.active
        self.assertEqual(s['A175'].value, TITRE)
        self.assertEqual(s['K17'].value, '=' + TOTAL_INITIAL[:-1] + ',K175)')
        self.assertEqual(s['K163'].value, f'=SUM($K164:$K173,$K{FIN + 1}:$K524)')
        self.assertEqual(s['K175'].value, f'=SUM($K176:$K{FIN})')
        for r in range(176, FIN + 1):
            self.assertEqual(s[f'K{r}'].value, Translator(s['K164'].value, origin='K164').translate_formula(f'K{r}'))
            for col in 'CDEFGHIJ':
                self.assertIsNone(s[f'{col}{r}'].value)
        self.assertEqual(s['E164'].value, 2)
        self.assertEqual(s['J164'].value, 100)
        self.assertEqual(s['A524'].value, 'Note conservée')
        with ZipFile(self.p) as z:
            for n, data in avant.items():
                if n not in ('xl/worksheets/sheet1.xml', 'xl/workbook.xml'):
                    self.assertEqual(z.read(n), data)
        self.assertFalse(ajouter_module_engins(self.p))

    def test_integration_dossiers_et_reouverture_sans_perdre_saisie(self):
        actualiser_liens(self.p)
        w = load_workbook(self.p)
        s = w.active
        cible = s['B178'].hyperlink.target
        self.assertEqual(analyser(cible)[1], '175')
        self.assertTrue((self.p.parent / 'Documents_installation/175_Mini_pelle').is_dir())
        s['J178'] = 123
        s['D178'] = 3
        w.save(self.p)
        actualiser_liens(self.p)
        self.assertEqual(load_workbook(self.p).active['J178'].value, 123)

    def test_ligne_occupee_preserve_le_fichier(self):
        w = load_workbook(self.p)
        w.active['A178'] = 'Poste personnel'
        w.save(self.p)
        avant = self.p.read_bytes()
        with self.assertRaises(ValueError):
            ajouter_module_engins(self.p)
        self.assertEqual(self.p.read_bytes(), avant)

    def test_extension_module_preserve_saisies(self):
        from unittest.mock import patch
        from module_engins_cout import POSTES
        with patch('module_engins_cout.FIN', 194), patch('module_engins_cout.POSTES', POSTES[:-1]):
            ajouter_module_engins(self.p)
        w = load_workbook(self.p)
        w.active['J178'] = 456
        w.save(self.p)
        self.assertTrue(ajouter_module_engins(self.p))
        w = load_workbook(self.p)
        self.assertEqual(w.active['J178'].value, 456)
        self.assertEqual(w.active['A195'].value, '192 · Grue à tour')
        self.assertEqual(w.active['K175'].value, '=SUM($K176:$K195)')
        self.assertFalse(ajouter_module_engins(self.p))

    def test_dossiers_bibliotheque_ajoutent_postes_excel_et_liens(self):
        bibliotheque = self.p.parent / 'Bibliotheque_Cout_securite'
        bibliotheque.mkdir()
        (bibliotheque / '193_Protection_des_accès_pour_engin').mkdir()
        w = load_workbook(self.p)
        w.active['A196'] = 'Ligne de clôture existante'
        w.active['K196'] = 42
        w.save(self.p)
        actualiser_liens(self.p, bibliotheque)
        s = load_workbook(self.p).active
        self.assertEqual(s['A196'].value, 'Ligne de clôture existante')
        self.assertEqual(s['K196'].value, 42)
        self.assertEqual(s['A525'].value, '193 · Protection des accès pour engin')
        self.assertEqual(s['K175'].value, '=SUM($K176:$K195,K525)')
        self.assertEqual(s['K163'].value, '=SUM($K164:$K173,$K196:$K524)')
        self.assertEqual(analyser(s['B525'].hyperlink.target)[1], '193')
        s['J525'] = 123
        s.parent.save(self.p)
        (bibliotheque / '194_Protection_complémentaire').mkdir()
        actualiser_liens(self.p, bibliotheque)
        s = load_workbook(self.p).active
        self.assertEqual(s['J525'].value, 123)
        self.assertEqual(s['A526'].value, '194 · Protection complémentaire')
        self.assertEqual(s['K175'].value, '=SUM($K176:$K195,K525,K526)')

    def test_poste_193_avec_total_excel_mis_en_forme_differemment(self):
        bibliotheque = self.p.parent / 'Bibliotheque_Cout_securite'
        bibliotheque.mkdir()
        (bibliotheque / '193_Protection_des_accès_pour_engin').mkdir()
        ajouter_module_engins(self.p)
        w = load_workbook(self.p)
        w.active['K17'] = '=' + TOTAL_INITIAL[:-1] + ', $K$175)'
        w.save(self.p)
        actualiser_liens(self.p, bibliotheque)
        s = load_workbook(self.p).active
        self.assertEqual(s['A525'].value, '193 · Protection des accès pour engin')
        self.assertEqual(s['K17'].value, '=' + TOTAL_INITIAL[:-1] + ', $K$175)')

    def test_libelle_193_existant_corrige_sans_perdre_les_valeurs(self):
        bibliotheque = self.p.parent / 'Bibliotheque_Cout_securite'
        bibliotheque.mkdir()
        (bibliotheque / '193_Protection_des_accès_pour_engin').mkdir()
        actualiser_liens(self.p, bibliotheque)
        w = load_workbook(self.p)
        w.active['A525'] = '193 · Ancien libellé'
        w.active['J525'] = 321
        w.save(self.p)
        actualiser_liens(self.p, bibliotheque)
        s = load_workbook(self.p).active
        self.assertEqual(s['A525'].value, '193 · Protection des accès pour engin')
        self.assertEqual(s['J525'].value, 321)
        self.assertEqual(analyser(s['B525'].hyperlink.target)[1], '193')

    def test_poste_193_existant_est_reintegre_au_sous_total(self):
        bibliotheque = self.p.parent / 'Bibliotheque_Cout_securite'
        bibliotheque.mkdir()
        (bibliotheque / '193_Protection_des_accès_pour_engin').mkdir()
        actualiser_liens(self.p, bibliotheque)
        w = load_workbook(self.p)
        w.active['J525'] = 321
        w.active['K175'] = '=SUM($K176:$K195)'
        w.save(self.p)
        actualiser_liens(self.p, bibliotheque)
        s = load_workbook(self.p).active
        self.assertEqual(s['A525'].value, '193 · Protection des accès pour engin')
        self.assertEqual(s['J525'].value, 321)
        self.assertEqual(s['K175'].value, '=SUM($K176:$K195,K525)')
        self.assertFalse(actualiser_liens(self.p, bibliotheque))

    def test_poste_193_repare_sans_imposer_formule_cloture_standard(self):
        bibliotheque = self.p.parent / 'Bibliotheque_Cout_securite'
        bibliotheque.mkdir()
        (bibliotheque / '193_Protection_des_accès_pour_engin').mkdir()
        actualiser_liens(self.p, bibliotheque)
        w = load_workbook(self.p)
        w.active['K163'] = '=SUM($K164:$K173,$K196:$K524,0)'
        w.active['K175'] = '=SUM($K176:$K195)'
        w.save(self.p)
        actualiser_liens(self.p, bibliotheque)
        s = load_workbook(self.p).active
        self.assertEqual(s['K163'].value, '=SUM($K164:$K173,$K196:$K524,0)')
        self.assertEqual(s['K175'].value, '=SUM($K176:$K195,K525)')

    def test_classeur_non_reconnu_signale_que_193_manque(self):
        bibliotheque = self.p.parent / 'Bibliotheque_Cout_securite'
        bibliotheque.mkdir()
        (bibliotheque / '193_Protection_des_accès_pour_engin').mkdir()
        w = Workbook()
        w.active['A1'] = 'Autre format'
        w.save(self.p)
        avant = self.p.read_bytes()
        with self.assertRaisesRegex(ValueError, '193.*Excel'):
            actualiser_liens(self.p, bibliotheque)
        self.assertEqual(self.p.read_bytes(), avant)

    def test_migration_numeros_preserve_photo_et_formules(self):
        from module_engins_cout import numeroter_articles_engins
        ajouter_module_engins(self.p)
        with ZipFile(self.p) as z:
            contenus = {n: z.read(n) for n in z.namelist()}
        n = 'xl/worksheets/sheet1.xml'
        texte = contenus[n].decode()
        for ligne in reversed(range(176, FIN + 1)):
            texte = texte.replace(f'>{ligne - 3:03d} · ', f'>{ligne:03d} · ')
        contenus[n] = texte.encode()
        with ZipFile(self.p, 'w') as z:
            for nom, data in contenus.items():
                z.writestr(nom, data)
        dossier = self.p.parent / 'Documents_installation/178_Mini_pelle'
        dossier.mkdir(parents=True)
        (dossier / 'photo.jpg').write_bytes(b'photo originale')
        self.assertTrue(numeroter_articles_engins(self.p))
        self.assertFalse(numeroter_articles_engins(self.p))
        actualiser_liens(self.p)
        w = load_workbook(self.p)
        self.assertEqual(w.active['A176'].value, '173 · Containers')
        self.assertEqual(w.active['A195'].value, '192 · Grue à tour')
        cible = w.active['B178'].hyperlink.target
        self.assertEqual((self.p.parent / 'Documents_installation/175_Mini_pelle' / 'photo.jpg').read_bytes(), b'photo originale')
        self.assertEqual(w.active['K175'].value, f'=SUM($K176:$K{FIN})')

    def test_migration_conserve_libelle_personnalise_du_poste_194(self):
        from module_engins_cout import numeroter_articles_engins
        ajouter_module_engins(self.p)
        w = load_workbook(self.p)
        w.active['A194'] = '194 · Transport spécial choisi pour ce chantier'
        w.save(self.p)
        dossier = self.p.parent / 'Documents_installation/194_Transport_special'
        dossier.mkdir(parents=True)
        (dossier / 'justificatif.pdf').write_bytes(b'document conserve')
        self.assertTrue(numeroter_articles_engins(self.p))
        self.assertEqual(load_workbook(self.p).active['A194'].value,
                         '191 · Transport spécial choisi pour ce chantier')
        self.assertEqual((self.p.parent / 'Documents_installation/191_Transport_special'
                          / 'justificatif.pdf').read_bytes(), b'document conserve')
        self.assertFalse(numeroter_articles_engins(self.p))

    def test_classeur_ouvert(self):
        self.p.with_name('~$' + self.p.name).touch()
        avant = self.p.read_bytes()
        with self.assertRaises(PermissionError):
            ajouter_module_engins(self.p)
        self.assertEqual(self.p.read_bytes(), avant)


if __name__ == '__main__':
    unittest.main()

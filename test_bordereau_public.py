import hashlib
import tempfile
import unittest
import zipfile
from pathlib import Path
from openpyxl import Workbook, load_workbook
from bordereau_public import preparer, exporter_prix, feuille, habiller_avancement
from main import _corriger_moins_etat_public, _selectionner_etat, _recalculer_etat, _reporter_quantites_cloture

class BordereauPublic(unittest.TestCase):
    def creer_indice_c_et_copie(self, dossier):
        original, copie = dossier / 'indice_C.xlsx', dossier / 'avancement.xlsx'
        w = Workbook(); s = w.active; s.title = 'LOT 1 - IND. C'
        for adresse, valeur in {'A1': 'N°', 'B1': 'DÉSIGNATION', 'C1': 'MESURAGE',
                                'D1': 'UNITÉ', 'J1': 'QUANTITÉ', 'K1': 'PRIX UNITAIRES',
                                'L1': 'PRIX', 'M1': 'PRIX OPTIONS'}.items():
            s[adresse] = valeur
        for r, code, libelle, quantite in ((13, '01', 'Poste', 12), (18, '02', 'OPTION 1 : Poste', 3)):
            for col, valeur in {'A': code, 'B': libelle, 'C': 'QP', 'D': 'M²',
                                'J': quantite, 'K': 0, 'L': f'=J{r}*K{r}',
                                'M': f'=L{r}' if r == 18 else None}.items():
                s[f'{col}{r}'] = valeur
        s['K13'] = None  # Cellule vide mais présente avec son style, XML autofermant.
        s['K13'].number_format = '#,##0.00'
        s['B20'] = 'Note officielle conservée'
        w.save(original); w.close()
        w = Workbook(); s = w.active
        s['B1'], s['K1'], s['L1'] = 'N°', 'QUANTITÉ', 'PRIX UNITAIRES'
        # Ordre et lignes différents ; les quantités officielles font foi dans l'export.
        for r, code, libelle, prix in ((4, '02', 'OPTION 1 : Poste', 0), (9, '01', 'Poste', 125.75)):
            for col, valeur in {'B': code, 'C': libelle, 'D': 'QP', 'E': 'M²', 'K': 999, 'L': prix}.items():
                s[f'{col}{r}'] = valeur
        w.save(copie); w.close()
        return original, copie

    def test_export_indice_c_par_article_et_preservation_du_document_officiel(self):
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            original, copie = self.creer_indice_c_et_copie(d)
            avant = original.read_bytes()
            sortie = d / 'export.xlsx'
            self.assertEqual(exporter_prix(original, copie, sortie, hashlib.sha256(avant).hexdigest()), 2)
            self.assertEqual(original.read_bytes(), avant)
            a, b = load_workbook(original), load_workbook(sortie)
            try:
                s = b.active
                self.assertEqual(s['K13'].value, 125.75)
                self.assertEqual(s['K18'].value, 0)
                self.assertEqual(s['J13'].value, 12)
                self.assertEqual(s['M18'].value, '=L18')
                for row in a.active:
                    for cellule in row:
                        cible = s[cellule.coordinate]
                        if cellule.coordinate not in {'K13', 'K18'}:
                            self.assertEqual(cellule.value, cible.value)
                        self.assertEqual(cellule._style, cible._style)
            finally:
                a.close(); b.close()
            with zipfile.ZipFile(original) as a, zipfile.ZipFile(sortie) as b:
                self.assertEqual(a.namelist(), b.namelist())
                for nom in a.namelist():
                    if nom != 'xl/worksheets/sheet1.xml':
                        self.assertEqual(a.read(nom), b.read(nom))

    def test_indice_c_refuse_les_correspondances_ou_prix_invalides(self):
        for cellule, valeur in (('B4', '01'), ('B4', 'inconnu'), ('E9', 'M³'),
                                ('C9', 'Autre poste'), ('L9', None), ('L9', '=1+2'), ('L9', True)):
            with self.subTest(cellule=cellule, valeur=valeur), tempfile.TemporaryDirectory() as tmp:
                d = Path(tmp)
                original, copie = self.creer_indice_c_et_copie(d)
                w = load_workbook(copie); w.active[cellule] = valeur; w.save(copie); w.close()
                avant = original.read_bytes()
                with self.assertRaises(ValueError):
                    exporter_prix(original, copie, d / 'export.xlsx', hashlib.sha256(avant).hexdigest())
                self.assertFalse((d / 'export.xlsx').exists())
                self.assertEqual(original.read_bytes(), avant)

    def test_indice_c_ne_remplace_pas_une_formule_de_prix(self):
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            original, copie = self.creer_indice_c_et_copie(d)
            w = load_workbook(original); w.active['K13'] = '=1+2'; w.save(original); w.close()
            with self.assertRaisesRegex(ValueError, 'Formule protégée en K13'):
                exporter_prix(original, copie, d / 'export.xlsx', hashlib.sha256(original.read_bytes()).hexdigest())
            self.assertFalse((d / 'export.xlsx').exists())

    def test_ancienne_synthese_publique_retire_moins_sans_effacer_saisie(self):
        ws = Workbook().active
        ws["P20"] = "Total des avenants cumulé"
        ws["T20"] = "=SUM('[1]Feuil1'!$D$3:$D$60)-SUM('[1]Feuil1'!$E$3:$E$60)"
        ws["P18"] = "Total état cumulé hors TVA"
        ws["U19"] = '="Montant de l’avenant "&Q9'
        ws["W19"] = ("=N(INDEX('[1]Feuil1'!$D$3:$D$60,1+3*($Q$9-1)))"
                     "-N(INDEX('[1]Feuil1'!$E$3:$E$60,1+3*($Q$9-1)))")
        ws["W20"] = "=SUM(W18:W19)"
        ws["P21"] = "Avenants cumulés en moins"
        ws["T21"] = 0
        _corriger_moins_etat_public(ws)
        self.assertEqual(ws["T20"].value, "=SUM('[1]Feuil1'!$D$3:$D$60)")
        self.assertIsNone(ws["U19"].value)
        self.assertIsNone(ws["W19"].value)
        self.assertEqual(ws["W20"].value, "=W18")
        self.assertIsNone(ws["P21"].value)
        self.assertIsNone(ws["T21"].value)
        ws["P21"], ws["T21"] = "Avenants cumulés en moins", 200
        with self.assertRaisesRegex(ValueError, "Vérifiez"):
            _corriger_moins_etat_public(ws)
        self.assertEqual(ws["T21"].value, 200)

    def test_circuit_prix_et_avancement(self):
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            original = d / 'original.xlsx'
            w = Workbook(); s = w.active; s.title = 'MARIEMONT'
            for cell, v in {'B1':'N°', 'K1':'QUANTITÉ', 'L1':'PRIX UNITAIRES',
                            'B13':'02.21.4a.01', 'C13':'Essais', 'D13':'QP', 'E13':'PCE',
                            'K13':3, 'L13':0, 'M13':'=L13*K13', 'O13':'Note officielle'}.items():
                s[cell] = v
            w.save(original); w.close()
            avant = original.read_bytes()
            copie = preparer(original, d)
            self.assertEqual(_selectionner_etat(d), copie)
            w = load_workbook(copie); s = feuille(w)
            self.assertEqual(s['P16'].value, 'Total état cumulé hors TVA')
            self.assertEqual(s['T16'].value, '=SUM(T13)')
            self.assertIsNone(s['U17'].value)
            self.assertIsNone(s['W17'].value)
            self.assertEqual(s['W18'].value, '=W16')
            self.assertEqual(s['W19'].value, '=W18*$W$9')
            self.assertEqual(s['W21'].value, '=SUM(W18:W20)')
            self.assertEqual(s['T22'].value, '=T16')  # Le mois n'est pas ajouté deux fois.
            self.assertNotIn('$E$3', s['T20'].value)  # Les moins restent hors de l'état exécuté.
            self.assertIsNone(s['P27'].value)  # Aucune ligne « avenants en moins » dans la synthèse.
            self.assertEqual(s['U21'].fill.fgColor.rgb[-6:], 'FFFF00')
            self.assertEqual([l.file_link.Target for l in w._external_links],
                             ['Avenants.xlsx', 'Revision_global.xlsx'])
            self.assertIn('$Q$9', s['W20'].value)
            self.assertNotIn('676', s['W21'].value)
            s['L13'] = 12.5; s['P13'] = 1; s['Q13'] = 1
            s['Q9'] = 3
            habiller_avancement(w)
            self.assertEqual(s['Q9'].value, 3)
            self.assertEqual(s['L13'].value, 12.5)
            self.assertEqual(s['P13'].value, 1)
            self.assertEqual(s['Q13'].value, 1)
            _recalculer_etat(s)
            self.assertEqual(s['T13'].value, 25)
            self.assertEqual(s['W13'].value, 12.5)
            _reporter_quantites_cloture(s)
            self.assertEqual(s['P13'].value, 2)
            self.assertEqual(s['Q13'].value, 0)
            w.save(copie); w.close()
            export = d / 'export.xlsx'
            self.assertEqual(exporter_prix(original, copie, export, hashlib.sha256(avant).hexdigest()), 1)
            self.assertEqual(original.read_bytes(), avant)
            with zipfile.ZipFile(original) as a, zipfile.ZipFile(export) as b:
                self.assertEqual(a.namelist(), b.namelist())
                for name in a.namelist():
                    if name != 'xl/worksheets/sheet1.xml':
                        self.assertEqual(a.read(name), b.read(name))
            w = load_workbook(export)
            self.assertEqual(w.active['L13'].value, 12.5)
            self.assertEqual(w.active['M13'].value, '=L13*K13')
            self.assertEqual(w.active['O13'].value, 'Note officielle')
            w.close()
            with self.assertRaises(FileExistsError):
                preparer(original, d)
            copie.unlink()
            with self.assertRaises(FileNotFoundError):
                _selectionner_etat(d)

    def test_refus_original_modifie(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / 'original.xlsx'; p.write_bytes(b'changed')
            with self.assertRaisesRegex(ValueError, 'original a changé'):
                exporter_prix(p, p, Path(tmp) / 'export.xlsx', 'ancienne empreinte')

if __name__ == '__main__':
    unittest.main()

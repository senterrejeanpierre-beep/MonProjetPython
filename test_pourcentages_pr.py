import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from pourcentages_pr import NS, lier_pourcentages_pr, lire_pourcentages


class PourcentagesPRTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name) / 'chantier' / 'data' / 'prix_de_revient.xlsx'
        self.path.parent.mkdir(parents=True)
        w = Workbook()
        s = w.active
        s.title = 'Chiffrage'
        # Les récapitulatifs ne sont pas supposés être espacés de 42 lignes.
        for ligne, installation, marge in ((28, 0, 0.05), (77, 0.12, 0.1)):
            s.cell(ligne, 5, 'IC %')
            s.cell(ligne + 1, 5, 'Marge vente %')
            s.cell(ligne, 6, installation).number_format = '0.00%'
            s.cell(ligne + 1, 6, marge).number_format = '0.00%'
            s.cell(ligne, 3, 100)
            s.cell(ligne, 10, f'=C{ligne}*F{ligne}')
            s.cell(ligne + 1, 10, f'=(C{ligne}+J{ligne})*F{ligne + 1}')
        s['A1'].number_format = '0.00'  # cellule vide autofermante avant une formule
        s['F31'] = 0.21
        s['J3'] = 0.03
        w.create_sheet('LIEN_BORDEREAU')['A1'] = '=Chiffrage!J28'
        w.calculation.iterate = True
        w.calculation.iterateCount = 20
        w.save(self.path)
        w.close()
        with ZipFile(self.path) as z:
            items = {n: z.read(n) for n in z.namelist()}
        for n in list(items):
            if n.startswith('xl/worksheets/'):
                items[n] = items[n].replace(b'</f><v></v>', b'</f><v>999</v>')
        items['xl/media/objet_existant.bin'] = b'objet ou image a conserver'
        with ZipFile(self.path, 'w') as z:
            for n, data in items.items():
                z.writestr(n, data)
        self.avant = self.path.read_bytes()
        self.backup = Path(self.tmp.name) / 'sauvegarde.xlsx'

    def sauvegarder(self, path):
        shutil.copy2(path, self.backup)

    def test_liens_premier_bloc_preservation_et_recalcul(self):
        self.assertEqual(lier_pourcentages_pr(self.path, self.sauvegarder), 4)
        self.assertEqual(self.backup.read_bytes(), self.avant)
        w = load_workbook(self.path)
        s = w['Chiffrage']
        self.assertEqual((s['F28'].value, s['F29'].value), (0, .05))
        self.assertFalse(s['F28'].protection.locked)
        self.assertFalse(s['F29'].protection.locked)
        self.assertEqual((s['F77'].value, s['F78'].value), ('=$F$28', '=$F$29'))
        self.assertEqual(s['F77'].number_format, '0.00%')
        self.assertEqual((s['F31'].value, s['J3'].value), (.21, .03))
        self.assertEqual(s['J78'].value, '=(C77+J77)*F78')
        self.assertTrue(w.calculation.fullCalcOnLoad)
        self.assertEqual(w.calculation.calcMode, 'auto')
        self.assertTrue(w.calculation.iterate)
        self.assertEqual(w.calculation.iterateCount, 20)
        w.close()
        w = load_workbook(self.path, data_only=True)
        self.assertEqual((w['Chiffrage']['F77'].value, w['Chiffrage']['F78'].value), (0, .05))
        self.assertIsNone(w['Chiffrage']['J78'].value)
        w.close()
        with ZipFile(self.backup) as old, ZipFile(self.path) as new:
            self.assertEqual(set(old.namelist()), set(new.namelist()))
            for name in old.namelist():
                if name not in ('xl/workbook.xml', 'xl/styles.xml') and not name.startswith('xl/worksheets/'):
                    self.assertEqual(old.read(name), new.read(name), name)

    def test_seconde_ouverture_ne_reecrit_pas(self):
        lier_pourcentages_pr(self.path, self.sauvegarder)
        avant = self.path.read_bytes()
        backup = Mock()
        self.assertEqual(lier_pourcentages_pr(self.path, backup), 0)
        backup.assert_not_called()
        self.assertEqual(self.path.read_bytes(), avant)

    def test_verrou_excel_pas_de_remplacement(self):
        self.path.with_name('~$' + self.path.name).touch()
        with self.assertRaisesRegex(ValueError, 'fermez'):
            lier_pourcentages_pr(self.path, self.sauvegarder)
        self.assertEqual(self.path.read_bytes(), self.avant)

    def test_echec_sauvegarde_et_modification_concurrente(self):
        def modifier(_path):
            self.path.write_bytes(self.avant + b'concurrent')
        for callback in (Mock(side_effect=OSError('sauvegarde indisponible')), modifier):
            with self.subTest(callback=callback), self.assertRaises((OSError, ValueError)):
                lier_pourcentages_pr(self.path, callback)
            self.assertTrue(self.path.read_bytes().startswith(self.avant))
            self.assertFalse(list(self.path.parent.glob('.pourcentages_pr_*')))
        self.assertEqual(self.path.read_bytes(), self.avant + b'concurrent')

    def test_premier_taux_invalide_et_bloc_incomplet_refuses(self):
        for ref, valeur in [('F28', '=0.05'), ('F28', None), ('E29', 'Autre calcul')]:
            self.path.write_bytes(self.avant)
            w = load_workbook(self.path)
            w['Chiffrage'][ref] = valeur
            w.save(self.path)
            w.close()
            with self.assertRaises(ValueError):
                lier_pourcentages_pr(self.path)

    def test_deplacement_et_separation_des_chantiers(self):
        autre = self.path.parents[2] / 'Autre chantier' / 'data' / self.path.name
        autre.parent.mkdir(parents=True)
        shutil.copy2(self.path, autre)
        deplace = self.path.parents[2] / 'Chantier déplacé'
        self.path.parents[1].rename(deplace)
        cible = deplace / 'data' / self.path.name
        lier_pourcentages_pr(cible, self.sauvegarder)
        self.assertEqual(autre.read_bytes(), self.avant)
        self.assertEqual(lire_pourcentages(cible)[1]['valeurs'], ('=$F$28', '=$F$29'))

    def test_creation_future_et_ouverture_utilisent_les_liens(self):
        from main import _creer_documents_chantier, open_pr
        base = self.path.parents[2] / 'Futurs'
        modele = base / 'Modele'
        modele.mkdir(parents=True)
        shutil.copy2(self.path, modele / 'prix_de_revient model.xlsx')
        _creer_documents_chantier(base, 'Nouveau', {})
        nouveau = base / 'Nouveau'
        self.assertEqual(lire_pourcentages(nouveau / 'data/prix_de_revient.xlsx')[1]['valeurs'],
                         ('=$F$28', '=$F$29'))
        from unittest.mock import patch
        with patch('main.ouvrir_chemin'), patch('main._backup_excel_before_write', self.sauvegarder):
            open_pr(self.path.parents[1])
        self.assertEqual(lire_pourcentages(self.path)[1]['valeurs'], ('=$F$28', '=$F$29'))


if __name__ == '__main__':
    unittest.main()

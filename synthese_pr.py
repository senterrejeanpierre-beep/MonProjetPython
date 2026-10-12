"""Synthèse en lecture seule des récapitulatifs enregistrés dans le PR."""
from pathlib import Path
import math
import re
import unicodedata
from xml.sax.saxutils import escape
from openpyxl import load_workbook


def normaliser(valeur):
    return ''.join(c for c in unicodedata.normalize('NFKD', str(valeur or '').casefold())
                   if not unicodedata.combining(c))


LIBELLES = {'matiere': 'Matière / matériaux (rubrique PR)', 'mo': 'Main-d’œuvre',
            'revient': 'Prix de revient total', 'ic': 'IC', 'marge': 'Marge de vente',
            'vente': 'Prix de vente HT', 'heures': 'Heures de main-d’œuvre', 'jours': 'Jours prévus au PR'}


def est_option(designation):
    return bool(re.match(r'^\s*options?(?=\b|\d)', normaliser(designation)))


def lire_synthese_pr(chemin):
    wb = load_workbook(chemin, read_only=True, data_only=True)
    try:
        if 'Chiffrage' not in wb:
            raise ValueError('Feuille Chiffrage absente du PR.')
        lignes = list(wb['Chiffrage'].iter_rows(max_col=16, values_only=True))
    finally:
        wb.close()
    # Distinguer une saisie laissée vide d'une formule dont Excel n'a pas
    # enregistré le résultat : seule la première vaut zéro pour la simulation.
    wb_formules = load_workbook(chemin, read_only=True, data_only=False)
    try:
        formules = {(cell.row, cell.column): cell.value
                    for row in wb_formules['Chiffrage'].iter_rows(max_col=16)
                    for cell in row if cell.data_type == 'f'}
    finally:
        wb_formules.close()
    postes, alertes = [], []
    taux_horaires, taux_marges = {False: [], True: []}, {False: [], True: []}
    def nombre(ligne, col, article):
        valeur = lignes[ligne - 1][col - 1] if ligne <= len(lignes) else None
        if isinstance(valeur, (int, float)) and not isinstance(valeur, bool) and math.isfinite(valeur):
            return float(valeur)
        if valeur in (None, '') and (ligne, col) not in formules:
            return 0.0
        from openpyxl.utils import get_column_letter
        alertes.append(f'{article} : {get_column_letter(col)}{ligne} absent ou sans résultat numérique enregistré.')
        return None
    for i, ligne in enumerate(lignes):
        if not normaliser(ligne[1]).startswith('article'):
            continue
        debut = i + 2
        if debut > len(lignes) or lignes[debut - 1][1] in (None, ''):
            continue
        article = str(lignes[debut - 1][1])
        if debut + 29 > len(lignes):
            raise ValueError(f'Bloc incomplet pour {article}.')
        for offset, col, attendu in ((25, 2, 'pr matiere'), (26, 2, 'pr mo'),
                                     (27, 2, 'pr total'), (27, 5, 'pv ht')):
            if not normaliser(lignes[debut + offset - 1][col - 1]).startswith(attendu):
                raise ValueError(f'Récapitulatif PR non reconnu pour {article}, ligne {debut + offset}.')
        quantite = nombre(debut, 5, article)
        poste = {'article': article, 'designation': str(lignes[debut - 1][2] or '')}
        option = poste['option'] = est_option(poste['designation'])
        for cle, offset, col in (('matiere',25,3), ('mo',26,3), ('revient',27,3),
                                  ('ic',25,10), ('marge',26,10), ('vente',27,6)):
            valeur = nombre(debut + offset, col, article)
            poste[cle] = valeur * quantite if valeur is not None and quantite is not None else None
        # I contient les heures de main-d'œuvre déjà calculées dans le PR.
        heures = [nombre(r, 9, article) for r in range(debut+13, debut+23)
                  if lignes[r-1][3] not in (None, '') or lignes[r-1][8] is not None]
        poste['heures'] = sum(heures) if all(v is not None for v in heures) else None
        for r in range(debut+13, debut+23):
            if lignes[r-1][3] not in (None, '') or lignes[r-1][8] is not None:
                poids = nombre(r, 9, article)
                taux_horaires[option].append((nombre(r, 10, article) if poids != 0 else 0, poids))
        base_marge = (poste['revient'] + poste['ic']
                      if poste['revient'] is not None and poste['ic'] is not None else None)
        taux_marges[option].append((nombre(debut+26, 6, article) if base_marge != 0 else 0, base_marge))
        # Les blocs récents portent un total explicite ; les anciens ont des jours par ligne MO.
        totaux = [r for r in range(debut+13, debut+23)
                  if normaliser(lignes[r-1][5]).startswith('total jours')]
        if totaux:
            # Total du poste déjà calculé au-dessus des jours : ne pas le
            # remplacer par les heures détaillées, parfois unitaires, de I.
            poste['heures'] = nombre(totaux[0] - 1, 7, article)
            total = totaux[0]
            mo = total - 1
            # Ancien total jours = heures prévues / nb personnes / 8.
            # Une équipe encore vide produit #DIV/0! dans Excel, mais n'a
            # aucun jour à ajouter à la simulation si ses heures sont nulles.
            equipe_vide = (lignes[total-1][6] == '#DIV/0!'
                           and formules.get((total, 7), '').replace(' ', '').upper()
                               == f'=I{mo}/C{mo}/8'
                           and lignes[mo-1][2] in (None, '', 0)
                           and lignes[mo-1][8] == 0)
            poste['jours'] = 0.0 if equipe_vide else nombre(total, 7, article)
        else:
            jours = [nombre(r,7,article) for r in range(debut+13,debut+23)
                     if lignes[r-1][3] not in (None,'') or lignes[r-1][6] is not None]
            poste['jours'] = sum(jours) if all(v is not None for v in jours) else None
        postes.append(poste)
    if not postes:
        raise ValueError('Aucun poste reconnu dans le PR.')
    def moyenne(taux):
        if any(valeur is None or poids is None for valeur, poids in taux):
            return None
        poids_total = sum(poids for _, poids in taux)
        return sum(valeur * poids for valeur, poids in taux) / poids_total if poids_total else 0.0
    def bilan(option):
        selection = [p for p in postes if p['option'] == option]
        return {'postes': selection,
                'totaux': {cle: sum(p[cle] for p in selection)
                           if all(p[cle] is not None for p in selection) else None for cle in LIBELLES},
                'taux_horaire': moyenne(taux_horaires[option]),
                'taux_marge': moyenne(taux_marges[option])}
    base = bilan(False)
    # Le planning conserve tous ses postes et leurs saisies. Seule la synthèse
    # distingue les totaux de base des options.
    return {**base, 'postes': postes, 'options': bilan(True), 'alertes': alertes}


def generer_pdf_synthese(resultat, destination, nom, logo_path=None):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    styles = getSampleStyleSheet()
    styles['Title'].textColor = colors.HexColor('#163A59')
    styles['Title'].fontName = 'Helvetica-Bold'
    options = resultat.get('options', {})
    avec_options = bool(options.get('postes'))
    def texte(s): return Paragraph(escape(str(s)), styles['BodyText'])
    def valeur(v, euros=True):
        return 'Non disponible' if v is None else f'{v:,.2f}'.replace(',', ' ').replace('.', ',') + (' €' if euros else '')
    def premiere_page(canvas, doc):
        from datetime import datetime
        from reportlab.lib.utils import ImageReader
        bleu = colors.HexColor('#163A59')
        def rect(x,y,w,h,couleur):
            canvas.setFillColor(colors.HexColor(couleur));canvas.setStrokeColor(colors.HexColor('#C6D5E3'))
            canvas.rect(x*mm,y*mm,w*mm,h*mm,fill=1,stroke=1)
        def label(x,y,contenu,taille=9,couleur=bleu,gras=False,droite=False):
            canvas.setFillColor(couleur);canvas.setFont('Helvetica-Bold' if gras else 'Helvetica',taille)
            (canvas.drawRightString if droite else canvas.drawString)(x*mm,y*mm,str(contenu))
        def bandeau(x,y,w,titre,couleur):
            rect(x,y,w,6,couleur);label(x+1.5,y+1.6,titre,11,colors.white,True)
        def ligne(x,y,w,titre,montant,fond='#FFFFFF'):
            rect(x,y,w,6,fond);label(x+1.5,y+1.7,titre,9)
            label(x+w-1.5,y+1.7,montant,9,gras=True,droite=True)
        t = resultat['totaux']
        horaire = valeur(resultat['taux_horaire']) + (' / h' if resultat['taux_horaire'] is not None else '')
        marge = (valeur(resultat['taux_marge']*100,False)+' %'
                 if resultat['taux_marge'] is not None else 'Non disponible')
        rect(16,190,265,14,'#F3F7FB')
        if logo_path is not None and Path(logo_path).is_file():
            canvas.drawImage(ImageReader(str(logo_path)),17*mm,191*mm,width=28*mm,height=12*mm,
                             preserveAspectRatio=True,anchor='c',mask='auto')
        label(91,194,'SYNTHÈSE DU PR',18,gras=True)
        label(178,195,'CHANTIER',8,gras=True)
        canvas.saveState()
        chantier = Paragraph(escape(str(nom)), styles['BodyText'])
        chantier.wrap(45*mm,12*mm);chantier.drawOn(canvas,198*mm,193*mm)
        canvas.restoreState()
        label(246,195,'DATE',8,gras=True);label(258,195,datetime.now().strftime('%d/%m/%Y'),8)
        x,w=42,131
        bandeau(x,182,w,'COÛTS DU PR — HORS TVA' + (' ET OPTIONS' if avec_options else ''),'#1F4E79')
        for y,cle in ((176,'matiere'),(170,'mo')):
            ligne(x,y,w,LIBELLES[cle],valeur(t[cle]))
        ligne(x,164,w,'Prix de revient total',valeur(t['revient']),'#FFF2CC')
        bandeau(x,157,w,'MAIN-D’ŒUVRE ET JOURS PR','#4F81BD')
        ligne(x,151,w,'Heures de main-d’œuvre',valeur(t['heures'],False)+' h' if t['heures'] is not None else 'Non disponible')
        ligne(x,145,w,'Taux horaire moyen appliqué HT',horaire,'#F3F7FB')
        ligne(x,139,w,'Jours prévus au PR (sans chevauchement)',valeur(t['jours'],False),'#FFF2CC')
        bandeau(x,132,w,'MARGE ET PRIX DE VENTE','#C0504D')
        ligne(x,126,w,'IC',valeur(t['ic']))
        ligne(x,120,w,'Marge de vente',valeur(t['marge']),'#F3F7FB')
        ligne(x,114,w,'Marge moyenne appliquée',marge,'#FFF2CC')
        ligne(x,108,w,'Prix de vente HT',valeur(t['vente']),'#D9EAD3')
        if avec_options:
            o = options['totaux']
            bandeau(x,101,w,'OPTIONS — HORS TOTAUX DE BASE','#163A59')
            for y, cle in ((95,'matiere'),(89,'mo'),(83,'revient'),(77,'vente'),
                           (71,'heures'),(65,'jours')):
                ligne(x,y,w,LIBELLES[cle],valeur(o[cle],cle not in ('heures','jours')))
        x,w=176,73
        bandeau(x,182,w,'RÉSUMÉ VISUEL','#163A59')
        rect(x,112,w,70,'#FFFFFF')
        valeurs=[t['matiere'],t['mo'],t['marge']]
        pied_barres = 134 if avec_options else 122
        hauteur_barres = 34 if avec_options else 46
        if all(v is not None and v >= 0 for v in valeurs):
            maximum=max(max(valeurs),1)
            for j in range(5):
                y=pied_barres+j*(hauteur_barres/4 if avec_options else 12)
                canvas.setStrokeColor(colors.HexColor('#E0EAF3'))
                canvas.line(179*mm,y*mm,246*mm,y*mm)
            for i,(v,nom_barre,couleur) in enumerate(zip(valeurs,['Matière','Main-d’œuvre','Marge'],['#00547D','#4F81BD','#70AD47'])):
                bx=181+i*22;hauteur=v/maximum*hauteur_barres
                rect(bx,pied_barres,13,hauteur,couleur)
                label(bx+16,pied_barres+2+hauteur,valeur(v),6.7,gras=True,droite=True)
                label(bx-1,pied_barres-5,nom_barre,7)
        else:
            label(179,148,'Valeurs incomplètes : voir le détail.',8)
        if avec_options:
            vente_base, vente_options = t['vente'], options['totaux']['vente']
            label(179,124,'RÉPARTITION DE LA VENTE HT',7,gras=True)
            if (vente_base is not None and vente_options is not None
                    and vente_base >= 0 and vente_options >= 0 and vente_base+vente_options > 0):
                part_options = vente_options / (vente_base+vente_options)
                part_base = 1-part_options
                bleu_base, couleur_options = '#00547D', '#D6A15D'
                rect(179,118.5,2.5,2.5,bleu_base)
                label(183,119,'Base : '+f'{100*part_base:.1f}'.replace('.',',')+' %',7)
                rect(215,118.5,2.5,2.5,couleur_options)
                label(219,119,'Options : '+f'{100*part_options:.1f}'.replace('.',',')+' %',7)
                if part_base > 0:
                    rect(179,114,67*part_base,3.5,bleu_base)
                if part_options > 0:
                    rect(179+67*part_base,114,67*part_options,3.5,couleur_options)
            else:
                label(179,116,'Répartition HT non disponible.',7)
        bandeau(x,105,w,'INDICATEURS CLÉS','#366092')
        for y,nom_indicateur,v in ((99,'PR total HT',valeur(t['revient'])),
                                  (93,'Vente HT',valeur(t['vente'])),
                                  (87,'Heures MO',valeur(t['heures'],False)),
                                  (81,'Taux horaire HT',horaire),
                                  (75,'Marge moyenne',marge),
                                  (69,'Jours PR',valeur(t['jours'],False))):
            ligne(x,y,w,nom_indicateur,v,'#D9EAD3' if y in (93,75) else '#F3F7FB')
        vente_options = options['totaux']['vente'] if avec_options else 0.0
        vente_generale = (t['vente'] + vente_options
                         if t['vente'] is not None and vente_options is not None else None)
        ligne(42,58,207,'Total général des prix de vente HT (options comprises)',
              valeur(vente_generale),'#D9EAD3')
        rect(42,42,207,15,'#FFF2CC')
        label(44,51,'Lecture du PR enregistré — tous les montants sont hors TVA.',10,gras=True)
        label(44,45,'Aucune modification des données ou des formules du PR depuis cette synthèse.',8)
        notes = ('Taux horaires pondérés par les heures ; marges pondérées par le PR + IC. '
                 'Délai calculé sans chevauchement des travaux : les jours des postes du PR sont additionnés, '
                 'sans tenir compte des travaux réalisables en parallèle ni convertir le total en durée calendaire. '
                 'Les travaux simultanés sont à organiser dans le planning de soumission. '
                 'La matière inclut les sous-traitances et locations chiffrées dans cette rubrique du PR.')
        if avec_options:
            notes = notes.replace('les jours des postes du PR sont additionnés',
                                  'seuls les jours des postes hors options sont additionnés')
        if resultat['alertes']:
            notes = 'SYNTHÈSE INCOMPLÈTE : cellules à vérifier dans les pages suivantes. ' + notes
        paragraphe=Paragraph(escape(notes),styles['BodyText'])
        _,h=paragraphe.wrap(207*mm,35*mm);paragraphe.drawOn(canvas,42*mm,38*mm-h)
    elements = [Spacer(1,1), PageBreak(), Paragraph('DÉTAIL PAR POSTE — HORS TVA', styles['Title'])]
    if resultat['alertes']:
        elements += [texte('SYNTHÈSE INCOMPLÈTE : vérifier les cellules indiquées, puis enregistrer le PR dans Excel.'),
                     *[texte(a) for a in resultat['alertes']],Spacer(1,4*mm)]
    entetes = ['Article / poste','Matière','MO','PR total','IC','Marge','Vente HT','Heures MO','Jours PR']
    def tableau(postes, totaux=None):
        rows = [entetes]
        for p in postes:
            rows.append([texte(p['article']+' — '+p['designation'])] +
                        [valeur(p[k],False) for k in LIBELLES])
        if totaux is not None:
            rows.append([texte('Total des options')] + [valeur(totaux[k],False) for k in LIBELLES])
        table = Table(rows, colWidths=[65*mm]+[25*mm]*8, repeatRows=1)
        table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#163A59')),
                               ('TEXTCOLOR',(0,0),(-1,0),colors.white),('FONTSIZE',(0,0),(-1,-1),8),
                               ('VALIGN',(0,0),(-1,-1),'TOP'),('ALIGN',(1,1),(-1,-1),'RIGHT'),
                               ('GRID',(0,0),(-1,-1),0.3,colors.lightgrey)]))
        return table
    elements.append(tableau([p for p in resultat['postes'] if not p.get('option',False)]))
    if avec_options:
        elements.extend([Spacer(1,4*mm), Paragraph('OPTIONS — HORS TOTAUX DE BASE',styles['Title']),
                         tableau(options['postes'],options['totaux'])])
    SimpleDocTemplate(str(destination),pagesize=landscape(A4),leftMargin=12*mm,rightMargin=12*mm).build(elements, onFirstPage=premiere_page)

import shutil
import re
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.styles import PatternFill
from planning_soumission_pr import (
    _generer_planning_soumission_base as generer_planning_soumission,
    _assurer_planning_soumission_base as assurer_planning_soumission, mettre_a_jour_grille,
)

POSTES = {'postes': [
    {'article': '01', 'designation': 'Dévégétalisation', 'jours': 8},
    {'article': '02', 'designation': 'Échafaudage', 'jours': 10},
]}


def ancien_compact(p):
    w = Workbook()
    s = w.active
    s.title = 'Planning soumission'
    s['B3'] = 'Test'
    s['H4'] = 15
    for c, valeur in enumerate(('Article', 'Poste', 'Jours PR', 'Période', 'Début', 'Durée', 'Fin'), 1):
        s.cell(8, c, valeur)
    for r, (article, n, debut, duree) in enumerate([
        ('01', 1, 1, 2), ('01', 2, 3, 3), ('01', 3, None, None), ('01', 4, 10, 3),
        ('02', 1, 1, .25), ('02', 2, None, None), ('02', 3, None, None), ('02', 4, None, None)
    ], 9):
        for c, valeur in enumerate((article, 'Poste '+article, 8, n, debut, duree, f'=E{r}+F{r}-1'), 1):
            s.cell(r, c, valeur)
        s.row_dimensions[r].hidden = n != 4
    s['H9'].fill = PatternFill('solid', fgColor='00B0F0')
    s.auto_filter.ref = 'D8:D16'
    s.auto_filter.add_filter_column(0, ['4'])
    w.save(p)
    w.close()


class PlanningSoumissionTests(unittest.TestCase):
    def test_couleurs_suivent_saisie_modification_et_effacement(self):
        with TemporaryDirectory() as d:
            p = Path(d)/'planning.xlsx'
            generer_planning_soumission(POSTES, p, 'Test')
            w = load_workbook(p)
            s = w.active
            rules = next(iter(s.conditional_formatting._cf_rules.values()))

            def couleur(ligne, jour):
                # Évalue la règle enregistrée, en décalant les références comme Excel.
                cible = s.cell(ligne, 11+jour).coordinate
                for rule in rules:
                    formule = Translator('='+rule.formula[0], origin='L9').translate_formula(cible)[1:]
                    def valeur(match):
                        ref = match.group(0).replace('$', '')
                        cellule = s[ref]
                        return repr(cellule.column-11 if cellule.row == 8 else cellule.value)
                    expression = re.sub(r'\$?[A-Z]+\$?\d+', valeur, formule)
                    # Les règles ne proviennent que du classeur de test généré ici.
                    try:
                        active = eval(expression, {'__builtins__': {}}, {
                            'AND': lambda *args: all(args),
                            'ISNUMBER': lambda v: isinstance(v, (int, float)) and not isinstance(v, bool)})
                    except TypeError:  # Une borne vide ne colore pas la période.
                        active = False
                    if active:
                        return rule.dxf.fill.bgColor.rgb
                return None

            for r in (9, 10):
                for c, v in zip(range(4, 12), (1, 2, 3, 5, 7, 8, 10, 12)):
                    s.cell(r, c).value = v
                attendu = {1: 'FF4F81BD', 2: 'FF4F81BD', 3: 'FF548235',
                           4: 'FF548235', 5: 'FF548235', 7: 'FFC55A11',
                           8: 'FFC55A11', 10: 'FF8064A2', 11: 'FF8064A2', 12: 'FF8064A2'}
                self.assertEqual({j: couleur(r, j) for j in range(1, 15)},
                                 {j: attendu.get(j) for j in range(1, 15)})
                s.cell(r, 6).value, s.cell(r, 7).value = 13, 14
                self.assertIsNone(couleur(r, 3))
                self.assertEqual(couleur(r, 14), 'FF548235')
                s.cell(r, 6).value = s.cell(r, 7).value = None
                self.assertIsNone(couleur(r, 14))
                self.assertEqual(couleur(r, 10), 'FF8064A2')
            w.close()
            with ZipFile(p) as z:
                ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                xml = ET.fromstring(z.read('xl/styles.xml'))
                fonds = xml.findall('s:dxfs/s:dxf/s:fill/s:patternFill/s:bgColor', ns)
                self.assertEqual([v.get('rgb') for v in fonds],
                                 ['FF4F81BD', 'FF548235', 'FFC55A11', 'FF8064A2'])

    def test_ouverture_repare_ancien_remplissage_sans_changer_saisies(self):
        with TemporaryDirectory() as d:
            p = Path(d)/'planning.xlsx'
            generer_planning_soumission(POSTES, p, 'Test')
            w = load_workbook(p)
            s = w.active
            s['J10'], s['K10'] = 10, 15
            regles = next(iter(s.conditional_formatting._cf_rules.values()))
            for regle, couleur in zip(regles, ('4F81BD', '548235', 'C55A11', '8064A2')):
                regle.dxf.fill = PatternFill('solid', fgColor=couleur)
            avant = [(c.coordinate, c.value, str(c._style)) for row in s for c in row]
            w.save(p)
            w.close()
            original = p.read_bytes()
            backup = Path(d)/'avant.xlsx'
            assurer_planning_soumission('absent.xlsx', p, 'Test', None, lambda f: shutil.copy2(f, backup))
            self.assertEqual(backup.read_bytes(), original)
            w = load_workbook(p)
            self.assertEqual([(c.coordinate, c.value, str(c._style)) for row in w.active for c in row], avant)
            regles = next(iter(w.active.conditional_formatting._cf_rules.values()))
            self.assertEqual([regle.dxf.fill.bgColor.rgb for regle in regles],
                             ['FF4F81BD', 'FF548235', 'FFC55A11', 'FF8064A2'])
            w.close()
            self.assertFalse(mettre_a_jour_grille(p, lambda _: self.fail('Sauvegarde inutile')))

    def test_une_ligne_quatre_paires_de_jours_et_calendrier(self):
        with TemporaryDirectory() as d:
            p = Path(d) / 'planning.xlsx'
            generer_planning_soumission(POSTES, p, 'Test')
            w = load_workbook(p)
            s = w.active
            self.assertEqual(w.sheetnames, ['Planning soumission'])
            self.assertEqual([s.cell(8, c).value for c in range(4, 12)], ['Du jour', 'Au jour']*4)
            self.assertEqual([s.cell(7, c).value for c in (4, 6, 8, 10)],
                             ['Période 1', 'Période 2', 'Période 3', 'Période 4'])
            self.assertEqual([s.cell(r, 1).value for r in (9, 10)], ['01', '02'])
            self.assertIsNone(s.auto_filter.ref)
            self.assertFalse(any(v.hidden for v in s.row_dimensions.values()))
            self.assertFalse(any(v.hidden for v in s.column_dimensions.values()))
            self.assertEqual(s['L8'].value, '=$L$4')
            self.assertEqual((s['D9'].value, s['E9'].value), (1, 8))
            self.assertEqual(len(s._images), 1)
            self.assertIn('MAX(E9:E10,G9:G10,I9:I10,K9:K10)', s['F4'].value)
            rules = next(iter(s.conditional_formatting._cf_rules.values()))
            for rule, debut, fin in zip(rules, 'DFHJ', 'EGIK'):
                self.assertIn(f'L$8>=${debut}9,L$8<${fin}9+1', rule.formula[0])
            w.close()
            with ZipFile(p) as z:
                self.assertIsNone(z.testzip())
                self.assertNotIn('xl/worksheets/sheet2.xml', z.namelist())

    def test_migration_compact_preserve_periode_vide_et_quart_jour(self):
        with TemporaryDirectory() as d:
            p = Path(d)/'planning.xlsx'
            ancien_compact(p)
            original = p.read_bytes()
            backup = Path(d)/'avant.xlsx'
            self.assertTrue(mettre_a_jour_grille(p, lambda f: shutil.copy2(f, backup)))
            self.assertEqual(backup.read_bytes(), original)
            w = load_workbook(p)
            s = w.active
            self.assertEqual([s.cell(9, c).value for c in range(4, 12)], [1, 2, 3, 5, None, None, 10, 12])
            self.assertEqual([s.cell(10, c).value for c in (4, 5)], [1, .25])
            self.assertEqual(s['L4'].value, 15)
            self.assertEqual(s['L9'].fill.fgColor.rgb, '0000B0F0')
            self.assertFalse(s.row_dimensions[9].hidden)
            w.close()
            self.assertFalse(mettre_a_jour_grille(p, lambda _: self.fail('Sauvegarde inutile')))

    def test_migration_anciennes_colonnes(self):
        with TemporaryDirectory() as d:
            p = Path(d)/'planning.xlsx'
            with patch('planning_soumission_pr.centraliser_periodes'):
                generer_planning_soumission(POSTES, p, 'Test')
            w = load_workbook(p)
            s = w.active
            s['D9'], s['E9'] = 2, 1
            s['G9'], s['H9'], s['M9'], s['N9'] = 3, 3, 3, 10
            w.save(p)
            w.close()
            mettre_a_jour_grille(p, lambda _: None)
            w = load_workbook(p)
            self.assertEqual([w.active.cell(9, c).value for c in range(4, 12)], [1, 2, 3, 5, None, None, 10, 12])
            w.close()

    def test_reimport_pr_preserve_jours_et_formules_meme_si_ordre_change(self):
        with TemporaryDirectory() as d:
            p = Path(d)/'planning.xlsx'
            generer_planning_soumission(POSTES, p, 'Test')
            w = load_workbook(p)
            s = w.active
            for c, value in zip(range(4, 12), [1, 2, 3, 5, None, None, 10, '=J9+2']):
                s.cell(9, c).value = value
            w.save(p)
            w.close()
            original = p.read_bytes()
            assurer_planning_soumission('absent.xlsx', p, 'Test', None, lambda _: self.fail('Sauvegarde inutile'))
            self.assertEqual(p.read_bytes(), original)
            postes = {'postes': list(reversed(POSTES['postes']))}
            with patch('synthese_pr.lire_synthese_pr', return_value=postes):
                assurer_planning_soumission('pr.xlsx', p, 'Test', None, lambda _: None, actualiser=True)
            w = load_workbook(p)
            self.assertEqual([w.active.cell(10, c).value for c in range(4, 12)],
                             [1, 2, 3, 5, None, None, 10, '=J10+2'])
            w.close()

    def test_reimport_ancien_compact_preserve_periode_quatre(self):
        with TemporaryDirectory() as d:
            p = Path(d)/'planning.xlsx'
            ancien_compact(p)
            with patch('synthese_pr.lire_synthese_pr', return_value=POSTES):
                assurer_planning_soumission('pr.xlsx', p, 'Test', None, lambda _: None, actualiser=True)
            w = load_workbook(p)
            self.assertEqual([w.active.cell(9, c).value for c in range(4, 12)], [1, 2, 3, 5, None, None, 10, 12])
            w.close()

    def test_verrou_excel_et_sauvegarde_refusee_conservent_original(self):
        with TemporaryDirectory() as d:
            p = Path(d)/'planning.xlsx'
            ancien_compact(p)
            original = p.read_bytes()
            verrou = p.with_name('~$'+p.name)
            verrou.touch()
            with self.assertRaises(PermissionError):
                mettre_a_jour_grille(p, lambda _: self.fail('Ne pas sauvegarder'))
            self.assertEqual(p.read_bytes(), original)
            verrou.unlink()
            with self.assertRaises(OSError):
                mettre_a_jour_grille(p, lambda _: (_ for _ in ()).throw(OSError('Sauvegarde impossible')))
            self.assertEqual(p.read_bytes(), original)

    def test_article_supprime_ne_perd_pas_ses_jours(self):
        with TemporaryDirectory() as d:
            p = Path(d)/'planning.xlsx'
            generer_planning_soumission(POSTES, p, 'Test')
            original = p.read_bytes()
            with patch('synthese_pr.lire_synthese_pr', return_value={'postes': POSTES['postes'][:1]}):
                with self.assertRaises(ValueError):
                    assurer_planning_soumission('pr.xlsx', p, 'Test', None, lambda _: None, actualiser=True)
            self.assertEqual(p.read_bytes(), original)

    def test_article_absent_ne_perd_pas_interventions(self):
        with TemporaryDirectory() as d:
            p = Path(d)/'planning.xlsx'
            with self.assertRaises(ValueError):
                generer_planning_soumission(POSTES, p, 'Test', interventions=[('inconnu', '', 3, 10)])
            self.assertFalse(p.exists())

    def test_document_existant_jamais_ecrase(self):
        with TemporaryDirectory() as d:
            p = Path(d)/'planning.xlsx'
            generer_planning_soumission(POSTES, p, 'Test')
            original = p.read_bytes()
            with self.assertRaises(FileExistsError):
                generer_planning_soumission(POSTES, p, 'Test')
            self.assertEqual(p.read_bytes(), original)

    def test_zero_et_duree_absente_distincts(self):
        with TemporaryDirectory() as d:
            p = Path(d)/'planning.xlsx'
            generer_planning_soumission({'postes': [
                {'article': '01', 'designation': 'À chiffrer', 'jours': None},
                {'article': '02', 'designation': 'Sans durée', 'jours': 0}]}, p, 'Test')
            w = load_workbook(p)
            self.assertIsNone(w.active['C9'].value)
            self.assertIsNone(w.active['E9'].value)
            self.assertEqual(w.active['C10'].value, 0)
            self.assertEqual(w.active['E10'].value, '=IF(D10="","",D10+0-1)')
            w.close()

    def test_duree_invalide_aucun_fichier(self):
        with TemporaryDirectory() as d:
            p = Path(d)/'planning.xlsx'
            for duree in (-1, float('nan'), True):
                with self.assertRaises(ValueError):
                    generer_planning_soumission({'postes': [
                        {'article': '01', 'designation': 'Poste', 'jours': duree}]}, p, 'Test')
                self.assertFalse(p.exists())


if __name__ == '__main__':
    unittest.main()

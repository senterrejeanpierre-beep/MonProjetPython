"""Parcours du menu ; les essais VBA utilisent un interpréteur, pas Excel installé."""
import ast
from pathlib import Path
import shutil
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter
from planning_soumission_pr import (
    generer_planning_soumission, assurer_planning_soumission, _generer_planning_soumission_base,
)
from planning_soumission_menu import lire_plages_menu, est_menu, fermer_planning
from planning_soumission_periodes import lire_plages
from test_planning_soumission_pr import POSTES

try:
    from pyopenvba.apps.excel import ExcelApplication
except ImportError:
    ExcelApplication = None


class MenuPlanningTests(unittest.TestCase):
    def test_creation_une_ligne_selecteur_et_deux_jours(self):
        with TemporaryDirectory() as d:
            p = generer_planning_soumission(POSTES, Path(d)/'Planning_soumission.xlsx', 'Test')
            self.assertEqual(p.suffix, '.xlsm')
            w = load_workbook(p, keep_vba=True)
            s = w.active
            self.assertEqual(w.sheetnames, ['Planning soumission'])
            self.assertTrue(est_menu(s))
            self.assertEqual([s.cell(r, 1).value for r in (9, 10)], ['01', '02'])
            self.assertIsNone(s.auto_filter.ref)
            self.assertFalse(any(v.hidden for v in s.row_dimensions.values()))
            self.assertEqual(s['G8'].value, '=$G$4')
            listes = [v for v in s.data_validations.dataValidation if v.type == 'list']
            self.assertEqual(len(listes), 2)
            self.assertEqual(str(listes[0].sqref), 'D9')
            self.assertEqual(listes[0].formula1, '$FL$9:$FO$9')
            self.assertEqual(listes[1].formula1, '$FL$10:$FO$10')
            self.assertFalse(listes[0].showDropDown)
            self.assertFalse(listes[0].showErrorMessage)
            self.assertTrue(all(s.column_dimensions[get_column_letter(c)].hidden for c in range(147, 172)))
            self.assertEqual(w.code_name, 'ThisWorkbook')
            fermer_planning(w)
            with ZipFile(p) as z:
                self.assertIsNone(z.testzip())
                self.assertEqual(z.read('xl/vbaProject.bin'),
                    Path('ressources/planning_soumission/vbaProject.bin').read_bytes())
                self.assertFalse(any('externalLinks' in n for n in z.namelist()))

    def test_reprise_preserve_saisies_et_archive_ancien(self):
        with TemporaryDirectory() as d:
            p = Path(d)/'Planning_soumission.xlsx'
            _generer_planning_soumission_base(POSTES, p, 'Test')
            w = load_workbook(p)
            s = w.active
            s['F9'], s['G9'], s['J9'], s['K9'] = 3, 5, 10, '=J9+2'
            s['L4'] = 7
            avant = lire_plages(s)
            w.save(p)
            fermer_planning(w)
            octets = p.read_bytes()
            backup = Path(d)/'avant.xlsx'
            q = assurer_planning_soumission('absent', p, 'Test', None, lambda f: shutil.copy2(f, backup))
            self.assertFalse(p.exists())
            self.assertEqual(backup.read_bytes(), octets)
            w = load_workbook(q, keep_vba=True)
            self.assertEqual(lire_plages_menu(w.active), avant)
            self.assertEqual(w.active['G4'].value, 7)
            fermer_planning(w)
            resultat = q.read_bytes()
            self.assertEqual(assurer_planning_soumission('absent', p, 'Test', None, lambda _: self.fail()), q)
            self.assertEqual(q.read_bytes(), resultat)

    def test_reimport_pr_preserve_choix_et_periodes(self):
        with TemporaryDirectory() as d:
            p = Path(d)/'Planning_soumission.xlsx'
            q = generer_planning_soumission(POSTES, p, 'Test')
            w = load_workbook(q, keep_vba=True)
            s = w.active
            s['D9'] = '4 – Finitions'
            s.cell(9, 159, 4)
            s.cell(9, 168, '1 – Préparation')
            s.cell(9, 171, '4 – Finitions')
            s.cell(9, 166, 10)
            s.cell(9, 167, '=FJ9+2')
            s['G4'] = 8
            w.save(q)
            fermer_planning(w)
            with patch('synthese_pr.lire_synthese_pr', return_value={'postes':list(reversed(POSTES['postes']))}):
                assurer_planning_soumission('pr', p, 'Test', None, lambda _: None, actualiser=True)
            w = load_workbook(q, keep_vba=True)
            s = w.active
            self.assertEqual(s['A10'].value, '01')
            self.assertEqual(s['D10'].value, '4 – Finitions')
            self.assertEqual(s.cell(10,168).value, '1 – Préparation')
            self.assertEqual(s.cell(10,171).value, '4 – Finitions')
            self.assertEqual(s.cell(10,166).value, 10)
            self.assertEqual(s.cell(10,167).value, '=FJ10+2')
            self.assertEqual(s['G4'].value, 8)
            self.assertEqual(s['E10'].value, '=IF(FJ10="","",FJ10)')
            fermer_planning(w)

    def test_ancien_menu_actualise_sans_changer_les_jours(self):
        with TemporaryDirectory() as d:
            p = generer_planning_soumission(POSTES, Path(d)/'planning.xlsm', 'Test')
            w = load_workbook(p, keep_vba=True)
            s = w.active
            s.cell(3,159,'HorizonPeriodes1')
            for r in (9,10):
                for c in range(168,172):
                    s.cell(r,c).value = None
            s['D9'] = 4
            s.cell(9,159,4)
            s.cell(9,166,10)
            s.cell(9,167,12)
            avant = lire_plages_menu(s)
            w.save(p)
            fermer_planning(w)
            original = p.read_bytes()
            backup = Path(d)/'avant.xlsm'
            assurer_planning_soumission('absent', p, 'Test', None, lambda f: shutil.copy2(f,backup))
            self.assertEqual(backup.read_bytes(), original)
            w = load_workbook(p, keep_vba=True)
            self.assertEqual(lire_plages_menu(w.active), avant)
            self.assertEqual(w.active['D9'].value, 4)
            self.assertEqual(w.active.cell(3,159).value, 'HorizonPeriodesTexte1')
            fermer_planning(w)
            nouveau = p.read_bytes()
            assurer_planning_soumission('absent', p, 'Test', None, lambda _: self.fail())
            self.assertEqual(p.read_bytes(), nouveau)

    @unittest.skipIf(ExcelApplication is None, 'Interpréteur VBA de développement absent')
    def test_ligne_uniquement_en_dessous_jusqu_a_fin_couleurs(self):
        with TemporaryDirectory() as d:
            p = generer_planning_soumission(POSTES, Path(d)/'planning.xlsm', 'Test')
            w = load_workbook(p,keep_vba=True)
            regle = next(v for regles in w.active.conditional_formatting._cf_rules.values()
                         for v in regles if v.dxf.border is not None)
            self.assertEqual(regle.dxf.border.bottom.style,'thin')
            for cote in ('left','right','top'):
                self.assertIsNone(getattr(regle.dxf.border,cote))
            fermer_planning(w)
            a = ExcelApplication.open(p)
            s = a.sheet(1)
            self.assertFalse(any(v.name.startswith(('HorizonGuide_','HorizonCadre_')) for v in s.shapes()))
            def dessous(ref):
                f = Translator('='+regle.formula[0],origin='A9').translate_formula(ref)
                # Evaluate est limité à 255 caractères ; chaque période est évaluée séparément.
                import re
                valeurs = [a.evaluate('Application.Evaluate("'+part.replace('"','""')+'")')
                           for part in re.findall(r'IFERROR\(.*?,FALSE\)',f)]
                self.assertTrue(all(isinstance(v,bool) for v in valeurs),valeurs)
                return any(valeurs)
            for ref,valeur in [('D9',1),('E9',10),('F9',12),('D9',2),('E9',3),('F9',5)]:
                s.set_value(ref,valeur)
            for ref in ('A9','B9','F9','G9','O9','R9'):
                self.assertTrue(dessous(ref),ref)
            self.assertFalse(dessous('S9'))
            self.assertFalse(dessous('A10'))
            s.set_value('D9',1)
            s.set_value('F9',15)
            self.assertTrue(dessous('U9'))
            self.assertFalse(dessous('V9'))
            s.set_value('D9','plateau')
            self.assertTrue(dessous('U9'))
            s.set_value('G4',12)
            self.assertTrue(dessous('J9'))
            self.assertFalse(dessous('K9'))
            a.add_module('Sub EffacerRepere()\nThisWorkbook.Worksheets(1).Range("E9:F9").ClearContents\nEnd Sub',name='EssaiReperes')
            a.run('EffacerRepere')
            self.assertFalse(dessous('A9'))
            s.set_value('G4',1)
            self.assertTrue(dessous('K9'))
            self.assertFalse(dessous('L9'))
            s.set_value('D9',2)
            a.run('EffacerRepere')
            self.assertFalse(dessous('A9'))
            # Une durée fractionnaire conserve exactement le même dernier jour que la couleur.
            s.set_value('E9',1)
            s.set_value('F9',.25)
            self.assertTrue(dessous('G9'))
            self.assertFalse(dessous('H9'))
            self.assertEqual(a.dialogs,[])

    def test_reprise_legendes_ajoute_reperes_preserve_donnees(self):
        from xml.etree import ElementTree as ET
        from planning_soumission_legendes import ecrire_legendes, X, A
        with TemporaryDirectory() as d:
            p = generer_planning_soumission(POSTES, Path(d)/'planning.xlsm', 'Test')
            w = load_workbook(p,keep_vba=True)
            s = w.active
            s.cell(4,159,'HorizonLegendes2')
            s.cell(9,160,3)
            s.cell(9,161,8)
            s.cell(9,168,'1 – plateau')
            s['D9'] = '1 – plateau'
            avant = lire_plages_menu(s)
            w.save(p)
            fermer_planning(w)
            ecrire_legendes(p)
            original = p.read_bytes()
            backup = Path(d)/'avant.xlsm'
            assurer_planning_soumission('absent',p,'Test',None,lambda f: shutil.copy2(f,backup))
            self.assertEqual(backup.read_bytes(),original)
            w = load_workbook(p,keep_vba=True)
            self.assertEqual(lire_plages_menu(w.active),avant)
            self.assertEqual(w.active['D9'].value,'plateau')
            self.assertEqual(w.active.cell(4,159).value,'HorizonLegendes5')
            fermer_planning(w)
            with ZipFile(p) as z:
                dessins = [ET.fromstring(z.read(n)) for n in z.namelist() if n.startswith('xl/drawings/') and n.endswith('.xml')]
            formes = {sp.find(f'{{{X}}}nvSpPr/{{{X}}}cNvPr').get('name'):sp
                      for dessin in dessins for sp in dessin.iter(f'{{{X}}}sp')}
            self.assertFalse(any(n.startswith(('HorizonGuide_', 'HorizonCadre_')) for n in formes))
            octets = p.read_bytes()
            assurer_planning_soumission('absent',p,'Test',None,lambda _: self.fail())
            self.assertEqual(p.read_bytes(),octets)

    @unittest.skipIf(ExcelApplication is None, 'Interpréteur VBA de développement absent')
    def test_texte_dans_liste_et_dans_barre_apres_deplacement(self):
        with TemporaryDirectory() as d:
            p = generer_planning_soumission(POSTES, Path(d)/'planning.xlsm', 'Test')
            a = ExcelApplication.open(p)
            s = a.sheet(1)
            for ref, valeur in [('D9','cour d’honneur'),('E9',3),('F9',8),
                                ('D9',2),('D9','mur'),('E9',12),('F9',16)]:
                s.set_value(ref, valeur)
            cour = s.shape('HorizonLegende_9_1')
            mur = s.shape('HorizonLegende_9_2')
            self.assertEqual((cour.text,mur.text),('cour d’honneur','mur'))
            self.assertGreater(cour.height, 1)
            haut = a.evaluate('ThisWorkbook.Worksheets(1).Range("I9").Top')
            hauteur = a.evaluate('ThisWorkbook.Worksheets(1).Range("I9").Height')
            self.assertAlmostEqual(cour.top,haut+1,places=3)
            self.assertAlmostEqual(cour.height,hauteur-2,places=3)
            self.assertFalse(any(v.name.startswith('HorizonFond_') for v in s.shapes()))
            self.assertEqual(s.value('D9'),'mur')
            self.assertEqual(s.value('FL9'),'cour d’honneur')
            self.assertLessEqual(cour.left+cour.width,mur.left)
            gauche = cour.left
            s.set_value('D9',1)
            s.set_value('E9',5)
            s.set_value('F9',10)
            self.assertGreater(s.shape('HorizonLegende_9_1').left,gauche)
            self.assertEqual(s.shape('HorizonLegende_9_2').text,'mur')
            # Le texte est aussi éditable directement dans sa zone dans la barre.
            a.add_module('''Sub TexteDirect()
ThisWorkbook.Worksheets("Planning soumission").Shapes("HorizonLegende_9_1").TextFrame.Characters.Text = "plateau"
End Sub
Sub SauverLegendes()
ThisWorkbook.Save
End Sub''',name='EssaiLegendes')
            a.run('TexteDirect')
            # Le moteur n'émet pas BeforeSave depuis Save ; on déclenche l'événement Excel.
            from pyopenvba.apps.excel._events import _handle
            _handle(a.workbook, 'Workbook_BeforeSave', [False, False])
            a.run('SauverLegendes')
            self.assertEqual(s.value('D9'),'plateau')
            self.assertEqual((s.value('E9'),s.value('F9')),(5,10))
            self.assertEqual(a.dialogs,[])
            a = ExcelApplication.open(p)
            s = a.sheet(1)
            self.assertEqual(s.shape('HorizonLegende_9_1').text,'plateau')
            s.set_value('D9',2)
            self.assertEqual(s.value('D9'),'mur')
            self.assertEqual((s.value('E9'),s.value('F9')),(12,16))

    @unittest.skipIf(ExcelApplication is None, 'Interpréteur VBA de développement absent')
    def test_texte_dans_liste_conserve_jours_et_independance_des_postes(self):
        with TemporaryDirectory() as d:
            p = generer_planning_soumission(POSTES, Path(d)/'planning.xlsm', 'Test')
            a = ExcelApplication.open(p)
            s = a.sheet(1)
            for cellule,valeur in [('D9','Préparation'),('E9',1),('F9',2),
                                    ('D9',2),('D9','Dépose, façade; côté cour'),('E9',3),('F9',5),
                                    ('D10','Installation')]:
                s.set_value(cellule,valeur)
            s.set_value('D9','Préparation')
            self.assertEqual((s.value('E9'),s.value('F9')),(1,2))
            s.set_value('D9','Préparation du mur')
            self.assertEqual(s.value('D9'),'Préparation du mur')
            s.set_value('D9','Dépose, façade; côté cour')
            self.assertEqual((s.value('E9'),s.value('F9')),(3,5))
            s.set_value('D9','Rejointoiement')
            self.assertEqual((s.value('E9'),s.value('F9')),(3,5))
            a.save(p)
            a = ExcelApplication.open(p)
            s = a.sheet(1)
            s.set_value('D9',1)
            self.assertEqual(s.value('D9'),'Préparation du mur')
            self.assertEqual((s.value('E9'),s.value('F9')),(1,2))
            s.set_value('D9',2)
            self.assertEqual(s.value('D9'),'Rejointoiement')
            self.assertEqual((s.value('E9'),s.value('F9')),(3,5))
            self.assertEqual(s.value('D10'),'Installation')
            # Effacer le texte conserve la période et ses jours.
            a.add_module('Sub EffacerTexte()\nThisWorkbook.Worksheets("Planning soumission").Range("D9").ClearContents\nEnd Sub', name='EssaiTexte')
            a.run('EffacerTexte')
            self.assertEqual(s.value('D9'),2)
            self.assertEqual((s.value('E9'),s.value('F9')),(3,5))
            # Texte réellement libre : ponctuation, espaces, chiffres et préfixes personnels.
            for texte in ('  cour d’honneur / plateau  ', '2026', '2 - mur', '=mur', 'mur\ncôté cour'):
                s.set_value('D9',texte)
                self.assertEqual(s.value('D9'),texte)
                self.assertEqual(s.value('FM9'),texte)
                self.assertEqual(s.shape('HorizonLegende_9_2').text,texte)
                self.assertEqual((s.value('E9'),s.value('F9')),(3,5))
            self.assertEqual(a.dialogs,[])

    def test_verrou_et_echec_sauvegarde_preservent_original(self):
        with TemporaryDirectory() as d:
            p = Path(d)/'Planning_soumission.xlsx'
            _generer_planning_soumission_base(POSTES, p, 'Test')
            original = p.read_bytes()
            verrou = p.with_name('~$'+p.name)
            verrou.touch()
            with self.assertRaises(PermissionError):
                assurer_planning_soumission('absent', p, 'Test', None, lambda _: self.fail())
            verrou.unlink()
            with self.assertRaises(OSError):
                assurer_planning_soumission('absent', p, 'Test', None,
                    lambda _: (_ for _ in ()).throw(OSError('Sauvegarde refusée')))
            self.assertEqual(p.read_bytes(), original)
            self.assertFalse(p.with_suffix('.xlsm').exists())

    def test_application_ouvre_le_fichier_retourne(self):
        arbre = ast.parse(Path('main.py').read_text())
        classe = next(n for n in arbre.body if isinstance(n, ast.ClassDef) and n.name == 'HorizonChantierApp')
        methode = next(n for n in classe.body if isinstance(n, ast.FunctionDef) and n.name == 'ouvrir_planning_soumission_pr')
        ns = {'Path':Path, '_normaliser_dossier_data':lambda p:p/'DATA',
              'dossier_base':lambda:Path('.'), '_backup_excel_before_write':lambda _:None}
        exec(compile(ast.Module(body=[methode], type_ignores=[]), 'main.py', 'exec'), ns)
        class App:
            def _nom_chantier_selectionne(self): return 'Test'
            def _dossier_chantier_selectionne(self): return Path('chantier déplacé')
            def _ouvrir_fichier_chantier(self, nom): self.ouvert = nom
        app = App()
        with patch('planning_soumission_pr.assurer_planning_soumission', return_value=Path('chantier déplacé/Planning_soumission.xlsm')):
            ns['ouvrir_planning_soumission_pr'](app)
        self.assertEqual(app.ouvert, 'Planning_soumission.xlsm')

    @unittest.skipIf(ExcelApplication is None, 'Interpréteur VBA de développement absent')
    def test_parcours_vba_selection_saisie_couleurs_enregistrement(self):
        with TemporaryDirectory() as d:
            p = generer_planning_soumission(POSTES, Path(d)/'Planning_soumission.xlsm', 'Test')
            a = ExcelApplication.open(p)
            s = a.sheet(1)
            # Le simulateur conserve les caches vides initiaux : invalide les jours.
            s.set_value('G4', 1)
            w = load_workbook(p, keep_vba=True)
            regles = next(iter(w.active.conditional_formatting._cf_rules.values()))
            fermer_planning(w)
            def couleur(jour):
                for regle in regles:
                    f = Translator('='+regle.formula[0], origin='G9').translate_formula(f'{get_column_letter(6+jour)}9')
                    if a.evaluate('Application.Evaluate("'+f.replace('"','""')+'")'):
                        return regle.dxf.fill.bgColor.rgb
                return None
            for periode, debut, fin in ((1,1,2), (2,3,5), (4,10,12)):
                s.set_value('D9', periode)
                s.set_value('E9', debut)
                s.set_value('F9', fin)
            for periode, jours in ((1,(1,2)), (2,(3,5)), (3,('','')), (4,(10,12))):
                s.set_value('D9', periode)
                self.assertEqual((s.value('E9'),s.value('F9')), jours)
                self.assertEqual([couleur(j) for j in (1,2,3,4,5,6,10,11,12,13)],
                    ['FF9DBAD5','FF9DBAD5','FF9DBAD5','FF9DBAD5','FF9DBAD5',None,
                     'FF9DBAD5','FF9DBAD5','FF9DBAD5',None])
            s.set_value('D9',2)
            s.set_value('E9',6)
            s.set_value('F9',7)
            self.assertIsNone(couleur(3))
            self.assertEqual(couleur(6),'FF9DBAD5')
            a.add_module('Sub EffacerTest()\nThisWorkbook.Worksheets("Planning soumission").Range("E9:F9").ClearContents\nEnd Sub', name='EssaiEffacement')
            a.run('EffacerTest')
            self.assertIsNone(couleur(6))
            self.assertEqual(couleur(10),'FF9DBAD5')
            a.save(p)
            a = ExcelApplication.open(p)
            s = a.sheet(1)
            s.set_value('D9',4)
            self.assertEqual((s.value('E9'),s.value('F9')),(10,12))
            s.set_value('D9',1)
            self.assertEqual((s.value('E9'),s.value('F9')),(1,2))


if __name__ == '__main__':
    unittest.main()

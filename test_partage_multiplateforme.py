import json
import os
import shutil
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch
from environnement_partage import trouver_base
from bordereau_public import lire_config
from main import ouvrir_chemin, excel_set_cell_value, _selectionner_etat

class Partage(unittest.TestCase):
    def test_deplacement_projet_et_chantier_avec_accents(self):
        with tempfile.TemporaryDirectory() as tmp, patch.dict(os.environ, {}, clear=True):
            racine = Path(tmp)
            projet = racine / 'Partage été' / 'MonProjetPython'; projet.mkdir(parents=True)
            base = projet.parent / 'Horizon_Chantier_Data'; chantier = base / 'Chantiers' / '004_Mariemont'
            chantier.mkdir(parents=True)
            (chantier / 'état_avancement.xlsx').touch()
            (chantier / 'bordereau_public.json').write_text(json.dumps(
                {'original': 'métré.xlsx', 'copie': 'état_avancement.xlsx'}, ensure_ascii=False), encoding='utf-8')
            self.assertEqual(trouver_base(projet, racine/'absent.json'), base)
            dest = racine/'Autre poste'; shutil.copytree(projet.parent,dest)
            trouve = trouver_base(dest/'MonProjetPython',racine/'absent.json')
            self.assertEqual(_selectionner_etat(trouve/'Chantiers'/'004_Mariemont').name,'état_avancement.xlsx')

    def test_partage_deconnecte_ne_bascule_pas_sur_un_autre_chantier(self):
        with tempfile.TemporaryDirectory() as tmp, patch.dict(os.environ, {}, clear=True):
            p=Path(tmp); config=p/'poste.json'
            config.write_text(json.dumps({'dossier':str(p/'absent')}),encoding='utf-8')
            with self.assertRaises(FileNotFoundError):trouver_base(p,config)

    def test_chemins_config_portables(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)
            for name in ('../original.xlsx', r'..\original.xlsx', r'C:\original.xlsx', '/original.xlsx'):
                (p/'bordereau_public.json').write_text(json.dumps({'original':name,'copie':'copie.xlsx'}),encoding='utf-8')
                with self.assertRaises(ValueError):lire_config(p)

    def test_pilotage_lit_les_montants_apres_les_libelles_fusionnes(self):
        from openpyxl import Workbook
        from main import _montant_a_droite, _nombre_metier
        ws = Workbook().active
        ws.merge_cells('P173:S173')
        ws['P173'] = 'Total état cumulé hors TVA'
        ws['T173'] = 142706.93
        ws.merge_cells('U173:V173')
        ws['U173'] = 'Montant de l’état 3'
        ws['W173'] = 8810
        self.assertEqual(_montant_a_droite(ws,173,16,_nombre_metier),142706.93)
        self.assertEqual(_montant_a_droite(ws,173,21,_nombre_metier),8810)

    def test_ouverture_windows_et_mac(self):
        start = Mock()
        with patch('main.sys.platform','win32'), patch('main.os',SimpleNamespace(name='nt',startfile=start)):
            ouvrir_chemin('chantier été.xlsx')
        start.assert_called_once_with('chantier été.xlsx')
        with patch('main.sys.platform','darwin'), patch('main.subprocess.Popen') as popen:
            ouvrir_chemin('chantier été.xlsx')
        popen.assert_called_once_with(['open','chantier été.xlsx'])

    def test_excel_windows_constante_modifiable_formule_protegee(self):
        cellule=SimpleNamespace(formula='ancien texte',value='ancien texte')
        livre=SimpleNamespace(name='PR.xlsx',sheets={'Feuil1':SimpleNamespace(range=lambda a:cellule)},activate=Mock())
        xw=SimpleNamespace(apps=[SimpleNamespace(books=[livre])])
        with patch('main.sys.platform','win32'),patch.dict('sys.modules',{'xlwings':xw}):
            excel_set_cell_value('PR.xlsx','Feuil1','P3','nouveau')
            self.assertEqual(cellule.value,'nouveau')
            cellule.formula='=1+1'
            with self.assertRaises(ValueError):excel_set_cell_value('PR.xlsx','Feuil1','P3','interdit')
            self.assertEqual(cellule.value,'nouveau')

if __name__=='__main__':unittest.main()

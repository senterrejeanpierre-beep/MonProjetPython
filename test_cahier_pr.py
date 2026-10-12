import json
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parent / '_vendor'))
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader, PdfWriter
import cahier_pr as cahier
import liens_horizon as liens


def creer_pdf(path, tailles=(100, 200, 300, 400)):
    writer = PdfWriter()
    for width in tailles:
        writer.add_blank_page(width=width, height=500)
    writer.write(path)


def creer_pr(path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Chiffrage'
    for row, code in ((3, '01.19.1a.01'), (45, '04.19.1a.01')):
        ws[f'B{row}'] = code
        ws[f'C{row}'] = 'Désignation ' + code
        ws[f'D{row+1}'] = 'Saisie matière à garder'
        ws[f'F{row+27}'] = f'=SUM(I{row}:I{row+8})'
    ws['C4'] = 'Cahier technique · p. 9'
    ws['C4'].hyperlink = "#'Anciennes pages'!A1"
    wb.create_sheet('Anciennes pages')['A1'] = 'Ne pas supprimer'
    wb.save(path)


class CahierPRTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.base = Path(self.tmp.name)
        self.chantier = self.base / '005_Chapelle'
        (self.chantier / 'data').mkdir(parents=True)
        self.pr = self.chantier / 'data/prix_de_revient.xlsx'
        self.pdf = self.chantier / 'data' / 'Cahier_technique.pdf'
        creer_pr(self.pr)
        creer_pdf(self.pdf)

    def test_pages_intervalles_et_validation(self):
        self.assertEqual(cahier.lire_pages('2–4, 1; 3', 4), [2, 3, 4, 1])
        for saisie in ('', '0', '5', '3-2', '-1', '1,,2', 'un', '1-999999999'):
            with self.subTest(saisie=saisie), self.assertRaises(ValueError):
                cahier.lire_pages(saisie, 4)

    def test_pdf_du_chantier_et_choix_manuel(self):
        self.assertEqual(cahier.trouver_cahier(self.chantier), self.pdf)
        self.assertEqual(cahier.choix_article(self.chantier, '01.19.1a.01', self.pdf), ('', False))
        extrait = cahier.preparer_pages(self.chantier, '01.19.1a.01', self.pdf, '3-4, 1')
        self.assertEqual([int(p.mediabox.width) for p in PdfReader(extrait).pages], [300, 400, 100])
        self.assertEqual(cahier.choix_article(self.chantier, '01.19.1a.01', self.pdf), ('3-4, 1', False))
        self.assertEqual(cahier.choix_article(self.chantier, '04.19.1a.01', self.pdf), ('', False))
        config = (self.chantier / cahier.CONFIG).read_text()
        self.assertNotIn(str(self.base), config)
        avant = self.pr.read_bytes()
        cahier.preparer_pages(self.chantier, '01.19.1a.01', self.pdf, '2')
        self.assertEqual(self.pr.read_bytes(), avant)

    def test_changement_de_version_invalide_les_anciennes_pages(self):
        cahier.preparer_pages(self.chantier, '1', self.pdf, '4')
        creer_pdf(self.pdf, (800, 900))
        self.assertEqual(cahier.choix_article(self.chantier, '1', self.pdf), ('', True))
        config_avant = (self.chantier / cahier.CONFIG).read_bytes()
        with self.assertRaises(ValueError):
            cahier.preparer_pages(self.chantier, '1', self.pdf, '4')
        self.assertEqual(config_avant, (self.chantier / cahier.CONFIG).read_bytes())
        extrait = cahier.preparer_pages(self.chantier, '1', self.pdf, '2')
        self.assertEqual(int(PdfReader(extrait).pages[0].mediabox.width), 900)

    def test_liens_existants_et_futurs_preservent_pr_et_sont_idempotents(self):
        with ZipFile(self.pr) as z:
            avant = {n: z.read(n) for n in z.namelist()}
        backups = []
        self.assertEqual(cahier.assurer_liens(self.chantier, self.pr, lambda p: backups.append(p.read_bytes())), 2)
        self.assertEqual(len(backups), 1)
        with ZipFile(self.pr) as z:
            for n, value in avant.items():
                if n not in ('xl/worksheets/sheet1.xml', 'xl/worksheets/_rels/sheet1.xml.rels'):
                    self.assertEqual(z.read(n), value, n)
        w = load_workbook(self.pr)
        for ref, code in [('C4', '01.19.1a.01'), ('C46', '04.19.1a.01')]:
            lien = w['Chiffrage'][ref].hyperlink.target
            self.assertEqual(liens.resoudre(self.base, lien), (self.chantier, code))
            self.assertNotIn(str(self.base), lien)
        self.assertEqual(w['Chiffrage']['F30'].value, '=SUM(I3:I11)')
        self.assertEqual(w['Chiffrage']['D4'].value, 'Saisie matière à garder')
        w.close()
        signature = cahier.empreinte(self.pr)
        self.assertEqual(cahier.assurer_liens(self.chantier, self.pr), 0)
        self.assertEqual(cahier.empreinte(self.pr), signature)
        w = load_workbook(self.pr)
        w['Chiffrage']['B87'], w['Chiffrage']['C87'] = '93.31.1x.01', 'Nouveau poste'
        w['Chiffrage']['D88'] = 'Matière'
        w.save(self.pr)
        self.assertEqual(cahier.assurer_liens(self.chantier, self.pr), 1)

    def test_deplacement_conserve_cahier_pages_et_liens(self):
        cahier.assurer_liens(self.chantier, self.pr)
        cahier.preparer_pages(self.chantier, '01.19.1a.01', self.pdf, '2-3')
        lien = liens.lien_cahier(self.chantier, '01.19.1a.01')
        nouvelle_base = self.base / 'Partage Windows été'
        nouvelle_base.mkdir()
        nouveau = nouvelle_base / '005 renommé'
        self.chantier.rename(nouveau)
        self.assertEqual(liens.resoudre(nouvelle_base, lien), (nouveau, '01.19.1a.01'))
        pdf = cahier.trouver_cahier(nouveau)
        self.assertEqual(pdf, (nouveau / 'data' / self.pdf.name).resolve())
        self.assertEqual(cahier.choix_article(nouveau, '01.19.1a.01', pdf), ('2-3', False))

    def test_deux_chantiers_meme_article_ne_partagent_pas_les_pages(self):
        autre = self.base / '009_autre'
        autre.mkdir()
        pdf = autre / 'Cahier_technique.pdf'
        creer_pdf(pdf, (800,))
        cahier.preparer_pages(self.chantier, '1', self.pdf, '4')
        extrait = cahier.preparer_pages(autre, '1', pdf, '1')
        self.assertEqual(int(PdfReader(extrait).pages[0].mediabox.width), 800)
        self.assertEqual(cahier.choix_article(self.chantier, '1', self.pdf), ('4', False))
        self.assertNotEqual(liens.lien_cahier(autre, '1'), liens.lien_cahier(self.chantier, '1'))

    def test_cahier_choisi_manquant_ne_bascule_pas_sur_autre_pdf(self):
        cahier.preparer_pages(self.chantier, '1', self.pdf, '1')
        self.pdf.unlink()
        creer_pdf(self.chantier / 'Cahier_technique.pdf')
        self.assertEqual(cahier.trouver_cahier(self.chantier), self.pdf.resolve())

    def test_pas_de_chemin_hors_chantier(self):
        for rel in ('../secret.pdf', '/secret.pdf', r'C:\secret.pdf', r'..\secret.pdf'):
            with self.subTest(rel=rel), self.assertRaises(ValueError):
                cahier.chemin_relatif(self.chantier, rel)

    def test_pr_ouvert_non_modifie(self):
        self.pr.with_name('~$' + self.pr.name).touch()
        signature = cahier.empreinte(self.pr)
        with self.assertRaises(PermissionError):
            cahier.assurer_liens(self.chantier, self.pr)
        self.assertEqual(cahier.empreinte(self.pr), signature)

    def test_nouveau_chantier_recoit_ses_propres_liens_sans_choix_herites(self):
        from main import _creer_documents_chantier
        modeles = self.base / 'Modele'
        modeles.mkdir()
        cahier.assurer_liens(self.chantier, self.pr)
        shutil.copy2(self.pr, modeles / 'prix_de_revient model.xlsx')
        shutil.copy2(self.chantier / liens.MARQUEUR, modeles / liens.MARQUEUR)
        (modeles / cahier.CONFIG).write_text('{"fichier":"ancien.pdf","articles":{}}')
        _creer_documents_chantier(self.base, '010_Nouveau', {})
        nouveau = self.base / '010_Nouveau'
        self.assertFalse((nouveau / cahier.CONFIG).exists())
        w = load_workbook(nouveau / 'data/prix_de_revient.xlsx')
        lien = w['Chiffrage']['C4'].hyperlink.target
        self.assertEqual(liens.resoudre(self.base, lien), (nouveau, '01.19.1a.01'))
        w.close()

    def test_reception_liens_cahier_et_refus_liens_invalides(self):
        with patch.object(liens, 'dossier_local', return_value=self.base / 'local'):
            lien = liens.lien_cahier(self.chantier, '01.19.1a.01')
            liens.deposer(lien)
            self.assertEqual(list(liens.commandes()), [lien])
        identifiant = liens.identifier(self.chantier)
        for suffix in ('../1', '%2Ftmp%2F1', '1?execute=oui', '1#page', '%5Csecret'):
            with self.subTest(suffix=suffix), self.assertRaises(ValueError):
                liens.analyser_cahier(f'horizonchantier://cahier/{identifiant}/{suffix}')

    def test_rangement_ancien_chantier_et_bouton_cahier_dans_data(self):
        from main import _chemin_fichier_chantier
        ancien = self.chantier / self.pdf.name
        self.pdf.rename(ancien)
        cahier.preparer_pages(self.chantier, '1', ancien, '2-3')
        pdf = _chemin_fichier_chantier(self.chantier, 'Cahier_technique.pdf')
        self.assertEqual(pdf, self.pdf.resolve())
        self.assertEqual(pdf.read_bytes(), ancien.read_bytes())
        self.assertEqual(cahier.choix_article(self.chantier, '1', pdf), ('2-3', False))
        config = (self.chantier / cahier.CONFIG).read_bytes()
        cahier.ranger_cahier(self.chantier)
        self.assertEqual((self.chantier / cahier.CONFIG).read_bytes(), config)
        self.assertEqual(len(list((self.chantier / 'data').glob('*.pdf'))), 1)

    def test_import_data_ne_remplace_pas_un_cahier_existant(self):
        source = self.base / self.pdf.name
        creer_pdf(source, (700,))
        avant = self.pdf.read_bytes()
        copie = cahier.ranger_cahier(self.chantier, source)
        self.assertEqual(copie.parent, (self.chantier / 'data').resolve())
        self.assertNotEqual(copie.name, self.pdf.name)
        self.assertEqual(self.pdf.read_bytes(), avant)
        self.assertEqual(cahier.trouver_cahier(self.chantier), copie)

    def test_recherche_technique_ne_prend_pas_administratif(self):
        self.pdf.unlink()
        creer_pdf(self.chantier / 'data/Cahier_administratif.pdf')
        self.assertIsNone(cahier.trouver_cahier(self.chantier))

    def test_recherche_pages_hors_sommaire_et_rubriques_communes(self):
        from reportlab.pdfgen.canvas import Canvas
        c = Canvas(str(self.pdf))
        for lignes in (
            ['TABLE DES MATIERES', '01.19.1a Installation ............ 2', '04.19.1a Travaux ............ 4'],
            ['Voir aussi 04.19.1a Travaux', '01.19.1a Installation', 'Texte du premier poste'],
            ['Suite du premier poste', '04.19.1a Travaux', 'Texte du second poste'],
            ['Suite du second poste'],
        ):
            texte = c.beginText(40, 800)
            for ligne in lignes:
                texte.textLine(ligne)
            c.drawText(texte)
            c.showPage()
        c.save()
        # L'ordre des postes du PR ne fixe pas la fin d'une rubrique PDF.
        articles = ['04.19.1a.01', '01.19.1a.02', '01.19.1a.01', '99.99.1a.01']
        self.assertEqual(cahier._pages_automatiques(self.pdf, articles), {
            '04.19.1a.01': '3-4', '01.19.1a.02': '2-3', '01.19.1a.01': '2-3'})

    def test_integration_preserve_objets_formules_styles_et_retour_exact(self):
        wb = load_workbook(self.pr)
        wb['Chiffrage']['C4'] = 'Suite de désignation'
        wb['Chiffrage']['C4'].hyperlink = None
        wb['Chiffrage']['C5'] = 'Cahier technique · p. 2'
        wb['Chiffrage']['C5']._style = wb['Chiffrage']['C4']._style
        wb.save(self.pr)
        wb.close()
        with ZipFile(self.pr) as z:
            contenu = {n: z.read(n) for n in z.namelist()}
        contenu['xl/worksheets/sheet1.xml'] = contenu['xl/worksheets/sheet1.xml'].replace(
            b'<f>SUM(I3:I11)</f><v></v>', b'<f>SUM(I3:I11)</f><v>123</v>')
        contenu['xl/customData/objet_a_conserver.bin'] = b'objet Excel intact'
        with ZipFile(self.pr, 'w') as z:
            for n, value in contenu.items():
                z.writestr(n, value)
        cahier.preparer_pages(self.chantier, '01.19.1a.01', self.pdf, '2-3')
        avant = self.pr.read_bytes()
        backups = []
        cahier.assurer_liens(self.chantier, self.pr, lambda p: backups.append(p.read_bytes()))
        self.assertEqual(backups, [avant])
        with ZipFile(self.pr) as z:
            for nom, value in contenu.items():
                if nom not in {'xl/workbook.xml', 'xl/_rels/workbook.xml.rels', '[Content_Types].xml',
                               'xl/worksheets/sheet1.xml', 'xl/worksheets/_rels/sheet1.xml.rels'}:
                    self.assertEqual(z.read(nom), value, nom)
        wb = load_workbook(self.pr)
        self.assertEqual(wb['Chiffrage']['C4'].value, 'Suite de désignation')
        self.assertEqual(wb['Chiffrage']['C5'].hyperlink.location, "'CSC 01.19.1a.01'!A1")
        self.assertEqual(wb['CSC 01.19.1a.01']['A2'].hyperlink.location, "'Chiffrage'!B3")
        self.assertEqual(len(wb['CSC 01.19.1a.01']._images), 2)
        wb.close()
        wb = load_workbook(self.pr, data_only=True)
        self.assertEqual(wb['Chiffrage']['F30'].value, 123)
        wb.close()
        resultat = self.pr.read_bytes()
        self.assertEqual(cahier.assurer_liens(self.chantier, self.pr), 0)
        self.assertEqual(self.pr.read_bytes(), resultat)

    def test_verrou_et_echec_sauvegarde_empechent_aussi_integration(self):
        cahier.preparer_pages(self.chantier, '01.19.1a.01', self.pdf, '2')
        avant = self.pr.read_bytes()
        verrou = self.pr.with_name('~$' + self.pr.name)
        verrou.touch()
        with self.assertRaises(PermissionError):
            cahier.assurer_liens(self.chantier, self.pr)
        self.assertEqual(self.pr.read_bytes(), avant)
        from main import open_pr
        with patch('main.ouvrir_chemin') as ouvrir:
            open_pr(self.chantier)
            ouvrir.assert_called_once_with(self.pr)
        self.assertEqual(self.pr.read_bytes(), avant)
        verrou.unlink()
        with self.assertRaises(PermissionError):
            cahier.assurer_liens(self.chantier, self.pr, lambda _: (_ for _ in ()).throw(PermissionError('sauvegarde')))
        self.assertEqual(self.pr.read_bytes(), avant)

    def test_changement_pdf_ne_remplace_pas_choix_utilisateur_par_detection(self):
        cahier.preparer_pages(self.chantier, '01.19.1a.01', self.pdf, '2')
        creer_pdf(self.pdf, (150, 250))
        with patch.object(cahier, '_pages_automatiques', return_value={'01.19.1a.01': '1'}):
            cahier.assurer_liens(self.chantier, self.pr)
        wb = load_workbook(self.pr)
        self.assertNotIn('CSC 01.19.1a.01', wb.sheetnames)
        wb.close()

    def test_nouveau_chantier_ne_recoit_pas_le_cahier_integre_du_modele(self):
        from main import _creer_documents_chantier
        cahier.preparer_pages(self.chantier, '01.19.1a.01', self.pdf, '2-3')
        cahier.assurer_liens(self.chantier, self.pr)
        modele = self.base / 'Modele'
        modele.mkdir()
        source = modele / 'prix_de_revient model.xlsx'
        shutil.copy2(self.pr, source)
        shutil.copy2(self.pdf, modele / 'Cahier_technique.pdf')
        avant = source.read_bytes()
        _creer_documents_chantier(self.base, '010_Nouveau', {})
        nouveau = self.base / '010_Nouveau'
        self.assertIsNone(cahier.trouver_cahier(nouveau))
        wb = load_workbook(nouveau / 'data/prix_de_revient.xlsx')
        self.assertFalse(any(n.startswith('CSC ') for n in wb.sheetnames))
        self.assertEqual(wb['Chiffrage']['C4'].value, cahier.LABEL)
        self.assertEqual(liens.resoudre(self.base, wb['Chiffrage']['C4'].hyperlink.target),
                         (nouveau, '01.19.1a.01'))
        wb.close()
        self.assertEqual(source.read_bytes(), avant)
        with ZipFile(nouveau / 'data/prix_de_revient.xlsx') as z:
            self.assertFalse(any(n.startswith('xl/media/cahier_') for n in z.namelist()))


if __name__ == '__main__':
    unittest.main()

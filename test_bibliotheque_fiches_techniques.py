import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from openpyxl import load_workbook

from bibliotheque_fiches_techniques import creer_bibliotheque_classee, copier_fiche_au_chantier
from main import HorizonChantierApp


class BibliothequeFichesTechniques(unittest.TestCase):
    def test_reprise_ancien_nom_sans_source_et_sans_perte(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            ancien = base / "FICHES TECHNIQUES CLASSEES"
            (ancien / "Toiture").mkdir(parents=True)
            (ancien / "Toiture/offre.pdf").write_bytes(b"offre")
            sortie, _, _ = creer_bibliotheque_classee(base)
            self.assertEqual(sortie.name, "Bibliotheque_Technique")
            self.assertEqual((sortie / "Toiture/offre.pdf").read_bytes(), b"offre")
            self.assertNotIn(ancien.name, [p.name for p in base.iterdir()])
            self.assertEqual(creer_bibliotheque_classee(base)[0], sortie)

    def test_reprise_nom_intermediaire(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            ancien = base / "Bibliotheque_technique"
            ancien.mkdir()
            (ancien / "offre.pdf").write_bytes(b"offre")
            sortie, _, _ = creer_bibliotheque_classee(base)
            self.assertEqual(sortie.name, "Bibliotheque_Technique")
            self.assertEqual((sortie / "offre.pdf").read_bytes(), b"offre")
            self.assertNotIn(ancien.name, [p.name for p in base.iterdir()])

    def test_bibliotheque_vierge_et_deux_dossiers_preserves(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            sortie, _, _ = creer_bibliotheque_classee(base)
            self.assertEqual(sortie.name, "Bibliotheque_Technique")
            self.assertTrue(sortie.is_dir())
            (sortie / "offre.pdf").write_bytes(b"nouvelle")
            ancien = base / "FICHES TECHNIQUES CLASSEES"
            ancien.mkdir()
            (ancien / "offre.pdf").write_bytes(b"ancienne")
            with self.assertRaises(ValueError):
                creer_bibliotheque_classee(base)
            self.assertEqual((sortie / "offre.pdf").read_bytes(), b"nouvelle")
            self.assertEqual((ancien / "offre.pdf").read_bytes(), b"ancienne")

    def test_classement_complet_et_copie_choisie(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            source = base / "FICHES TECHNIQUES"
            chaux = source / "Fiches techniques 1/Chaux"
            chaux.mkdir(parents=True)
            (chaux / "Produit A.pdf").write_bytes(b"a")
            (chaux / "Produit B.pdf").write_bytes(b"b")
            copie = source / "FICHES TECHNIQUES 2/Fiches techniques/Chaux"
            copie.mkdir(parents=True)
            (copie / "Produit A copie.pdf").write_bytes(b"a")
            echafaudage = source / "Fiches techniques 1/offre ADRIA"
            echafaudage.mkdir(parents=True)
            (echafaudage / "location échafaudage.pdf").write_bytes(b"c")
            (echafaudage / "image.jpg").write_bytes(b"image")

            classee, categories, documents = creer_bibliotheque_classee(base, alimenter_depuis_source=True)
            self.assertEqual((categories, documents), (2, 3))
            wb = load_workbook(classee / "Index_fiches_techniques.xlsx")
            try:
                self.assertEqual([r[0].value for r in wb["Catégories"].iter_rows(min_row=2)],
                                 ["001", "002"])
                self.assertEqual(wb["Documents"].max_row, 4)
            finally:
                wb.close()
            choix = next(classee.rglob("Produit B.pdf"))
            chantier = base / "Chantiers/Exemple"
            destination, ajoutee = copier_fiche_au_chantier(choix, classee, chantier)
            self.assertTrue(ajoutee)
            self.assertEqual(destination.read_bytes(), b"b")
            self.assertEqual(copier_fiche_au_chantier(choix, classee, chantier),
                             (destination, False))
            self.assertFalse((chantier / "Technique/Produit A.pdf").exists())

    def test_ouverture_ne_recherche_pas_dans_les_anciens_dossiers_source(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            source = base / 'FICHES TECHNIQUES'
            source.mkdir()
            (source / 'fiche.pdf').write_bytes(b'source')
            bibliotheque, categories, documents = creer_bibliotheque_classee(base)
            self.assertEqual((categories, documents), (0, 0))
            self.assertEqual(list(bibliotheque.iterdir()), [])
            self.assertEqual((source / 'fiche.pdf').read_bytes(), b'source')



if __name__ == "__main__":
    unittest.main()

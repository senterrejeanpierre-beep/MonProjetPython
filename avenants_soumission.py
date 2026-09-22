"""Affiche la soumission HT dans Avenants sans recalculer le classeur Excel."""

import os
import re
import shutil
import tempfile
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook


LIBELLE = "Soumission hors TVA"
FEUILLE_XML = "xl/worksheets/sheet1.xml"


def lire_soumission_etat(chemin_etat: Path) -> float:
    """Lit le résultat enregistré dans la synthèse de l'état, zéro inclus."""
    wb = load_workbook(chemin_etat, read_only=True, data_only=True)
    try:
        for ws in wb:
            for row in ws.iter_rows(min_col=16, max_col=min(23, ws.max_column)):
                for cellule in row:
                    if not isinstance(cellule.value, str) or cellule.value.strip().casefold() != "total soumission hors tva":
                        continue
                    for col in range(cellule.column + 1, min(cellule.column + 5, ws.max_column + 1)):
                        valeur = ws.cell(cellule.row, col).value
                        if isinstance(valeur, (int, float)) and not isinstance(valeur, bool):
                            return float(valeur)
                    raise ValueError("Résultat de la soumission absent dans l'état. Enregistrez l'état dans Excel.")
        raise ValueError("Soumission hors TVA introuvable dans la synthèse de l'état.")
    finally:
        wb.close()


def renseigner_soumission_avenants(chemin_avenants: Path, montant: float | None) -> bool:
    """Change D1 et E1 sans toucher aux formules ni à leurs valeurs calculées."""
    chemin_avenants = Path(chemin_avenants)
    if chemin_avenants.with_name("~$" + chemin_avenants.name).exists():
        raise ValueError("Fermez Avenants.xlsx dans Excel avant de l'actualiser.")
    signature = (chemin_avenants.stat().st_size, chemin_avenants.stat().st_mtime_ns)
    with ZipFile(chemin_avenants) as source:
        xml = source.read(FEUILLE_XML).decode("utf-8")
        cellule_d1 = r'<c r="D1"[^>]*?/>|<c r="D1"[^>]*?>.*?</c>'
        cellule_e1 = r'<c r="E1"[^>]*?/>|<c r="E1"[^>]*?>.*?</c>'
        if not re.search(cellule_d1, xml) or not re.search(cellule_e1, xml):
            raise ValueError("Cellules D1/E1 absentes du classeur Avenants.")
        etiquette = f'<c r="D1" s="9" t="inlineStr"><is><t>{LIBELLE}</t></is></c>'
        valeur = (f'<c r="E1" s="14"><v>{montant:.12g}</v></c>'
                  if montant is not None else '<c r="E1" s="14"/>')
        nouveau = re.sub(cellule_d1, etiquette, xml, count=1)
        nouveau = re.sub(cellule_e1, valeur, nouveau, count=1)
        if nouveau == xml:
            return False
        fd, nom_temp = tempfile.mkstemp(prefix=".avenants_", suffix=".xlsx", dir=chemin_avenants.parent)
        os.close(fd)
        temporaire = Path(nom_temp)
        try:
            with ZipFile(temporaire, "w") as cible:
                for entree in source.infolist():
                    cible.writestr(entree, nouveau.encode("utf-8") if entree.filename == FEUILLE_XML else source.read(entree.filename))
            shutil.copymode(chemin_avenants, temporaire)
            if (chemin_avenants.stat().st_size, chemin_avenants.stat().st_mtime_ns) != signature:
                raise ValueError("Avenants.xlsx a changé pendant l'actualisation.")
            os.replace(temporaire, chemin_avenants)
        finally:
            temporaire.unlink(missing_ok=True)
    return True

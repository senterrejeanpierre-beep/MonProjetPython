"""Les actions existantes doivent viser uniquement l'état choisi."""
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unicodedata
import unittest
from unittest.mock import Mock, patch

from openpyxl import load_workbook

from main import (
    HorizonChantierApp, _choisir_etat_avancement, _etats_avancement_disponibles,
    _fermer_classeur, _selectionner_etat,
)
from test_stabilite_ciblee import etat


class EtatSupplement(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.dossier = Path(self.tmp.name) / '001_Chantier'
        self.dossier.mkdir()
        self.principal = self.dossier / 'Etat_avancement_01.xlsm'
        self.supplement = (self.dossier / 'Supplément aux travaux principaux'
                           / 'Etat_avancement_Supplement.xlsm')
        self.supplement.parent.mkdir()
        for chemin in (self.principal, self.supplement):
            wb = etat()
            wb.save(chemin)
            _fermer_classeur(wb)
        self.original = self.principal.read_bytes()
        self.app = SimpleNamespace(
            _chemin_chantier_selectionne=lambda: self.dossier.with_suffix('.json'),
            _dossier_chantier_selectionne=lambda: self.dossier,
            _confirmer_cloture_etat=lambda: True,
        )

    def test_disponibilite_et_pilotage_principal_inchange(self):
        self.assertEqual(_etats_avancement_disponibles(self.dossier), [
            ('Travaux principaux', self.principal),
            ('Supplément aux travaux principaux', self.supplement),
        ])
        self.assertEqual(_selectionner_etat(self.dossier), self.principal)

    def test_nom_unicode_mac_et_dossier_deplace(self):
        nom = unicodedata.normalize('NFD', 'Supplément aux travaux principaux')
        self.supplement.parent.rename(self.dossier / nom)
        nouveau = self.dossier.with_name('002_Dossier déplacé')
        self.dossier.rename(nouveau)
        choix = _etats_avancement_disponibles(nouveau)
        self.assertEqual(len(choix), 2)
        self.assertTrue(choix[1][1].is_file())
        self.assertEqual(choix[1][1].parents[1], nouveau)

    def test_sans_supplement_pas_de_fenetre(self):
        self.supplement.unlink()
        with patch('main.tk.Toplevel') as fenetre:
            self.assertEqual(_choisir_etat_avancement(None, self.dossier, 'Calcul état'), self.principal)
        fenetre.assert_not_called()

    def test_calcul_et_cloture_supplement_seulement(self):
        with patch('main._choisir_etat_avancement', return_value=self.supplement), \
                patch('main.messagebox.showerror') as erreur, patch('main.messagebox.showinfo'):
            self.assertTrue(HorizonChantierApp.calcul_etat_avancement(self.app))
            wb = load_workbook(self.supplement, keep_vba=True)
            self.assertEqual(wb.active['R24'].value, 5)
            self.assertEqual(wb.active['T24'].value, 125)
            self.assertEqual(wb.active['W24'].value, 50)
            _fermer_classeur(wb)
            HorizonChantierApp.mise_a_zero_etat(self.app)
        erreur.assert_not_called()
        wb = load_workbook(self.supplement, keep_vba=True)
        for adresse, attendu in {'P24': 5, 'Q24': 0, 'R24': 5, 'T24': 125, 'W24': 0}.items():
            self.assertEqual(wb.active[adresse].value, attendu)
        _fermer_classeur(wb)
        self.assertEqual(self.principal.read_bytes(), self.original)
        self.assertEqual(len(list((self.supplement.parent / 'Sauvegarde').glob('*.xlsm'))), 1)
        self.assertFalse((self.dossier / 'Sauvegarde').exists())

    def test_annuler_ne_modifie_aucun_etat(self):
        avant = self.supplement.read_bytes()
        with patch('main._choisir_etat_avancement', return_value=None):
            self.assertFalse(HorizonChantierApp.calcul_etat_avancement(self.app))
            HorizonChantierApp.mise_a_zero_etat(self.app)
        self.assertEqual(self.supplement.read_bytes(), avant)
        self.assertEqual(self.principal.read_bytes(), self.original)

    def test_ouverture_du_supplement_par_bouton_etat(self):
        self.app._nom_chantier_selectionne = lambda: '001_Chantier'
        self.app._preparer_ouverture_document = Mock(return_value=True)
        self.app._doit_rappeler_pdf_historique = lambda _: False
        with patch('main._choisir_etat_avancement', return_value=self.supplement), \
                patch('main.ouvrir_chemin') as ouvrir:
            HorizonChantierApp._ouvrir_premier_fichier_chantier(self.app, 'Etat_avancement*.xlsm')
        self.app._preparer_ouverture_document.assert_called_once_with(self.supplement)
        ouvrir.assert_called_once_with(self.supplement)


if __name__ == '__main__':
    unittest.main()

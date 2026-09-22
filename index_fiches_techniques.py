"""Index de la bibliothèque partagée FICHES TECHNIQUES."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


NOM_BIBLIOTHEQUE = "FICHES TECHNIQUES"
NOM_INDEX = "Index_fiches_techniques.xlsx"


def _feuille(wb, nom, sous_titre, entetes, largeurs):
    ws = wb.create_sheet(nom)
    derniere = chr(64 + len(entetes))
    ws.merge_cells(f"A1:{derniere}1")
    ws["A1"] = nom.upper()
    ws["A1"].font = Font(name="Aptos", size=17, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="174A72")
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 34
    ws.merge_cells(f"A2:{derniere}2")
    ws["A2"] = sous_titre
    ws["A2"].font = Font(name="Aptos", italic=True, color="536579")
    ws.row_dimensions[2].height = 25
    for colonne, titre in enumerate(entetes, 1):
        cellule = ws.cell(4, colonne, titre)
        cellule.fill = PatternFill("solid", fgColor="DCEAF4")
        cellule.font = Font(name="Aptos", bold=True, color="173A53")
    for colonne, largeur in largeurs.items():
        ws.column_dimensions[colonne].width = largeur
    ws.freeze_panes = "B5"
    return ws


def _finir(ws, colonnes):
    ws.auto_filter.ref = f"A4:{colonnes}{max(ws.max_row, 5)}"
    for ligne in range(5, ws.max_row + 1):
        ws.row_dimensions[ligne].height = 23
        if ligne % 2 == 0:
            for cellule in ws[ligne]:
                cellule.fill = PatternFill("solid", fgColor="F4F8FB")


def _lier(cellule, cible, dossier=False):
    cellule.value = "Ouvrir"
    cellule.hyperlink = cible.as_posix() + ("/" if dossier else "")
    cellule.style = "Hyperlink"


def creer_index_fiches(base: str | Path, sortie: str | Path) -> dict[str, int]:
    """Crée un index sans déplacer ni renommer les documents source."""
    base, sortie = Path(base), Path(sortie)
    bibliotheque = base / NOM_BIBLIOTHEQUE
    if not bibliotheque.is_dir():
        raise FileNotFoundError(f"Bibliothèque de fiches techniques introuvable : {bibliotheque}")
    if sortie.resolve() == bibliotheque.resolve():
        raise ValueError("L'index doit être enregistré hors de la bibliothèque.")
    dossiers = sorted((p for p in bibliotheque.rglob("*") if p.is_dir()),
                      key=lambda p: p.relative_to(bibliotheque).as_posix().casefold())
    categories = [p for p in dossiers if len(p.relative_to(bibliotheque).parts) == 2
                  or (len(p.relative_to(bibliotheque).parts) == 1
                      and not p.name.casefold().startswith("fiches techniques"))]
    racine = sorted((p for p in bibliotheque.iterdir()
                     if p.is_file() and not p.name.startswith(".")), key=lambda p: p.name.casefold())
    wb = Workbook()
    del wb[wb.sheetnames[0]]
    ws_cat = _feuille(wb, "Catégories", "Choisir une catégorie : le lien ouvre son dossier et toutes ses fiches.",
                      ("N°", "Catégorie", "Groupe", "Fichiers", "Accès", "Emplacement relatif"),
                      {"A": 8, "B": 43, "C": 27, "D": 12, "E": 14, "F": 85})
    ws_tous = _feuille(wb, "Tous les dossiers", "Recherche par nom ou filtre ; chaque lien mène au dossier exact.",
                       ("N°", "Dossier", "Dossier parent", "Fichiers", "Accès", "Emplacement relatif"),
                       {"A": 8, "B": 45, "C": 38, "D": 12, "E": 14, "F": 100})
    ws_racine = _feuille(wb, "Fichiers à la racine", "Documents placés directement dans FICHES TECHNIQUES.",
                         ("N°", "Document", "Type", "Accès", "Emplacement relatif"),
                         {"A": 8, "B": 82, "C": 15, "D": 14, "E": 100})

    def fichier_compte(dossier):
        return sum(p.is_file() and not p.name.startswith(".") for p in dossier.rglob("*"))

    for numero, dossier in enumerate(categories, 1):
        ligne = numero + 4
        relatif = dossier.relative_to(base)
        ws_cat.cell(ligne, 1, f"{numero:03d}")
        ws_cat.cell(ligne, 2, dossier.name)
        ws_cat.cell(ligne, 3, dossier.parent.name if dossier.parent != bibliotheque else "Racine")
        ws_cat.cell(ligne, 4, fichier_compte(dossier))
        _lier(ws_cat.cell(ligne, 5), relatif, dossier=True)
        ws_cat.cell(ligne, 6, relatif.as_posix())
    for numero, dossier in enumerate(dossiers, 1):
        ligne = numero + 4
        relatif = dossier.relative_to(base)
        ws_tous.cell(ligne, 1, f"{numero:03d}")
        ws_tous.cell(ligne, 2, dossier.name)
        ws_tous.cell(ligne, 3, dossier.parent.name if dossier.parent != bibliotheque else "Racine")
        ws_tous.cell(ligne, 4, fichier_compte(dossier))
        _lier(ws_tous.cell(ligne, 5), relatif, dossier=True)
        ws_tous.cell(ligne, 6, relatif.as_posix())
    for numero, fichier in enumerate(racine, 1):
        ligne = numero + 4
        relatif = fichier.relative_to(base)
        ws_racine.cell(ligne, 1, f"{numero:03d}")
        ws_racine.cell(ligne, 2, fichier.name)
        ws_racine.cell(ligne, 3, fichier.suffix.lstrip(".").upper())
        _lier(ws_racine.cell(ligne, 4), relatif)
        ws_racine.cell(ligne, 5, relatif.as_posix())
    _finir(ws_cat, "F")
    _finir(ws_tous, "F")
    _finir(ws_racine, "E")
    wb.active = 0
    try:
        wb.save(sortie)
    finally:
        wb.close()
    return {"categories": len(categories), "dossiers": len(dossiers), "fichiers_racine": len(racine)}

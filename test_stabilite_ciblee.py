import unittest
import os
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from main import (
    HorizonChantierApp, _charger_valeurs_excel, _corriger_avenants_externes_etat_prive,
    _fermer_classeur, _montant_a_droite,
    _nombre_metier, _prepare_excel_file_before_write, _publier_pdf_atomique,
    _recalculer_etat, _reporter_quantites_cloture, _sauver_excel_atomique,
    _selectionner_etat, _signature_excel, _verifier_blocs_pr, ecrire_json,
)


def etat(public=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Bordereau"
    if public:
        ws["E8"], ws["H8"] = "Quantités calculées par", "Prix unitaires"
        ws["H9"], ws["I9"], ws["J9"] = "en chiffres", "en lettres", "Somme"
    else:
        ws["J1"], ws["L1"] = "Quantité", "P.U."
    ws["B24"] = "01.02"
    ws["E24" if public else "J24"] = 10
    ws["H24" if public else "L24"] = 25
    ws["P24"], ws["Q24"] = 3, 2
    return wb


class CalculsEtat(unittest.TestCase):
    def test_avenant_contractuel_ne_devient_pas_execute(self):
        ws = Workbook().active
        ws["V55"], ws["W55"] = "Montant de l'avenant état 1", "=SUM(M54)"
        ws["M54"], ws["W56"] = "=SUM(Avenants!D62)", "=SUM(W54:W55)"
        ws["V65"], ws["W65"] = "Montant de l'avenant état 3", "=SUM([1]Feuil1!$D$3)"
        ws["W66"] = "=SUM(W64:W65)"
        ws["V60"], ws["W60"] = "Montant de l'avenant état 2", "=SUM(W35:W59)"
        _corriger_avenants_externes_etat_prive(ws)
        self.assertIsNone(ws["V55"].value)
        self.assertIsNone(ws["W55"].value)
        self.assertEqual(ws["W56"].value, "=W54")
        self.assertIsNone(ws["W65"].value)
        self.assertEqual(ws["W66"].value, "=W64")
        self.assertEqual(ws["W60"].value, "=SUM(W35:W59)")

    def test_synthese_obligatoire_distingue_zero_et_absence(self):
        ws = Workbook().active
        ws["S25"], ws["T25"] = "Total du mois hors TVA", 0
        self.assertEqual(_montant_a_droite(ws, 25, 19, _nombre_metier, requis=True), 0)
        ws["T25"] = None
        with self.assertRaisesRegex(ValueError, "absent ou illisible"):
            _montant_a_droite(ws, 25, 19, _nombre_metier, requis=True)

    def test_pilotage_ignore_seulement_cache_pourcentage_absent(self):
        with TemporaryDirectory() as tmp:
            chemin = Path(tmp) / "etat.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Bordereau"
            ws["S25"] = '=IFERROR(R25/J25,0)'
            ws["T25"] = 42
            wb.save(chemin)
            _fermer_classeur(wb)
            valeurs = _charger_valeurs_excel(chemin, "Bordereau", "État d'avancement",
                                              max_col=23, colonnes_formules_facultatives=("S",))
            self.assertEqual(valeurs["Bordereau"]["T25"].value, 42)
            valeurs.close()
            wb = load_workbook(chemin)
            wb.active["T25"] = "=40+2"
            wb.save(chemin)
            _fermer_classeur(wb)
            with self.assertRaisesRegex(ValueError, "T25"):
                _charger_valeurs_excel(chemin, "Bordereau", "État d'avancement",
                                       max_col=23, colonnes_formules_facultatives=("S",))

    def test_lecture_refuse_classeur_modifie_pendant_controle(self):
        with TemporaryDirectory() as tmp:
            chemin = Path(tmp) / "etat.xlsx"
            wb = Workbook()
            wb.active.title = "Bordereau"
            wb.save(chemin)
            _fermer_classeur(wb)
            with patch("main._signature_excel", side_effect=[(1,), (2,)]):
                with self.assertRaisesRegex(ValueError, "changé pendant la lecture"):
                    _charger_valeurs_excel(chemin, "Bordereau", "État d'avancement")

    def test_poste_ajoute_sans_code_numerique(self):
        ws = etat().active
        ws["B25"], ws["J25"], ws["L25"] = "Ajout A", 4, 30
        ws["P25"], ws["Q25"] = 1, 2
        ws["J26"], ws["L26"], ws["Q26"] = 4, 30, 2
        ws["Q27"], ws["R27"] = "Total soumission hors TVA", 500
        _recalculer_etat(ws)
        self.assertEqual((ws["R25"].value, ws["T25"].value, ws["W25"].value), (3, 90, 60))
        self.assertEqual((ws["R26"].value, ws["T26"].value, ws["W26"].value), (2, 60, 60))
        self.assertEqual(ws["R27"].value, 500)

    def test_public_quantite_et_cumul_sans_double_compte(self):
        ws = etat(True).active
        ws["J24"] = "=E24*H24"
        ws["S30"], ws["T30"] = "Total état cumulé hors TVA", "=SUM(T24:T29)"
        ws["S32"], ws["T32"] = "Total du mois hors TVA", "=W24"
        ws["S34"], ws["T34"] = "Total exécuté au", "=SUM(T30:T32)"
        _recalculer_etat(ws)
        self.assertIn("R24/E24", ws["S24"].value)
        self.assertEqual(ws["R24"].value, 5)
        self.assertEqual(ws["T24"].value, 125)
        self.assertEqual(ws["W24"].value, 50)
        self.assertEqual(ws["T34"].value, "=T30")
        self.assertEqual(ws["J24"].value, "=E24*H24")

    def test_prive_quantite_formule_et_cloture_apres_synthese(self):
        ws = etat().active
        valeurs = Workbook().active
        ws["J24"] = "=SUM(J25:J26)"
        valeurs["J24"] = 10
        ws["AC881"] = "Montant après formule de révision"
        ws["B890"], ws["J890"], ws["L890"] = "02.03", 20, 5
        ws["P890"], ws["Q890"] = 4, 3
        ws["Q900"], ws["R900"] = "Total soumission hors TVA", 350
        _recalculer_etat(ws, valeurs)
        _reporter_quantites_cloture(ws)
        _recalculer_etat(ws, valeurs)
        self.assertEqual(ws["P890"].value, 7)
        self.assertEqual(ws["Q890"].value, 0)
        self.assertEqual(ws["T890"].value, 35)
        self.assertEqual(ws["W890"].value, 0)
        self.assertEqual(ws["P24"].value, 5)
        self.assertIn("R24/J24", ws["S24"].value)
        self.assertEqual(ws["J24"].value, "=SUM(J25:J26)")
        self.assertEqual(ws["Q900"].value, "Total soumission hors TVA")
        self.assertEqual(ws["R900"].value, 350)

    def test_formule_sans_resultat_refusee(self):
        ws = etat().active
        ws["J24"] = "=SUM(J25:J26)"
        with self.assertRaisesRegex(ValueError, "Résultat Excel absent"):
            _recalculer_etat(ws)

    def test_nombres_metier_et_erreurs(self):
        self.assertEqual(_nombre_metier("1\u202f234,50 €"), 1234.5)
        self.assertEqual(_nombre_metier(" −12,5 "), -12.5)
        self.assertEqual(_nombre_metier("\xa0 - \xa0"), 0)
        for valeur in ("#VALUE!", "=A1", "nan", "inf"):
            with self.assertRaises(ValueError):
                _nombre_metier(valeur)


class FichiersMetier(unittest.TestCase):
    def test_publication_pdf_preserve_ancien_en_cas_echec(self):
        destination = Path(self.base) / "pilotage.pdf"
        destination.write_bytes(b"ancien rapport")

        class Document:
            def build(self, _elements):
                Path(self.filename).write_bytes(b"rapport incomplet")
                raise OSError("génération interrompue")

        with self.assertRaisesRegex(OSError, "interrompue"):
            _publier_pdf_atomique(Document(), [], destination)
        self.assertEqual(destination.read_bytes(), b"ancien rapport")
        self.assertFalse(list(self.base.glob(".horizon_pdf_*")))

    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.base = Path(self.tmp.name)
        self.dossier = self.base / "001_Test"
        (self.dossier / "data").mkdir(parents=True)
        self.app = SimpleNamespace(_dossier_chantier_selectionne=lambda: self.dossier)
        self.patch_base = patch("main.dossier_chantiers", return_value=self.base)
        self.patch_base.start()
        self.addCleanup(self.patch_base.stop)

    def creer_pr(self, prix=125):
        p = self.dossier / "data/prix_de_revient.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Chiffrage"
        ws["B3"], ws["F30"], ws["F32"] = "1.2", prix, 151.25
        ws["F30"].number_format = "0.00 €"
        wb.save(p)
        _fermer_classeur(wb)
        return p

    def creer_etat(self, public=False):
        p = self.dossier / ("Etat_avancement_Public.xlsm" if public else "Etat_avancement_Privé.xlsm")
        wb = etat(public)
        wb.save(p)
        _fermer_classeur(wb)
        return p

    def test_choix_public_prive_unicode_et_ambiguite(self):
        self.creer_etat(True)
        prive = self.creer_etat(False)
        decompose = prive.with_name("Etat_avancement_Prive\u0301.xlsm")
        prive.rename(decompose)
        with self.assertRaisesRegex(ValueError, "Plusieurs états"):
            _selectionner_etat(self.dossier)
        ecrire_json(self.base / "001_Test.json", {"type_etat": "Privé"})
        self.assertEqual(_selectionner_etat(self.dossier), decompose)

    def test_import_pr_ht_preserve_source_et_autre_etat(self):
        pr = self.creer_pr()
        prive = self.creer_etat(False)
        public = self.creer_etat(True)
        ecrire_json(self.base / "001_Test.json", {"type_etat": "Public"})
        pr_avant, prive_avant = pr.read_bytes(), prive.read_bytes()
        count = HorizonChantierApp.importer_pr_dans_etat(self.app)
        self.assertEqual(count, 1)
        self.assertEqual(pr.read_bytes(), pr_avant)
        self.assertEqual(prive.read_bytes(), prive_avant)
        wb = load_workbook(public, keep_vba=True)
        self.assertEqual(wb["Bordereau"]["H24"].value, 125)  # HT, pas F32 TTC.
        self.assertEqual(wb["Bordereau"]["E24"].value, 10)
        _fermer_classeur(wb)
        self.assertTrue(list((self.dossier / "Historique_Sauvegardes").glob("*.xlsm")))

    def test_import_prive_colonne_l(self):
        self.creer_pr(37.5)
        p = self.creer_etat()
        self.assertEqual(HorizonChantierApp.importer_pr_dans_etat(self.app), 1)
        wb = load_workbook(p, keep_vba=True)
        self.assertEqual(wb["Bordereau"]["L24"].value, 37.5)
        self.assertIsNone(wb["Bordereau"]["H24"].value)
        _fermer_classeur(wb)

    def test_pr_ouvert_en_lecture_et_priorite_historique_des_doublons(self):
        pr = self.creer_pr(25)
        wb = load_workbook(pr)
        wb.active["B45"], wb.active["F72"] = "01.02", 30
        wb.save(pr)
        _fermer_classeur(wb)
        pr.with_name("~$" + pr.name).touch()
        p = self.creer_etat()
        avant = pr.read_bytes()
        self.assertEqual(HorizonChantierApp.importer_pr_dans_etat(self.app), 1)
        wb = load_workbook(p, keep_vba=True)
        self.assertEqual(wb.active["L24"].value, 30)
        _fermer_classeur(wb)
        self.assertEqual(self.app._dernier_import_pr_doublons, ["1.2"])
        self.assertEqual(pr.read_bytes(), avant)

    def test_pas_de_substitution_du_seul_etat_prive_au_public(self):
        self.creer_etat(False)
        ecrire_json(self.base / "001_Test.json", {"type_etat": "Public"})
        with self.assertRaisesRegex(ValueError, "ne correspond pas"):
            _selectionner_etat(self.dossier)

    def test_cloture_complete_sur_fichier_isole(self):
        p = self.creer_etat()
        self.app._confirmer_cloture_etat = lambda: True
        self.app._chemin_chantier_selectionne = lambda: self.base / "001_Test.json"
        with patch("main.messagebox.showerror") as erreur, patch("main.messagebox.showinfo"):
            HorizonChantierApp.mise_a_zero_etat(self.app)
        erreur.assert_not_called()
        wb = load_workbook(p, keep_vba=True)
        self.assertEqual(wb.active["P24"].value, 5)
        self.assertEqual(wb.active["Q24"].value, 0)
        self.assertEqual(wb.active["T24"].value, 125)
        self.assertEqual(wb.active["W24"].value, 0)
        _fermer_classeur(wb)
        self.assertEqual(len(list((self.dossier / "Sauvegarde").glob("*.xlsm"))), 1)

    def test_pr_sans_cache_aucune_ecriture(self):
        pr = self.creer_pr("=100+25")
        p = self.creer_etat()
        avant = {f: f.read_bytes() for f in (pr, p)}
        with self.assertRaisesRegex(ValueError, "indisponible"):
            HorizonChantierApp.importer_pr_dans_etat(self.app)
        for f, contenu in avant.items():
            self.assertEqual(f.read_bytes(), contenu)
        self.assertFalse((self.dossier / "Historique_Sauvegardes").exists())

    def test_aucune_correspondance_ne_sauve_pas(self):
        self.creer_pr()
        p = self.creer_etat()
        wb = load_workbook(p, keep_vba=True)
        wb.active["B24"] = "9.99"
        wb.save(p)
        _fermer_classeur(wb)
        avant = p.read_bytes()
        self.assertEqual(HorizonChantierApp.importer_pr_dans_etat(self.app), 0)
        self.assertEqual(p.read_bytes(), avant)

    def test_echec_enregistrement_preserve_original(self):
        p = self.creer_etat()
        signature = _signature_excel(p)
        avant = p.read_bytes()
        wb = etat()
        def echec(destination):
            Path(destination).write_bytes(b"ecriture partielle")
            raise OSError("disque indisponible")
        with patch.object(wb, "save", side_effect=echec):
            with self.assertRaises(OSError):
                _sauver_excel_atomique(wb, p, signature)
        self.assertEqual(p.read_bytes(), avant)
        self.assertFalse(list(self.dossier.glob(".horizon_*")))

    def test_preparation_classeur_sauvegarde_avant_remplacement(self):
        p = self.dossier / "document.xlsx"
        wb = Workbook()
        wb.active["A1"] = "=1+2"
        wb.save(p)
        _fermer_classeur(wb)
        original = p.read_bytes()
        _prepare_excel_file_before_write(p)
        sauvegardes = list((self.dossier / "Historique_Sauvegardes").glob("*.xlsx"))
        self.assertEqual(len(sauvegardes), 1)
        self.assertEqual(sauvegardes[0].read_bytes(), original)
        wb = load_workbook(p)
        self.assertEqual(wb.active["A1"].value, "=1+2")
        self.assertTrue(wb.active["A1"].protection.locked)
        _fermer_classeur(wb)

    def test_modification_concurrente_et_verrou_excel(self):
        p = self.creer_etat()
        signature = _signature_excel(p)
        p.write_bytes(b"nouvelle version concurrente")
        with self.assertRaisesRegex(ValueError, "modifié"):
            _sauver_excel_atomique(etat(), p, signature)
        self.assertEqual(p.read_bytes(), b"nouvelle version concurrente")
        p.with_name("~$" + p.name).touch()
        with self.assertRaisesRegex(ValueError, "Fermez"):
            _signature_excel(p)

    def test_signature_detecte_changement_meme_taille_et_date(self):
        p = self.dossier / "document.xlsx"
        p.write_bytes(b"version A")
        signature = _signature_excel(p)
        stat = p.stat()
        p.write_bytes(b"version B")
        os.utime(p, ns=(stat.st_atime_ns, stat.st_mtime_ns))
        self.assertNotEqual(_signature_excel(p), signature)

    def test_verification_deuxieme_bloc_et_commentaire(self):
        ws = Workbook().active
        ws["B3"], ws["F30"] = "01.02", 125
        ws["B45"], ws["F72"] = "01.03", 150
        ws["C16"] = "Note sans heures ni coût"
        ws["I45"], ws["J45"], ws["H45"] = 10, 0.1, 2
        ws["K45"], ws["L45"] = 11, 21  # Attendu 22.
        nb, erreurs = _verifier_blocs_pr(ws)
        self.assertEqual(nb, 2)
        self.assertEqual(len(erreurs), 1)
        self.assertIn("Ligne 45", erreurs[0])

    def test_pr_prix_matiere_saisi_controle_le_total_reel(self):
        valeurs = Workbook().active
        formules = Workbook().active
        for ws in (valeurs, formules):
            ws["B3"], ws["B30"] = "01.02", "RÉCAPITULATIF POSTE"
            ws["I3"], ws["J3"], ws["H3"] = 1, 0.03, 2
            ws["K3"] = 650
        valeurs["L3"] = 1300
        formules["L3"] = "=H3*K3"
        nb, erreurs = _verifier_blocs_pr(valeurs, formules)
        self.assertEqual((nb, erreurs), (1, []))
        valeurs["L3"] = 1299
        _, erreurs = _verifier_blocs_pr(valeurs, formules)
        self.assertEqual(erreurs, ["Ligne 3 : total matière à vérifier"])
        formules["K3"] = "=I3*(1+J3)"
        _, erreurs = _verifier_blocs_pr(valeurs, formules)
        self.assertIn("Ligne 3 : PU matière à vérifier", erreurs)

    def test_echec_fiche_json_preserve_original(self):
        p = self.base / "fiche.json"
        ecrire_json(p, {"client": "Original"})
        avant = p.read_bytes()
        with patch("main.os.replace", side_effect=OSError("interruption")):
            with self.assertRaises(OSError):
                ecrire_json(p, {"client": "Modifié"})
        self.assertEqual(p.read_bytes(), avant)
        self.assertFalse(list(self.base.glob(".fiche_*")))


if __name__ == "__main__":
    unittest.main()

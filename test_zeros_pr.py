from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import Mock
from xml.etree import ElementTree as ET
from zipfile import ZipFile
from openpyxl import Workbook, load_workbook

from zeros_pr import completer_zeros_pr, NS


def creer_pr(path):
    wb=Workbook();s=wb.active;s.title='Chiffrage'
    s['B1']='MATIÈRE'
    entetes=('Article : code','Désignation','Désignation matière','Qté/pcs bordereau',
             'Unité bordereau','Quantité','Qté totale','PU achat','% pertes','PU corrigé',
             'Total matière','Transport (total)','PU transport','Total transport','Total revient')
    for col,v in enumerate(entetes,2):s.cell(2,col,v)
    s['B3']='01.1';s['C3']='Maçonnerie';s['D3']='Briques';s['E3']=10;s['F3']='m²'
    s['I3']=123;s['P3']='=I3*E3';s['P12']='=SUM(P3:P11)'
    s['O13']='Sous-total matière (revient)';s['P13']='=SUM(P12/E3)'
    s['B14']='MAIN-D’ŒUVRE'
    for col,v in {2:'Métier/Équipe',3:'Nb pers.',4:'Désignation matière',5:'Mode (h/u ou u/h)',
                  6:'Qté poste',7:'Jours',8:'Heures calc.',9:'Heures prévues',10:'Coût horaire',
                  11:'Coût MO',12:'Transport (total)',13:'PU transport',14:'Total transport',16:'Total revient'}.items():
        s.cell(15,col,v)
    s['I16']='=H16*C16';s['G17']='=I16/C16/8';s['F17']='Total jours P'
    s['P25']='=SUM(P16:P24)';s['O26']='Sous-total MO (revient)';s['P26']='=SUM(P25/E3)'
    for row,label in ((28,'PR matière (revient)'),(29,'PR MO (revient)'),(30,'PR total')):s.cell(row,2,label)
    s['E28']='IC %';s['E29']='Marge vente %';s['E30']='PV HT'
    s['I28']='IC €';s['I29']='Marge €';s['E31']='TVA %';s['E32']='PV TTC'
    s['B38']='SUIVI EN EXÉCUTION (réalisé vs prévu)'
    for col,v in {2:'Date',3:'Commentaire / Situation',4:'Qté réalisée',5:'% avancement',
                  6:'Heures réelles',7:'Prix u MO',8:'Coût réel MO',9:'Matière réelle',10:'Coût réel matière',
                  11:'Coût réel total',12:'PR total (réf)',13:'Écart coût',14:'Observations',15:'PV HT (réf)',16:'Écart marge%'}.items():s.cell(39,col,v)
    s['E40']='=D40/$E$3';s['F40']='=G17/8/D40';s['P40']='=(K40/L40)-E40'
    s['G41']='Gain/Perte MO';s['K41']='Gain/Perte Total';s['Q44']='Hors périmètre'
    s.merge_cells('G5:H5')
    wb.create_sheet('Autre')['A1']='Ne pas modifier'
    wb.save(path);wb.close()


class ZerosPRTests(unittest.TestCase):
    def test_zeros_preservent_saisies_structure_styles_et_formules_utiles(self):
        with TemporaryDirectory() as d:
            p=Path(d)/'pr.xlsx';creer_pr(p)
            with ZipFile(p) as z:avant={n:z.read(n) for n in z.namelist()}
            sauvegarde=Mock()
            self.assertGreater(completer_zeros_pr(p,sauvegarde),0)
            sauvegarde.assert_called_once_with(p)
            w=load_workbook(p);s=w['Chiffrage']
            for ref in ('C16','H16','J16','I4','G5','C28','F28','F29','J29','D40','H41'):
                self.assertEqual(s[ref].value,0,ref)
            for ref in ('B4','C4','D4','F4','B17','C17','B40','C40','N40','H5','E12'):
                self.assertIsNone(s[ref].value,ref)
            self.assertEqual(s['I3'].value,123);self.assertEqual(s['E3'].value,10)
            self.assertEqual(s['P3'].value,'=I3*E3')
            self.assertEqual(s['G17'].value,'=IF(OR(C16=0,C16=""),0,I16/C16/8)')
            self.assertEqual(s['P40'].value,'=IF(OR(L40=0,L40=""),0,(K40/L40)-E40)')
            w.close()
            with ZipFile(p) as z:
                self.assertIsNone(z.testzip())
                for n,data in avant.items():
                    if n not in ('xl/worksheets/sheet1.xml','xl/workbook.xml'):
                        self.assertEqual(z.read(n),data,n)
            b=p.read_bytes();self.assertEqual(completer_zeros_pr(p),0);self.assertEqual(p.read_bytes(),b)

    def test_classeur_ouvert_et_modification_concurrente_proteges(self):
        with TemporaryDirectory() as d:
            p=Path(d)/'pr.xlsx';creer_pr(p);avant=p.read_bytes()
            verrou=p.with_name('~$'+p.name);verrou.write_text('ouvert')
            with self.assertRaises(PermissionError):completer_zeros_pr(p)
            self.assertEqual(p.read_bytes(),avant)
            verrou.unlink()
            def modifier(path):path.write_bytes(avant+b' modification externe')
            with self.assertRaises(ValueError):completer_zeros_pr(p,modifier)
            self.assertEqual(p.read_bytes(),avant+b' modification externe')

    def test_ancien_entete_rendement_et_diviseur_saisi_conserves(self):
        with TemporaryDirectory() as d:
            p=Path(d)/'pr.xlsx';creer_pr(p)
            w=load_workbook(p);s=w['Chiffrage']
            s['D15']='Rendement';s['H16']='=SUM(E3/R16)';s['R16']=4
            w.save(p);w.close()
            completer_zeros_pr(p)
            w=load_workbook(p);s=w['Chiffrage']
            self.assertEqual(s['C16'].value,0)
            self.assertEqual(s['R16'].value,4)
            self.assertEqual(s['H16'].value,'=IF(OR(R16=0,R16=""),0,SUM(E3/R16))')
            self.assertIsNone(s['D16'].value)
            w.close()

    def test_division_vide_cache_zero_et_resultats_renseignes_conserves(self):
        with TemporaryDirectory() as d:
            p=Path(d)/'pr.xlsx';creer_pr(p)
            tmp=Path(d)/'cache.xlsx'
            with ZipFile(p) as src, ZipFile(tmp,'w') as dst:
                for info in src.infolist():
                    data=src.read(info.filename)
                    if info.filename=='xl/worksheets/sheet1.xml':
                        root=ET.fromstring(data)
                        for ref,val,typ in (('G17','#DIV/0!','e'),('P13','123','n'),('I16','0','n')):
                            cell=root.find(f'.//{{{NS}}}c[@r="{ref}"]')
                            cell.set('t',typ);cell.find('{'+NS+'}v').text=val
                        data=ET.tostring(root)
                    dst.writestr(info,data)
            completer_zeros_pr(tmp)
            w=load_workbook(tmp,data_only=True)
            self.assertEqual(w['Chiffrage']['G17'].value,0)
            self.assertEqual(w['Chiffrage']['P13'].value,123)
            self.assertEqual(w['Chiffrage']['I16'].value,0)
            w.close()

    def test_creation_future_recoit_les_zeros(self):
        from main import _creer_documents_chantier
        with TemporaryDirectory() as d:
            base=Path(d);modele=base/'Modele';modele.mkdir()
            creer_pr(modele/'prix_de_revient model.xlsx')
            _creer_documents_chantier(base,'Futur',{})
            w=load_workbook(base/'Futur/data/prix_de_revient.xlsx')
            self.assertEqual(w['Chiffrage']['C16'].value,0)
            self.assertEqual(w['Chiffrage']['I4'].value,0)
            self.assertTrue(w['Chiffrage']['G17'].value.startswith('=IF(OR('))
            w.close()


if __name__=='__main__':unittest.main()

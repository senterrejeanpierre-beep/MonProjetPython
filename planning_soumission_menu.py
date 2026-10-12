"""Saisie des périodes par liste déroulante, sur une seule ligne par poste."""
from copy import copy
from io import BytesIO
from pathlib import Path
import re
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl.formatting.rule import FormulaRule
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.styles import PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.filters import AutoFilter

from planning_soumission_periodes import lire_plages, remplissage_automatique

STOCKAGE = 160
META = 159
PREMIER_JOUR = 7
NB_JOURS = 140
BLEU_DOUX = '9DBAD5'


def fermer_planning(wb):
    wb.close()
    if wb.vba_archive is not None:
        wb.vba_archive.close()


def est_menu(ws):
    return (ws['D8'].value, ws['E8'].value, ws['F8'].value) == ('Période', 'Du jour', 'Au jour')


def incorporer_macros(wb):
    fichier = Path(__file__).parent / 'ressources' / 'planning_soumission' / 'vbaProject.bin'
    memoire = BytesIO()
    with ZipFile(memoire, 'w', ZIP_DEFLATED) as archive:
        archive.writestr('xl/vbaProject.bin', fichier.read_bytes())
        archive.writestr('_rels/.rels',
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>')
        archive.writestr('[Content_Types].xml',
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Override PartName="/xl/vbaProject.bin" ContentType="application/vnd.ms-office.vbaProject"/>'
            '</Types>')
    memoire.seek(0)
    if wb.vba_archive is not None:
        wb.vba_archive.close()
    wb.vba_archive = ZipFile(memoire)
    wb.code_name = 'ThisWorkbook'
    wb['Planning soumission'].sheet_properties.codeName = 'Sheet1'


def _texte_libre(label, periode, ancien):
    if ancien and isinstance(label, str):
        for separateur in (' – ', ' - '):
            prefixe = str(periode) + separateur
            if label.startswith(prefixe):
                return label[len(prefixe):]
    return label


def lire_textes_menu(ws, convertir=True):
    nombre, fin = int(ws.cell(1, META).value), int(ws.cell(2, META).value)
    premiere = STOCKAGE+2*nombre
    ancien = convertir and ws.cell(4, META).value not in ('HorizonLegendes5', 'HorizonLegendes6')
    return {str(ws.cell(r, 1).value): [_texte_libre(ws.cell(r, premiere+n).value, n+1, ancien)
            if ws.cell(r, premiere+n).value is not None else n+1
            for n in range(nombre)] for r in range(9, fin+1)}


def selection_menu(ws, ligne):
    valeur = ws.cell(ligne, 4).value
    nombre = int(ws.cell(1, META).value)
    if isinstance(valeur, (int, float)) and 1 <= valeur <= nombre and int(valeur) == valeur:
        return int(valeur)
    return int(ws.cell(ligne, META).value)


def assurer_textes_periodes(ws):
    """Autorise le texte directement dans le sélecteur, sans nouveau champ visible."""
    if ws.cell(3, META).value == 'HorizonPeriodesTexte1':
        return assurer_legendes(ws)
    nombre, fin = int(ws.cell(1, META).value), int(ws.cell(2, META).value)
    premiere = STOCKAGE+2*nombre
    textes = lire_textes_menu(ws, convertir=False)
    ws.data_validations.dataValidation = [v for v in ws.data_validations.dataValidation
        if not (v.type == 'list' and any(p.min_col <= 4 <= p.max_col and p.max_row >= 9
                                        for p in v.sqref.ranges))]
    for r in range(9, fin+1):
        labels = textes[str(ws.cell(r, 1).value)]
        for n, label in enumerate(labels):
            c = ws.cell(r, premiere+n, label)
            if isinstance(label, str):
                c.data_type = 's'
        selection = selection_menu(ws, r)
        ws.cell(r, META, selection)
        ws.cell(r, 4).value = labels[selection-1]
        formule = f'${get_column_letter(premiere)}${r}:${get_column_letter(premiere+nombre-1)}${r}'
        v = DataValidation(type='list', formula1=formule, allow_blank=True)
        v.showDropDown = False
        # Une saisie qui n'est pas encore dans la liste renomme la période choisie.
        v.showErrorMessage = False
        ws.add_data_validation(v)
        v.add(f'D{r}')
    for c in range(premiere, premiere+nombre):
        ws.column_dimensions[get_column_letter(c)].hidden = True
    ws.column_dimensions['D'].width = 26
    ws['A6'] = 'Choisissez une période, ou écrivez son texte, puis saisissez du jour / au jour.'
    ws.cell(3, META, 'HorizonPeriodesTexte1')
    assurer_legendes(ws)
    return True


def assurer_legendes(ws):
    couleur_modifiee = False
    remplissage = remplissage_automatique(BLEU_DOUX)
    plage = f'G9:EP{int(ws.cell(2, META).value)}'
    for zone, regles in ws.conditional_formatting._cf_rules.items():
        if str(zone.sqref) == plage:
            for regle in regles:
                if regle.dxf is not None and regle.dxf.fill is not None and regle.dxf.fill != remplissage:
                    regle.dxf.fill = copy(remplissage)
                    couleur_modifiee = True
    if ws.cell(4, META).value == 'HorizonLegendes5':
        return couleur_modifiee
    from planning_soumission_legendes import texte_legende
    nombre, fin = int(ws.cell(1, META).value), int(ws.cell(2, META).value)
    ancien = ws.cell(4, META).value not in ('HorizonLegendes5', 'HorizonLegendes6')
    for r in range(9, fin+1):
        for n in range(nombre):
            cellule = ws.cell(r, STOCKAGE+2*nombre+n)
            cellule.value = _texte_libre(cellule.value, n+1, ancien)
            cellule.number_format = '@'
            if isinstance(cellule.value, str):
                cellule.data_type = 's'
            ws.cell(r, STOCKAGE+3*nombre+n, texte_legende(cellule.value))
        ws.cell(r, 4).number_format = '@'
        ws.cell(r, 4).value = ws.cell(r, STOCKAGE+2*nombre+selection_menu(ws,r)-1).value
        if isinstance(ws.cell(r, 4).value, str):
            ws.cell(r, 4).data_type = 's'
    for c in range(STOCKAGE+3*nombre, STOCKAGE+4*nombre):
        ws.column_dimensions[get_column_letter(c)].hidden = True
    # Une bordure de cellule reste exactement sous la ligne, même avec le zoom Excel.
    plages = []
    for n in range(nombre):
        d = f'${get_column_letter(STOCKAGE+2*n)}9'
        f = f'${get_column_letter(STOCKAGE+2*n+1)}9'
        plages.append(f'IFERROR(AND(ISNUMBER({d}),ISNUMBER({f}),{d}>=1,'
                      f'{f}>{d}-1,ROUNDUP({d},0)<{f}+1,{d}<=$G$4+139,'
                      f'{f}+1>$G$4,COLUMN(A9)-7+$G$4<{f}+1),FALSE)')
    formule = 'OR('+','.join(plages)+')'
    if not any(regle.formula == [formule]
               for regles in ws.conditional_formatting._cf_rules.values() for regle in regles):
        ws.conditional_formatting.add(f'A9:EP{fin}', FormulaRule(
            formula=[formule], border=Border(bottom=Side(style='thin', color='163A59'))))
    for validation in ws.data_validations.dataValidation:
        if validation.type == 'list' and any(p.min_col <= 4 <= p.max_col and p.max_row >= 9
                                            for p in validation.sqref.ranges):
            validation.showErrorMessage = False
            validation.showDropDown = False
    ws.cell(4, META, 'HorizonLegendes5')
    ws['A6'] = 'Choisissez la période, tapez votre texte dans la même cellule, puis Entrée : il apparaît dans la barre.'
    incorporer_macros(ws.parent)
    return True


def restaurer_textes_menu(ws, textes):
    from planning_soumission_legendes import texte_legende
    nombre, fin = int(ws.cell(1, META).value), int(ws.cell(2, META).value)
    premiere = STOCKAGE+2*nombre
    for r in range(9, fin+1):
        for n, label in enumerate(textes.get(str(ws.cell(r, 1).value), [])):
            c = ws.cell(r, premiere+n, label)
            if isinstance(label, str):
                c.data_type = 's'
            ws.cell(r, STOCKAGE+3*nombre+n, texte_legende(label))
        ws.cell(r, 4).value = ws.cell(r, premiere+selection_menu(ws, r)-1).value
        if isinstance(ws.cell(r, 4).value, str):
            ws.cell(r, 4).data_type = 's'


def _deplacer_formule(valeur, colonnes):
    if not isinstance(valeur, str) or not valeur.startswith('='):
        return valeur
    morceaux = []
    for t in Tokenizer(valeur).items:
        if t.type == 'OPERAND' and t.subtype == 'RANGE' and '!' not in t.value:
            def deplacer(m):
                absolu, lettre, ligne = m.groups()
                return absolu + colonnes.get(lettre, lettre) + ligne
            morceaux.append(re.sub(r'(\$?)([A-Z]+)(\$?\d+)', deplacer, t.value))
        else:
            morceaux.append(t.value)
    return '=' + ''.join(morceaux)


def lire_plages_menu(ws):
    nombre = int(ws.cell(1, META).value)
    fin = int(ws.cell(2, META).value)
    colonnes = {get_column_letter(STOCKAGE+i): get_column_letter(4+i) for i in range(2*nombre)}
    return {str(ws.cell(r, 1).value): [
        (f'{get_column_letter(4+i)}{r}', _deplacer_formule(ws.cell(r, STOCKAGE+i).value, colonnes))
        for i in range(2*nombre)] for r in range(9, fin+1)}


def installer_menu(ws):
    """Convertit les quatre paires en un sélecteur et deux cellules de saisie."""
    if est_menu(ws):
        return False
    plages = lire_plages(ws)
    if not plages:
        raise ValueError('Aucun poste à reprendre dans le planning.')
    nombre = max(len(p) // 2 for p in plages.values())
    if nombre < 1 or STOCKAGE+2*nombre > 16384:
        raise ValueError('Nombre de périodes invalide.')
    ancien_premier = 4 + 2*nombre
    jour = ws.cell(4, ancien_premier).value or 1
    fin = 8+len(plages)
    lignes = {str(ws.cell(r, 1).value): r for r in range(9, fin+1)}
    grille = {a: [copy(ws.cell(r, c)) for c in range(ancien_premier, ancien_premier+NB_JOURS)]
              for a, r in lignes.items()}
    titres = copy(ws['A8']._style)
    saisie = copy(ws['D9']._style)
    style_depart = copy(ws.cell(4, ancien_premier)._style)
    source = ws.cell(fin+2, 1).value
    style_source = copy(ws.cell(fin+2, 1)._style)
    for fusion in list(ws.merged_cells.ranges):
        if fusion.min_row >= 7 or fusion.min_col >= 7:
            ws.unmerge_cells(str(fusion))
    for row in ws.iter_rows(min_row=7, min_col=4):
        for c in row:
            c.value = None
            c._style = None
            c.comment = None
    for row in ws.iter_rows(min_row=1, max_row=6, min_col=7):
        for c in row:
            c.value = None
            c._style = None
    for cle in list(ws.column_dimensions):
        if cle not in ('A', 'B', 'C'):
            del ws.column_dimensions[cle]
    ws.conditional_formatting._cf_rules.clear()
    ws.data_validations.dataValidation.clear()
    ws.auto_filter = AutoFilter()
    ws.sheet_properties.filterMode = False
    ws['A6'] = 'Choisissez la période, puis saisissez du jour / au jour. Les couleurs suivent automatiquement.'
    for c, titre, largeur in ((4, 'Période', 12), (5, 'Du jour', 14), (6, 'Au jour', 16)):
        ws.cell(8, c, titre)._style = copy(titres)
        ws.column_dimensions[get_column_letter(c)].width = largeur
    validation = DataValidation(type='list', formula1='"'+','.join(str(n) for n in range(1, nombre+1))+'"', allow_blank=False)
    validation.showDropDown = False
    validation.showErrorMessage = True
    validation.error = 'Choisissez une période dans la liste.'
    ws.add_data_validation(validation)
    validation.add(f'D9:D{fin}')
    jours = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1=0, allow_blank=True)
    jours.showErrorMessage = True
    jours.error = 'Saisissez un numéro de jour.'
    ws.add_data_validation(jours)
    jours.add(f'E9:F{fin}')
    ws.cell(1, META, nombre)
    ws.cell(2, META, fin)
    ws.cell(3, META, 'HorizonPeriodes1')
    mapping = {get_column_letter(4+i): get_column_letter(STOCKAGE+i) for i in range(2*nombre)}
    for a, cellules in plages.items():
        r = lignes[a]
        ws.cell(r, 4, 1)
        ws.cell(r, META, 1)
        for i, (_, valeur) in enumerate(cellules):
            ws.cell(r, STOCKAGE+i).value = _deplacer_formule(valeur, mapping)
        for i in range(2):
            ref = f'{get_column_letter(STOCKAGE+i)}{r}'
            ws.cell(r, 5+i, f'=IF({ref}="","",{ref})')
        for c in range(4, 7):
            ws.cell(r, c)._style = copy(saisie)
            ws.cell(r, c).number_format = '0.##' if c > 4 else '0'
            ws.cell(r, c).fill = PatternFill('solid', fgColor='FFF2CC')
        for i, ancien in enumerate(grille[a], PREMIER_JOUR):
            c = ws.cell(r, i)
            c._style = copy(ancien._style)
            c.value = ancien.value
            c.comment = copy(ancien.comment)
    ws.merge_cells('G3:M3')
    ws['G3'] = 'À PARTIR DU JOUR'
    ws['G3']._style = copy(ws['E3']._style)
    ws.merge_cells('G4:I4')
    ws['G4'] = jour
    ws['G4']._style = style_depart
    v = DataValidation(type='whole', operator='greaterThanOrEqual', formula1=1)
    v.showErrorMessage = True
    ws.add_data_validation(v)
    v.add('G4')
    dernier = PREMIER_JOUR+NB_JOURS-1
    for c in range(PREMIER_JOUR, dernier+1):
        col = get_column_letter(c)
        ws.column_dimensions[col].width = 3.5
        ws.cell(8, c, '=$G$4' if c == PREMIER_JOUR else f'={get_column_letter(c-1)}8+1')._style = copy(titres)
        ws.cell(8, c).number_format = '0'
    controles, fins = [], []
    for n in range(nombre):
        d, f = (get_column_letter(STOCKAGE+2*n+i) for i in range(2))
        controles.extend([
            f'SUMPRODUCT(--(({d}9:{d}{fin}="")<>({f}9:{f}{fin}="")))',
            f'SUMPRODUCT(--({d}9:{d}{fin}<>""),--({d}9:{d}{fin}<1))',
            f'SUMPRODUCT(--({f}9:{f}{fin}<>""),--({f}9:{f}{fin}<{d}9:{d}{fin}-1))'])
        fins.append(f'{f}9:{f}{fin}')
        formule = (f'AND(ISNUMBER(${d}9),ISNUMBER(${f}9),${d}9>=1,'
                   f'${f}9>${d}9-1,G$8>=${d}9,G$8<${f}9+1)')
        ws.conditional_formatting.add(f'G9:{get_column_letter(dernier)}{fin}',
            FormulaRule(formula=[formule], fill=remplissage_automatique(BLEU_DOUX)))
    ws['F4'] = f'=IF({"+".join(controles)}>0,"À compléter",MAX({",".join(fins)}))'
    # La mémoire reste dans le même document, hors de l'affichage et de l'impression.
    for c in range(dernier+1, STOCKAGE+2*nombre):
        ws.column_dimensions[get_column_letter(c)].hidden = True
    if source:
        ws.merge_cells(start_row=fin+2, start_column=1, end_row=fin+2, end_column=6)
        ws.cell(fin+2, 1, source)._style = style_source
    ws.freeze_panes = 'D9'
    ws.print_area = f'A1:{get_column_letter(dernier)}{fin+2}'
    assurer_textes_periodes(ws)
    return True

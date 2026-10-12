"""Les taux IC et marge du premier bloc pilotent tous les blocs du même PR."""
from copy import deepcopy
import hashlib
import math
import os
from pathlib import Path
import posixpath
import re
import shutil
import tempfile
import unicodedata
from xml.etree import ElementTree as ET
from zipfile import ZipFile
from openpyxl import load_workbook

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def _signature(path, verifier_verrou=False):
    if verifier_verrou and path.with_name('~$' + path.name).exists():
        raise ValueError('Enregistrez et fermez le PR dans Excel avant de préparer ses pourcentages communs.')
    avant = path.stat()
    h = hashlib.sha256()
    with path.open('rb') as stream:
        for bloc in iter(lambda: stream.read(1024 * 1024), b''):
            h.update(bloc)
    apres = path.stat()
    if (avant.st_mtime_ns, avant.st_size, avant.st_ino) != (apres.st_mtime_ns, apres.st_size, apres.st_ino):
        raise ValueError('Le PR a changé pendant sa lecture. Relancez l’opération.')
    return apres.st_mtime_ns, apres.st_size, apres.st_ino, h.digest()


def _libelle(value):
    return ''.join(c for c in unicodedata.normalize('NFKD', str(value or '').casefold()) if c.isalnum())


def lire_pourcentages(path):
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        if 'Chiffrage' not in wb:
            return []
        rows = list(wb['Chiffrage'].iter_rows(max_col=6, values_only=True))
    finally:
        wb.close()
    postes, marges = [], set()
    for i, row in enumerate(rows):
        if _libelle(row[4]) == 'margevente':
            marges.add(i)
        if _libelle(row[4]) != 'ic':
            continue
        if i + 1 >= len(rows) or _libelle(rows[i + 1][4]) != 'margevente':
            raise ValueError(f'Récapitulatif IC / marge incomplet à la ligne {i + 1}.')
        postes.append({'installation': f'F{i + 1}', 'marge': f'F{i + 2}',
                       'valeurs': (row[5], rows[i + 1][5])})
    if marges != {int(p['marge'][1:]) - 1 for p in postes}:
        raise ValueError('Une case de marge est présente sans son récapitulatif IC.')
    if postes:
        for valeur in postes[0]['valeurs']:
            if isinstance(valeur, bool) or not isinstance(valeur, (float, int)) or not math.isfinite(valeur):
                raise ValueError('Les taux IC et marge du premier bloc doivent être des valeurs numériques saisies.')
    return postes


def _sans_resultats_formules(xml):
    def nettoyer(match):
        cellule = match[0]
        if re.search(r'<(?:\w+:)?f\b', cellule):
            cellule = re.sub(r'<(?:\w+:)?v\b[^>]*?(?:/>|>.*?</(?:\w+:)?v>)', '', cellule, flags=re.S)
        return cellule
    return re.sub(r'<(?:\w+:)?c\b[^>]*?(?:/>|>.*?</(?:\w+:)?c>)', nettoyer, xml, flags=re.S)


def _recalcul_excel(workbook):
    calc = re.search(r'<((?:\w+:)?calcPr)\b[^>]*?(?:/>|>.*?</(?:\w+:)?calcPr>)', workbook, re.S)
    nom = calc[1] if calc else re.search(r'<((?:\w+:)?workbook)\b', workbook)[1][:-8] + 'calcPr'
    attrs = re.match(r'<[^\s/>]+([^>]*?)/?>', calc[0])[1] if calc else ''
    for attribut in ('calcId', 'calcMode', 'fullCalcOnLoad', 'forceFullCalc'):
        attrs = re.sub(r'\s+' + attribut + r'="[^"]*"', '', attrs)
    ajout = f'<{nom}{attrs} calcId="0" calcMode="auto" fullCalcOnLoad="1" forceFullCalc="1"/>'
    if calc:
        return workbook[:calc.start()] + ajout + workbook[calc.end():]
    pos = re.search(r'<(?:\w+:)?(?:oleSize|customWorkbookViews|pivotCaches|smartTagPr|smartTagTypes|webPublishing|fileRecoveryPr|webPublishObjects|extLst)\b|</(?:\w+:)?workbook>', workbook)
    return workbook[:pos.start()] + ajout + workbook[pos.start():]


def _feuille_chiffrage(src):
    workbook = ET.fromstring(src.read('xl/workbook.xml'))
    feuille = next(s for s in workbook.find(f'{{{NS}}}sheets') if s.get('name') == 'Chiffrage')
    relations = ET.fromstring(src.read('xl/_rels/workbook.xml.rels'))
    target = next(r.get('Target') for r in relations if r.get('Id') == feuille.get(f'{{{REL}}}id'))
    return target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/' + target)


def _styles_saisie(src, premier):
    """Déverrouille seulement les deux entrées, sans modifier leur présentation."""
    styles = src.read('xl/styles.xml').decode('utf-8')
    root = ET.fromstring(styles)
    xfs = root.find(f'{{{NS}}}cellXfs')
    feuille = ET.fromstring(src.read(_feuille_chiffrage(src)))
    ajouts, variantes, cellules = [], {}, {}
    for ref in (premier['installation'], premier['marge']):
        cell = feuille.find(f'.//{{{NS}}}c[@r="{ref}"]')
        if cell is None:
            raise ValueError(f'Case {ref} absente du premier bloc.')
        index = int(cell.get('s', '0'))
        xf = xfs[index]
        protection = xf.find(f'{{{NS}}}protection')
        if protection is not None and protection.get('locked') in ('0', 'false'):
            continue
        if index not in variantes:
            nouveau = deepcopy(xf)
            nouveau.set('applyProtection', '1')
            protection = nouveau.find(f'{{{NS}}}protection')
            if protection is None:
                protection = ET.Element(f'{{{NS}}}protection')
                extension = nouveau.find(f'{{{NS}}}extLst')
                nouveau.insert(list(nouveau).index(extension) if extension is not None else len(nouveau), protection)
            protection.set('locked', '0')
            variantes[index] = len(xfs) + len(ajouts)
            ajouts.append(ET.tostring(nouveau, encoding='unicode'))
        cellules[ref] = variantes[index]
    if ajouts:
        debut = re.search(r'<(?:\w+:)?cellXfs\b[^>]*>', styles)
        ouverture = re.sub(r'\bcount="\d+"', f'count="{len(xfs) + len(ajouts)}"', debut[0])
        styles = styles[:debut.start()] + ouverture + styles[debut.end():]
        fin = re.search(r'</(?:\w+:)?cellXfs>', styles).start()
        styles = styles[:fin] + ''.join(ajouts) + styles[fin:]
        ET.fromstring(styles)
    return styles.encode('utf-8'), cellules


def lier_pourcentages_pr(path, sauvegarder=None):
    """Pose des références absolues internes ; aucune fenêtre ni macro nécessaire."""
    path = Path(path)
    signature = _signature(path, verifier_verrou=True)
    postes = lire_pourcentages(path)
    changements = {}
    if not postes:
        return 0
    premier = postes[0]
    for poste in postes[1:]:
        for i, cle in enumerate(('installation', 'marge')):
            formule = '=$F$' + premier[cle][1:]
            if poste['valeurs'][i] != formule:
                changements[poste[cle]] = (formule[1:], premier['valeurs'][i])
    with ZipFile(path) as src:
        styles, styles_cellules = _styles_saisie(src, premier)
    if not changements and not styles_cellules:
        return 0

    def verifier():
        if _signature(path, verifier_verrou=True) != signature:
            raise ValueError('Le PR a été modifié pendant la préparation. Aucun remplacement effectué.')

    fd, nom = tempfile.mkstemp(prefix='.pourcentages_pr_', suffix=path.suffix, dir=path.parent)
    os.close(fd)
    temporaire = Path(nom)
    try:
        with ZipFile(path) as src, ZipFile(temporaire, 'w') as dst:
            workbook = src.read('xl/workbook.xml').decode('utf-8')
            membre = _feuille_chiffrage(src)
            for info in src.infolist():
                data = src.read(info.filename)
                if info.filename.startswith('xl/worksheets/') and info.filename.endswith('.xml'):
                    xml = data.decode('utf-8')
                    if changements:
                        xml = _sans_resultats_formules(xml)
                    if info.filename == membre:
                        restantes = set(changements)
                        def remplacer(match):
                            cellule = match[0]
                            ref = re.search(r'\br="([A-Z]+\d+)"', cellule)[1]
                            if ref in styles_cellules:
                                balise = re.match(r'<[^>]*>', cellule)[0]
                                balise = re.sub(r'\s+s="\d+"', '', balise)
                                balise = balise.replace('>', f' s="{styles_cellules[ref]}">', 1)
                                return balise + cellule[cellule.index('>') + 1:]
                            if ref not in changements:
                                return cellule
                            nom, attrs = re.match(r'<((?:\w+:)?c)\b([^>]*?)/?>', cellule).groups()
                            attrs = re.sub(r'\s+t="[^"]*"', '', attrs)
                            prefixe = nom[:-1]
                            formule, valeur = changements[ref]
                            restantes.remove(ref)
                            return f'<{nom}{attrs}><{prefixe}f>{formule}</{prefixe}f><{prefixe}v>{valeur!r}</{prefixe}v></{nom}>'
                        xml = re.sub(r'<(?:\w+:)?c\b[^>]*?(?:/>|>.*?</(?:\w+:)?c>)', remplacer, xml, flags=re.S)
                        if restantes:
                            raise ValueError('Cases de pourcentage absentes : ' + ', '.join(sorted(restantes)))
                    ET.fromstring(xml)
                    data = xml.encode('utf-8')
                elif info.filename == 'xl/workbook.xml':
                    xml = _recalcul_excel(workbook)
                    ET.fromstring(xml)
                    data = xml.encode('utf-8')
                elif info.filename == 'xl/styles.xml':
                    data = styles
                dst.writestr(deepcopy(info), data)
        shutil.copymode(path, temporaire)
        verifier()
        if sauvegarder:
            sauvegarder(path)
        verifier()
        os.replace(temporaire, path)
    finally:
        temporaire.unlink(missing_ok=True)
    return len(changements) + len(styles_cellules)

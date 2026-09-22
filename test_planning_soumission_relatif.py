import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

from planning_soumission_relatif import adapter_planning_soumission


class PlanningSoumissionRelatif(unittest.TestCase):
    def test_jours_relatifs_sans_alterer_autre_feuille(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'Planning_soumission.xlsx'
            wb = Workbook()
            ws = wb.active
            ws.title = 'PlanningProjet'
            ws['C17'] = 'Début du projet'
            ws['E17'] = datetime(2026, 1, 1)
            ws['E17'].number_format = 'dd/mm/yyyy'
            ws['E18'] = 1
            ws['H18'] = '=H19'
            ws['H19'] = '=Début_Projet-WEEKDAY(Début_Projet,1)+2+7*(Semaine_Affichage-1)'
            ws['I19'] = '=H19+1'
            ws['H20'] = '=LEFT(TEXT(H19,"jjj"),1)'
            ws['E23'] = '=Début_Projet+0'
            ws['F23'] = '=E23+3'
            wb.create_sheet('À propos de')['A1'] = 'Conservé'
            from openpyxl.workbook.defined_name import DefinedName
            wb.defined_names.add(DefinedName('Début_Projet', attr_text='PlanningProjet!$E$17'))
            wb.defined_names.add(DefinedName('Semaine_Affichage', attr_text='PlanningProjet!$E$18'))
            wb.save(path)
            wb.close()
            with ZipFile(path) as z:
                autre = z.read('xl/worksheets/sheet2.xml')
            self.assertTrue(adapter_planning_soumission(path))
            self.assertFalse(adapter_planning_soumission(path))
            wb = load_workbook(path, read_only=True)
            ws = wb['PlanningProjet']
            self.assertEqual(ws['E17'].value, 0)
            self.assertEqual(ws['H19'].value, '=1+7*(Semaine_Affichage-1)')
            self.assertIsNone(ws['H20'].value)
            self.assertEqual(ws['E23'].value, '=Début_Projet+1')
            wb.close()
            with ZipFile(path) as z:
                self.assertEqual(z.read('xl/worksheets/sheet2.xml'), autre)


if __name__ == '__main__':
    unittest.main()

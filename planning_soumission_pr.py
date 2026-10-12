"""Planning Excel modifiable, établi à partir des postes de la synthèse PR."""
from datetime import date
import math
import os
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.views import Selection
from planning_soumission_periodes import (
    centraliser_periodes, est_compact, est_sur_une_ligne, lire_plages, restaurer_plages,
    assurer_couleurs_automatiques,
)

COULEURS_PERIODES = ('4F81BD', '548235', 'C55A11', '8064A2')


def compacter_debut(ws):
    """La première intervention n'affiche que le début ; ses calculs sont conservés."""
    from copy import copy
    modifie = False
    for col in ('D', 'F'):
        if not ws.column_dimensions[col].hidden:
            ws.column_dimensions[col].hidden = True
            modifie = True
    valeurs = {
        'E8': 'Début',
        'A6': 'Début du poste, puis interventions 2, 3 et 4 : saisissez les cases jaunes.',
        # La colonne F masquée contient aussi la date et le total de l'en-tête.
        'G3': '=F3',
        'G4': '=F4',
    }
    for ref, valeur in valeurs.items():
        if ws[ref].value != valeur:
            ws[ref] = valeur
            if ref in ('G3', 'G4'):
                ws[ref]._style = copy(ws[f'F{ws[ref].row}']._style)
            modifie = True
    return modifie


def appliquer_couleurs_periodes(ws):
    """Colore les règles automatiques et leurs en-têtes, sans toucher aux saisies."""
    from copy import copy
    modifie = False
    nombre = _nombre_periodes(ws)
    lettre = get_column_letter(4 + 3*nombre)
    for n in range(nombre):
        couleur = COULEURS_PERIODES[n % len(COULEURS_PERIODES)]
        remplissage = PatternFill('solid', fgColor=couleur)
        d, e = get_column_letter(4+3*n), get_column_letter(5+3*n)
        formule = (f'AND(ISNUMBER(${e}9),ISNUMBER(${d}9),${e}9>=1,${d}9>0,'
                   f'{lettre}$8>=${e}9,{lettre}$8<${e}9+${d}9)')
        for regles in ws.conditional_formatting._cf_rules.values():
            for regle in regles:
                if regle.formula == [formule] and regle.dxf is not None:
                    if regle.dxf.fill != remplissage:
                        # À la lecture, plusieurs règles peuvent partager le même style.
                        regle.dxf = copy(regle.dxf)
                        regle.dxf.fill = copy(remplissage)
                        modifie = True
        for col in range(4+3*n, 7+3*n):
            cellule = ws.cell(8, col)
            if cellule.fill != remplissage:
                cellule.fill = copy(remplissage)
                modifie = True
    return modifie


def corriger_volets(w):
    """Conserve une sélection cohérente par volet figé."""
    modifie = False
    for ws in w:
        pane = ws.sheet_view.pane
        if pane is None or not pane.xSplit or not pane.ySplit:
            continue
        col = get_column_letter(int(pane.xSplit) + 1)
        ligne = int(pane.ySplit) + 1
        selections = [Selection(pane=nom, activeCell=ref, sqref=ref)
                      for nom, ref in (('topRight', f'{col}1'), ('bottomLeft', f'A{ligne}'),
                                       ('bottomRight', f'{col}{ligne}'))]
        if ws.sheet_view.selection != selections:
            ws.sheet_view.selection = selections
            modifie = True
    return modifie


def _nombre_periodes(ws):
    nombre = 1
    while ws.cell(8, 4 + 3 * nombre).value == f'Durée {nombre + 1}':
        nombre += 1
    return nombre


def _interventions_saisies(w):
    if 'Planning soumission' in w.sheetnames and est_compact(w['Planning soumission']):
        return [(a, poste, duree, debut)
                for a, poste, _, periode, debut, duree, _ in
                w['Planning soumission'].iter_rows(min_row=9, max_col=7, values_only=True)
                if isinstance(periode, (int, float)) and periode > 1]
    if 'Interventions' in w.sheetnames:
        return [tuple(row) for row in w['Interventions'].iter_rows(
            min_row=5, max_col=4, values_only=True)
            if row[2] is not None or row[3] is not None]
    ws = w['Planning soumission']
    periodes = []
    for r in range(9, ws.max_row + 1):
        if ws.cell(r, 1).value is None or ws.cell(r, 6).data_type != 'f':
            continue
        for n in range(1, _nombre_periodes(ws)):
            duree, debut = (ws.cell(r, 4 + 3*n).value, ws.cell(r, 5 + 3*n).value)
            periodes.append((ws.cell(r, 1).value, ws.cell(r, 2).value, duree, debut))
    return periodes


def ajouter_interventions(ws, interventions=None):
    """Quatre périodes sur la ligne du poste, sans feuille auxiliaire."""
    if _nombre_periodes(ws) >= 4:
        return False
    from copy import copy
    w = ws.parent
    interventions = _interventions_saisies(w) if interventions is None else interventions
    lignes = [r for r in range(9, ws.max_row + 1)
              if ws.cell(r, 1).value is not None and ws.cell(r, 6).data_type == 'f']
    if not lignes:
        raise ValueError('Aucun poste reconnu dans le planning.')
    articles = {str(ws.cell(r, 1).value) for r in lignes}
    if any(str(a) not in articles for a, _, _, _ in interventions):
        raise ValueError('Une intervention concerne un article absent du PR. Le planning est conservé.')
    # Une ancienne saisie de plus de quatre périodes est conservée intégralement.
    nombre = max(4, 1 + max((sum(str(p[0]) == a for p in interventions)
                            for a in articles), default=0))
    jour = ws['G4'].value or 1
    for plage in list(ws.merged_cells.ranges):
        if plage.min_col >= 7:
            ws.unmerge_cells(str(plage))
    ws.delete_cols(7, max(1, ws.max_column - 6))
    for cle in list(ws.column_dimensions):
        if ws.column_dimensions[cle].min and ws.column_dimensions[cle].min >= 7:
            del ws.column_dimensions[cle]
    ws.conditional_formatting._cf_rules.clear()
    ws.data_validations.dataValidation = [v for v in ws.data_validations.dataValidation
                                          if str(v.sqref) != 'G4']
    for r in lignes:
        existantes = [p for p in interventions if str(p[0]) == str(ws.cell(r, 1).value)]
        for n in range(1, nombre):
            d, e, f = (get_column_letter(4 + 3*n + k) for k in range(3))
            for c, origine in zip((d, e, f), ('D', 'E', 'F')):
                ws[f'{c}{r}']._style = copy(ws[f'{origine}{r}']._style)
            if n <= len(existantes):
                ws[f'{d}{r}'], ws[f'{e}{r}'] = existantes[n-1][2:]
            ws[f'{f}{r}'] = (
                f'=IF(AND({d}{r}="",{e}{r}=""),"",'
                f'IF(OR(NOT(ISNUMBER({d}{r})),NOT(ISNUMBER({e}{r}))),"À compléter",'
                f'IF(OR({d}{r}<0,{e}{r}<1),"À vérifier",'
                f'IF({d}{r}=0,"",{e}{r}+{d}{r}-1))))')
    fin = max(lignes)
    for n in range(1, nombre):
        for k, titre in enumerate(('Durée', 'Début', 'Fin')):
            c = get_column_letter(4 + 3*n + k)
            ws[f'{c}8'] = f'{titre} {n+1}'
            ws[f'{c}8']._style = copy(ws.cell(8, 4+k)._style)
            ws.column_dimensions[c].width = 11
            if k < 2:
                validation = DataValidation(type='decimal' if k == 0 else 'whole',
                    operator='greaterThanOrEqual', formula1=k, allow_blank=True)
                validation.showErrorMessage = True
                ws.add_data_validation(validation)
                validation.add(f'{c}9:{c}{fin}')
    durees = [f'{get_column_letter(4+3*n)}9:{get_column_letter(4+3*n)}{fin}' for n in range(nombre)]
    fins = [f'{get_column_letter(6+3*n)}9:{get_column_letter(6+3*n)}{fin}' for n in range(nombre)]
    erreurs = '+'.join(f'COUNTIF({f},"À*")' for f in fins)
    positifs = '+'.join(f'COUNTIF({d},">0")' for d in durees)
    negatifs = '+'.join(f'COUNTIF({d},"<0")' for d in durees)
    refs = ','.join(fins)
    ws['F4'] = (f'=IF(OR(COUNT(D9:D{fin})<>COUNTA(A9:A{fin}),{erreurs}>0,'
                f'{negatifs}>0,COUNT({refs})<>{positifs}),"À compléter",MAX({refs}))')
    ws['A6'] = 'Intervention 1 : durée retenue et début. Interventions 2, 3 et 4 : cases jaunes à droite.'
    ws['A7'] = 'Répartissez la durée entre vos périodes. Les barres se colorent automatiquement sur la même ligne.'
    if 'Interventions' in w.sheetnames:
        del w['Interventions']
    for nom in ('ArticlesPlanning', 'PostesPlanning', 'PeriodesArticles', 'PeriodesDurees',
                'PeriodesDebuts', 'PeriodesFins', 'PeriodesEtat'):
        w.defined_names.pop(nom, None)
    premier = get_column_letter(4 + 3*nombre)
    ws[f'{premier}4'] = jour
    ajouter_grille_jours(ws)
    w.active = w.index(ws)
    return True


def ajouter_grille_jours(ws):
    """Barres des périodes séparées sur une seule ligne par poste."""
    nombre = _nombre_periodes(ws)
    premier = 4 + 3*nombre
    lettre = get_column_letter(premier)
    if ws.cell(8, premier).value == f'=${lettre}$4':
        return appliquer_couleurs_periodes(ws)
    lignes = [r for r in range(9, ws.max_row + 1)
              if ws.cell(r, 1).value is not None and ws.cell(r, 6).data_type == 'f']
    if not lignes:
        raise ValueError('Aucun poste reconnu pour afficher les barres du planning.')
    fin = max(lignes)
    ws.merge_cells(start_row=3, start_column=premier, end_row=3, end_column=premier+6)
    ws.cell(3, premier, 'À PARTIR DU JOUR').font = Font(name='Helvetica', size=10, bold=True, color='163A59')
    c = ws.cell(4, premier)
    if c.value is None:
        c.value = 1
    c.fill = PatternFill('solid', fgColor='FFF2CC')
    c.font = Font(name='Helvetica', size=10, color='163A59')
    c.alignment = Alignment(horizontal='center')
    ws.merge_cells(start_row=6, start_column=premier, end_row=6, end_column=premier+13)
    ws.cell(6, premier, '1 case = 1 jour — les barres suivent vos débuts et durées.').font = Font(name='Helvetica', size=10, color='163A59')
    validation = DataValidation(type='whole', operator='greaterThanOrEqual', formula1=1)
    validation.showErrorMessage = True
    ws.add_data_validation(validation)
    validation.add(f'{lettre}4')
    trait = Side(style='thin', color='C6D5E3')
    dernier = premier + 139
    for col in range(premier, dernier + 1):
        c = get_column_letter(col)
        ws.column_dimensions[c].width = 3.5
        entete = ws.cell(8, col, f'=${lettre}$4' if col == premier else f'={get_column_letter(col-1)}8+1')
        entete.number_format = '0'
        entete.font = Font(name='Helvetica', size=9, bold=True, color='FFFFFF')
        entete.fill = PatternFill('solid', fgColor='163A59')
        entete.alignment = Alignment(horizontal='center', vertical='center')
        for ligne in lignes:
            cell = ws.cell(ligne, col)
            cell.fill = PatternFill('solid', fgColor='F3F7FB' if ligne % 2 else 'FFFFFF')
            cell.border = Border(left=trait, right=trait, top=trait, bottom=trait)
    for n in range(nombre):
        d, e = get_column_letter(4+3*n), get_column_letter(5+3*n)
        formule = (f'AND(ISNUMBER(${e}9),ISNUMBER(${d}9),${e}9>=1,${d}9>0,'
                   f'{lettre}$8>=${e}9,{lettre}$8<${e}9+${d}9)')
        ws.conditional_formatting.add(f'{lettre}9:{get_column_letter(dernier)}{fin}',
            FormulaRule(formula=[formule], fill=PatternFill('solid', fgColor='4F81BD')))
    # Le poste reste visible en défilant vers les périodes et le calendrier.
    ws.freeze_panes = 'D9'
    ws.print_area = f'A1:{get_column_letter(dernier)}{fin+2}'
    ws.print_title_cols = 'A:C'
    appliquer_couleurs_periodes(ws)
    return True


def mettre_a_jour_grille(chemin, sauvegarder):
    chemin = Path(chemin)
    w = load_workbook(chemin)
    try:
        ws = w['Planning soumission']
        modifie = False
        if est_sur_une_ligne(ws):
            modifie = assurer_couleurs_automatiques(ws)
        else:
            if not est_compact(ws):
                modifie = ajouter_interventions(ws)
            modifie = centraliser_periodes(ws) or modifie
        modifie = corriger_volets(w) or modifie
        if not modifie:
            return False
        if chemin.with_name('~$' + chemin.name).exists():
            raise PermissionError('Fermez le planning dans Excel avant d’ajouter la grille de jours.')
        with TemporaryDirectory(prefix='.grille_jours_', dir=chemin.parent) as temporaire:
            nouveau = Path(temporaire) / chemin.name
            w.save(nouveau)
            sauvegarder(chemin)
            os.replace(nouveau, chemin)
    finally:
        w.close()
    return True


def _generer_planning_soumission_base(resultat, destination, nom, logo_path=None, source='PR enregistré',
                               saisies=None, interventions=None, plages=None):
    postes = resultat['postes']
    if not postes:
        raise ValueError('Aucun poste à reprendre dans le planning.')
    for poste in postes:
        jours = poste['jours']
        if jours is not None and (isinstance(jours, bool) or not isinstance(jours, (int, float))
                                  or not math.isfinite(jours) or jours < 0):
            raise ValueError(f"Durée PR invalide pour {poste['article']}.")
    destination = Path(destination)
    if destination.exists():
        raise FileExistsError('Le planning existe déjà : ouvrez-le pour conserver vos saisies.')
    wb = Workbook()
    ws = wb.active
    ws.title = 'Planning soumission'
    wb.calculation = CalcProperties(calcMode='auto', fullCalcOnLoad=True)
    bleu, pale, jaune = '163A59', 'F3F7FB', 'FFF2CC'
    trait = Side(style='thin', color='C6D5E3')
    for col, largeur in zip('ABCDEF', (21, 78, 15, 17, 17, 17)):
        ws.column_dimensions[col].width = largeur
    ws.sheet_view.showGridLines = False
    ws.merge_cells('B1:F2')
    ws['B1'] = 'PLANNING DE SOUMISSION'
    ws['B1'].font = Font(name='Helvetica', size=18, bold=True, color=bleu)
    ws['B1'].alignment = Alignment(vertical='center', horizontal='center')
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 24
    if logo_path is None or not Path(logo_path).is_file():
        logo_path = Path(__file__).with_name('logo_jt_bati_planning.jpg')
    if Path(logo_path).is_file():
        logo = Image(str(logo_path))
        logo.height = logo.height * 140 / logo.width
        logo.width = 140
        ws.add_image(logo, 'A1')
    ws['A3'], ws['B3'] = 'CHANTIER', str(nom)
    ws['E3'], ws['F3'] = 'DATE', date.today()
    ws['F3'].number_format = 'dd/mm/yyyy'
    ws['A4'], ws['B4'] = 'COMMENCEMENT', 'Jour 1 — sans date de commande'
    ws['E4'] = 'DÉLAI FINAL'
    debut, fin = 9, 8 + len(postes)
    # Le délai reste incomplet si une durée ou un début nécessaire manque.
    ws['F4'] = (f'=IF(COUNT(D{debut}:D{fin})<>COUNTA(A{debut}:A{fin}),"À compléter",'
                f'IF(COUNTIF(D{debut}:D{fin},"<0")>0,"À vérifier",'
                f'IF(COUNT(F{debut}:F{fin})<>COUNTIF(D{debut}:D{fin},">0"),"À compléter",'
                f'IF(COUNTIF(D{debut}:D{fin},">0")=0,0,MAX(F{debut}:F{fin})))))')
    ws['F4'].number_format = '0.00" jours"'
    ws.merge_cells('A6:F6')
    ws['A6'] = 'Cases jaunes : durée retenue et jour de début. Vous choisissez les chevauchements.'
    ws.merge_cells('A7:F7')
    ws['A7'] = 'Fin = début + durée − 1. Exemple : 8 jours à partir du Jour 1 → fin au Jour 8.'
    for ref in ('A6', 'A7'):
        ws[ref].font = Font(name='Helvetica', size=10, color=bleu)
    entetes = ('Article', 'Poste', 'Jours PR', 'Durée retenue', 'Début (jour)', 'Fin (jour)')
    for col, titre in enumerate(entetes, 1):
        c = ws.cell(8, col, titre)
        c.fill = PatternFill('solid', fgColor=bleu)
        c.font = Font(name='Helvetica', size=10, bold=True, color='FFFFFF')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for ligne, poste in enumerate(postes, debut):
        ws.cell(ligne, 1, str(poste['article'])).data_type = 's'
        ws.cell(ligne, 2, str(poste['designation'])).data_type = 's'
        ws.cell(ligne, 3, poste['jours'])
        if poste['jours'] is not None:
            ws.cell(ligne, 4, f'=C{ligne}')
        if ligne == debut:
            ws.cell(ligne, 5, 1)
        if saisies and str(poste['article']) in saisies:
            duree, commencement = saisies[str(poste['article'])]
            ws.cell(ligne, 5).value = commencement
            if duree is not None:
                ws.cell(ligne, 4).value = duree
        ws.cell(ligne, 6,
                f'=IF(OR(D{ligne}="",E{ligne}="",D{ligne}=0),"",'
                f'IF(OR(D{ligne}<0,E{ligne}<1),"À vérifier",E{ligne}+D{ligne}-1))')
        ws.row_dimensions[ligne].height = max(32, 15 * math.ceil(len(str(poste['designation'])) / 70))
        for col in range(1, 7):
            c = ws.cell(ligne, col)
            c.font = Font(name='Helvetica', size=10, color=bleu)
            c.fill = PatternFill('solid', fgColor=jaune if col in (4, 5) else (pale if ligne % 2 else 'FFFFFF'))
            c.border = Border(left=trait, right=trait, top=trait, bottom=trait)
            c.alignment = Alignment(vertical='center', wrap_text=True,
                                    horizontal='left' if col <= 2 else 'center')
            if col >= 3:
                c.number_format = '0.00'
    for col, minimum, type_saisie in (('D', 0, 'decimal'), ('E', 1, 'whole')):
        validation = DataValidation(type=type_saisie, operator='greaterThanOrEqual',
                                    formula1=minimum, allow_blank=True)
        validation.showErrorMessage = True
        validation.error = f'Saisissez un nombre supérieur ou égal à {minimum}.'
        validation.errorTitle = 'Saisie à vérifier'
        ws.add_data_validation(validation)
        validation.add(f'{col}{debut}:{col}{fin}')
    for ligne in (3, 4):
        ws.row_dimensions[ligne].height = 24
        for c in ws[ligne]:
            c.font = Font(name='Helvetica', size=10, color=bleu, bold=c.column in (1, 5, 6))
            c.fill = PatternFill('solid', fgColor='D9EAD3' if c.coordinate == 'F4' else pale)
            c.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[8].height = 30
    ws.merge_cells(start_row=fin+2, start_column=1, end_row=fin+2, end_column=6)
    ws.cell(fin+2, 1, f'Source : {source}. Les jours PR restent la référence ; les durées retenues sont modifiables.')
    ws.cell(fin+2, 1).font = Font(name='Helvetica', size=9, color=bleu)
    ws.cell(fin+2, 1).alignment = Alignment(wrap_text=True)
    ws.row_dimensions[fin+2].height = 28
    ws.freeze_panes = 'C9'
    ws.print_title_rows = '1:8'
    ws.print_area = f'A1:F{fin+2}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=.25, right=.25, top=.35, bottom=.35, header=.15, footer=.15)
    ws.oddFooter.center.text = 'Page &P / &N'
    ws.oddFooter.right.text = str(nom).replace('&', '&&')
    if plages:
        # Conserve également les rares anciens plannings de plus de quatre périodes.
        interventions = [(a, '', None, None) for a, cellules in plages.items()
                         for _ in range(len(cellules)//2-1)]
    ajouter_interventions(ws, interventions)
    centraliser_periodes(ws)
    if plages is not None:
        restaurer_plages(ws, plages)
    corriger_volets(wb)
    # Mode exclusif : une seconde création ne peut pas écraser le travail saisi.
    cree = False
    try:
        with destination.open('xb') as fichier:
            cree = True
            wb.save(fichier)
    except Exception:
        if cree:
            destination.unlink(missing_ok=True)
        raise
    finally:
        wb.close()
    return destination


def _lire_saisies(chemin, nom):
    """Identifie notre tableau, puis reprend seulement les choix de l'utilisateur."""
    w = load_workbook(chemin, read_only=True, data_only=False)
    try:
        s = w['Planning soumission'] if 'Planning soumission' in w.sheetnames else w.active
        if s['B3'].value == nom and est_compact(s):
            return {str(a.value): (None if duree.value == f'=C{duree.row}' else duree.value, debut.value)
                    for a, _, _, periode, debut, duree, _ in s.iter_rows(min_row=9, max_col=7)
                    if periode.value == 1 and a.value is not None}
        if s['B3'].value != nom or s['D8'].value != 'Durée retenue':
            return None
        saisies = {}
        for row in s.iter_rows(min_row=9, max_col=6):
            article, _, jours, duree, debut, fin = row
            if article.value is None or fin.data_type != 'f':
                continue
            valeur = None if duree.value == f'=C{duree.row}' else duree.value
            saisies[str(article.value)] = (valeur, debut.value)
        return saisies
    finally:
        w.close()


def _lire_interventions(chemin):
    w = load_workbook(chemin, read_only=True, data_only=False)
    try:
        return _interventions_saisies(w)
    finally:
        w.close()


def _assurer_planning_soumission_base(pr, destination, nom, logo_path, sauvegarder, actualiser=False):
    """Remplace une fois l'ancien planning au nom habituel, puis conserve les saisies."""
    from synthese_pr import lire_synthese_pr
    destination = Path(destination)
    plages = None
    if destination.is_file():
        with_plages = load_workbook(destination, read_only=True)
        try:
            if 'Planning soumission' in with_plages:
                ws = with_plages['Planning soumission']
                if ws['B3'].value == nom and est_sur_une_ligne(ws):
                    plages = lire_plages(ws)
        finally:
            with_plages.close()
        if plages is not None and not actualiser:
            mettre_a_jour_grille(destination, sauvegarder)
            return destination
    saisies = _lire_saisies(destination, nom) if destination.is_file() else None
    interventions = _lire_interventions(destination) if saisies is not None else []
    if saisies is not None and not actualiser:
        mettre_a_jour_grille(destination, sauvegarder)
        return destination
    if destination.with_name('~$' + destination.name).exists():
        raise PermissionError('Enregistrez et fermez le planning de soumission dans Excel avant sa mise à jour.')
    if saisies is None and plages is None:
        # Reprise du document créé sous le mauvais nom dans la version précédente.
        precedent = destination.with_name('Planning_execution_PR.xlsx')
        if precedent.is_file():
            saisies = _lire_saisies(precedent, nom)
            if saisies is not None:
                interventions = _lire_interventions(precedent)
    resultat = lire_synthese_pr(pr)
    with TemporaryDirectory(prefix='.planning_soumission_', dir=destination.parent) as temporaire:
        nouveau = Path(temporaire) / destination.name
        _generer_planning_soumission_base(resultat, nouveau, nom, logo_path, saisies=saisies,
                                   interventions=interventions, plages=plages)
        if destination.exists():
            sauvegarder(destination)
        os.replace(nouveau, destination)
    return destination


def generer_planning_soumission(resultat, destination, nom, logo_path=None, source='PR enregistré',
                               saisies=None, interventions=None, plages=None):
    """Crée le planning avec sélecteur de période et mémorisation dans Excel."""
    from planning_soumission_menu import installer_menu, fermer_planning
    from planning_soumission_legendes import ecrire_legendes
    destination = Path(destination).with_suffix('.xlsm')
    if destination.exists():
        raise FileExistsError('Le planning existe déjà : ouvrez-le pour conserver vos saisies.')
    with TemporaryDirectory(prefix='.planning_menu_', dir=destination.parent) as temporaire:
        base = Path(temporaire)/'base.xlsx'
        _generer_planning_soumission_base(resultat, base, nom, logo_path, source,
                                          saisies, interventions, plages)
        wb = load_workbook(base)
        try:
            installer_menu(wb['Planning soumission'])
            corriger_volets(wb)
            nouveau = Path(temporaire)/destination.name
            wb.save(nouveau)
            ecrire_legendes(nouveau)
        finally:
            fermer_planning(wb)
        # Création exclusive : ne pas écraser un planning apparu entre-temps.
        cree = False
        try:
            with destination.open('xb') as fichier:
                cree = True
                fichier.write(nouveau.read_bytes())
        except Exception:
            if cree:
                destination.unlink(missing_ok=True)
            raise
    return destination


def assurer_planning_soumission(pr, destination, nom, logo_path, sauvegarder, actualiser=False):
    from planning_soumission_legendes import ecrire_legendes
    from planning_soumission_menu import (
        est_menu, installer_menu, lire_plages_menu, fermer_planning, META,
        assurer_textes_periodes, lire_textes_menu, restaurer_textes_menu, selection_menu,
    )
    from synthese_pr import lire_synthese_pr
    ancien = Path(destination).with_suffix('.xlsx')
    destination = ancien.with_suffix('.xlsm')
    if destination.is_file():
        wb = load_workbook(destination, keep_vba=True)
        try:
            ws = wb['Planning soumission']
            if not est_menu(ws) or ws['B3'].value != nom:
                raise ValueError('Le planning existant ne correspond pas au chantier. Il est conservé.')
            if not actualiser:
                if assurer_textes_periodes(ws):
                    if destination.with_name('~$'+destination.name).exists():
                        raise PermissionError('Enregistrez et fermez le planning dans Excel avant sa mise à jour.')
                    with TemporaryDirectory(prefix='.planning_texte_', dir=destination.parent) as temporaire:
                        nouveau = Path(temporaire)/destination.name
                        wb.save(nouveau)
                        ecrire_legendes(nouveau)
                        sauvegarder(destination)
                        os.replace(nouveau, destination)
                return destination
            plages = lire_plages_menu(ws)
            textes = lire_textes_menu(ws)
            selections = {str(ws.cell(r, 1).value): selection_menu(ws, r)
                          for r in range(9, int(ws.cell(2, META).value)+1)}
            jour = ws['G4'].value
        finally:
            fermer_planning(wb)
        if destination.with_name('~$'+destination.name).exists():
            raise PermissionError('Enregistrez et fermez le planning dans Excel avant sa mise à jour.')
        with TemporaryDirectory(prefix='.planning_menu_', dir=destination.parent) as temporaire:
            nouveau = generer_planning_soumission(lire_synthese_pr(pr), Path(temporaire)/destination.name,
                                                  nom, logo_path, plages=plages)
            wb = load_workbook(nouveau, keep_vba=True)
            try:
                ws = wb['Planning soumission']
                ws['G4'] = jour
                for r in range(9, int(ws.cell(2, META).value)+1):
                    selection = selections.get(str(ws.cell(r, 1).value), 1)
                    ws.cell(r, 4, selection)
                    ws.cell(r, META, selection)
                    for i in range(2):
                        ref = f'{get_column_letter(160+2*(selection-1)+i)}{r}'
                        ws.cell(r, 5+i, f'=IF({ref}="","",{ref})')
                restaurer_textes_menu(ws, textes)
                wb.save(nouveau)
                ecrire_legendes(nouveau)
            finally:
                fermer_planning(wb)
            sauvegarder(destination)
            os.replace(nouveau, destination)
        return destination
    if not ancien.is_file():
        return generer_planning_soumission(lire_synthese_pr(pr), destination, nom, logo_path)
    if ancien.with_name('~$'+ancien.name).exists():
        raise PermissionError('Enregistrez et fermez le planning dans Excel avant sa mise à jour.')
    # Travaille sur une copie ; l'ancien document reste intact jusqu'à la sauvegarde finale.
    import shutil
    with TemporaryDirectory(prefix='.planning_menu_', dir=ancien.parent) as temporaire:
        copie = Path(temporaire)/ancien.name
        shutil.copy2(ancien, copie)
        _assurer_planning_soumission_base(pr, copie, nom, logo_path, lambda _: None, actualiser)
        wb = load_workbook(copie)
        try:
            installer_menu(wb['Planning soumission'])
            corriger_volets(wb)
            nouveau = Path(temporaire)/destination.name
            wb.save(nouveau)
            ecrire_legendes(nouveau)
        finally:
            fermer_planning(wb)
        sauvegarder(ancien)
        if destination.exists():
            raise FileExistsError('Un planning a été créé entre-temps. Les deux fichiers sont conservés.')
        os.replace(nouveau, destination)
        ancien.unlink()
    return destination
